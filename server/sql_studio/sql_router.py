from __future__ import annotations

import os
from io import BytesIO
from typing import Any, Dict, List, Optional

import pandas as pd
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from ..datasets.dataset_registry import create_dataset_from_dataframe, link_project
from ..security.security_store import audit_event, require_user
from .sql_engine import run_preview_sql, validate_sql
from .sql_store import delete_query, get_query, list_queries, save_query

router = APIRouter(dependencies=[Depends(require_user)])
ALLOW_SAMPLE_SQL_REGISTRATION = os.getenv("ATLAS_ALLOW_SAMPLE_SQL_REGISTRATION", "false").strip().lower() in {"1", "true", "yes"}


class DatasetRef(BaseModel):
    datasetId: str
    versionId: str
    alias: str


class SqlValidationRequest(BaseModel):
    sql: str
    sourceType: str = "dataset_library_preview"
    datasetRefs: List[DatasetRef] = Field(default_factory=list)
    limit: int = 100


class SqlPreviewRequest(SqlValidationRequest):
    pass


class SaveQueryRequest(BaseModel):
    queryId: Optional[str] = None
    name: str
    description: str = ""
    sql: str
    sourceType: str = "dataset_library_preview"
    datasetRefs: List[Dict[str, Any]] = Field(default_factory=list)
    parameters: List[Dict[str, Any]] = Field(default_factory=list)
    linkedProjectId: str = ""
    tags: List[str] = Field(default_factory=list)


class RegisterSqlResultRequest(BaseModel):
    queryName: str
    sql: str
    sourceType: str = "dataset_library_preview"
    datasetRefs: List[Dict[str, Any]] = Field(default_factory=list)
    resultRows: List[Dict[str, Any]] = Field(default_factory=list)
    resultColumns: List[Dict[str, Any]] = Field(default_factory=list)
    description: str = ""
    linkedProjectId: str = ""
    tags: List[str] = Field(default_factory=list)


class ExportSqlResultRequest(BaseModel):
    queryName: str = "SQL Data Studio Result"
    sql: str
    rows: List[Dict[str, Any]] = Field(default_factory=list)
    columns: List[Dict[str, Any]] = Field(default_factory=list)
    schemaProfile: List[Dict[str, Any]] = Field(default_factory=list)
    datasetRefs: List[Dict[str, Any]] = Field(default_factory=list)
    warnings: List[str] = Field(default_factory=list)
    lineage: Dict[str, Any] = Field(default_factory=dict)


@router.post("/sql-studio/validate")
def post_validate(payload: SqlValidationRequest) -> Dict[str, Any]:
    return validate_sql(payload.sql, [ref.model_dump() for ref in payload.datasetRefs], payload.limit)


@router.post("/sql-studio/run-preview")
def post_run_preview(payload: SqlPreviewRequest) -> Dict[str, Any]:
    if payload.sourceType != "dataset_library_preview":
        return {
            "ok": False,
            "status": "blocked",
            "rowCount": 0,
            "returnedRows": 0,
            "columns": [],
            "rows": [],
            "runtimeMs": 0,
            "warnings": [],
            "error": "Connector not configured in this local build. Use Dataset Library Preview SQL or register source metadata.",
        }
    return run_preview_sql(payload.sql, [ref.model_dump() for ref in payload.datasetRefs], payload.limit)


@router.get("/sql-studio/queries")
def get_queries() -> Dict[str, Any]:
    return {"ok": True, "queries": list_queries()}


@router.post("/sql-studio/queries")
def post_query(payload: SaveQueryRequest) -> Dict[str, Any]:
    return {"ok": True, "query": save_query(payload.model_dump())}


@router.get("/sql-studio/queries/{query_id}")
def get_one_query(query_id: str) -> Dict[str, Any]:
    query = get_query(query_id)
    if not query:
        raise HTTPException(status_code=404, detail="Saved SQL query not found.")
    return {"ok": True, "query": query}


@router.delete("/sql-studio/queries/{query_id}")
def delete_one_query(query_id: str) -> Dict[str, Any]:
    return delete_query(query_id)


@router.post("/sql-studio/register-result")
def post_register_result(payload: RegisterSqlResultRequest) -> Dict[str, Any]:
    if payload.sourceType == "dataset_library_preview" and not ALLOW_SAMPLE_SQL_REGISTRATION:
        raise HTTPException(
            status_code=409,
            detail="SQL preview results are computed from sample rows and cannot be registered as governed datasets. Configure a full-data SQL connector before registration.",
        )
    if not payload.resultRows:
        raise HTTPException(status_code=400, detail="Run a SQL preview before registering a result dataset.")
    df = pd.DataFrame(payload.resultRows)
    source_info = {
        "datasetRole": "sql_result",
        "sourceType": "sql_studio",
        "sqlQuery": payload.sql,
        "sourceDatasets": payload.datasetRefs,
        "queryName": payload.queryName,
        "linkedProjectId": payload.linkedProjectId,
        "rowCount": int(df.shape[0]),
        "columnCount": int(df.shape[1]),
        "previewLimitation": "Phase 12A registers preview/sample SQL results only.",
        "tags": payload.tags,
    }
    dataset = create_dataset_from_dataframe(
        df=df,
        name=f"SQL Result - {payload.queryName}",
        description=payload.description or "SQL Data Studio governed result dataset. Review source lineage before modeling or reporting.",
        domain="SQL Data Studio",
        source_type="sql_studio",
        source_info=source_info,
        dataset_role="sql_result",
    )
    version_id = dataset.get("latestVersion", {}).get("versionId")
    if payload.linkedProjectId and version_id:
        link_project(payload.linkedProjectId, dataset["datasetId"], version_id, "sql_result")
    audit_event("sql_studio.result.registered", "dataset", dataset.get("datasetId"), "success", metadata=source_info)
    return {"ok": True, "dataset": dataset, "lineage": source_info}


@router.post("/sql-studio/export/excel")
def post_export_excel(payload: ExportSqlResultRequest) -> StreamingResponse:
    try:
        from openpyxl import Workbook
    except Exception as exc:
        raise HTTPException(status_code=503, detail="openpyxl is required for SQL Data Studio Excel export.") from exc
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Query Name", payload.queryName])
    summary.append(["Rows", len(payload.rows)])
    summary.append(["Columns", len(payload.columns)])
    summary.append(["Limitation", "Preview/sample-row execution only in Phase 12A."])
    sql_sheet = workbook.create_sheet("SQL")
    sql_sheet.append(["SQL"])
    sql_sheet.append([payload.sql])
    result_sheet = workbook.create_sheet("Result Preview")
    headers = [str(column.get("name") or column) for column in payload.columns] or sorted({key for row in payload.rows for key in row.keys()})
    result_sheet.append(headers)
    for row in payload.rows:
        result_sheet.append([row.get(header) for header in headers])
    schema_sheet = workbook.create_sheet("Schema Profile")
    schema_sheet.append(["Column", "Type", "Missing", "Distinct", "Min", "Max", "Mean", "Sample Values"])
    for column in payload.schemaProfile:
        schema_sheet.append([
            column.get("name"),
            column.get("type"),
            column.get("missingCount"),
            column.get("distinctCount"),
            column.get("min"),
            column.get("max"),
            column.get("mean"),
            ", ".join(map(str, column.get("sampleValues") or [])),
        ])
    source_sheet = workbook.create_sheet("Source Datasets")
    source_sheet.append(["Dataset ID", "Version ID", "Alias"])
    for ref in payload.datasetRefs:
        source_sheet.append([ref.get("datasetId"), ref.get("versionId"), ref.get("alias")])
    lineage_sheet = workbook.create_sheet("Lineage")
    lineage_sheet.append(["Key", "Value"])
    for key, value in (payload.lineage or {}).items():
        lineage_sheet.append([key, str(value)])
    warnings_sheet = workbook.create_sheet("Warnings")
    warnings_sheet.append(["Warning"])
    for warning in payload.warnings:
        warnings_sheet.append([warning])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="atlas-sql-data-studio-result.xlsx"'},
    )
