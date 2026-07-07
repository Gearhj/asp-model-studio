from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any, Dict, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from ..security.security_store import audit_event, require_user
from .chart_store import chart_diagnostics, create_chart_insight, list_chart_insights


router = APIRouter(dependencies=[Depends(require_user)])


class ChartInsightRequest(BaseModel):
    title: str = "Untitled Chart"
    summary: str = ""
    datasetId: str
    datasetName: str = ""
    versionId: str = ""
    projectId: Optional[str] = None
    linkedProjectId: Optional[str] = None
    businessQuestion: str = ""
    chartType: str = "bar"
    roles: Dict[str, Any] = Field(default_factory=dict)
    chartData: list[Dict[str, Any]] = Field(default_factory=list)
    interpretation: str = ""
    generatedPython: str = ""
    generatedSql: str = ""
    generatedR: str = ""
    generatedSas: str = ""
    metadata: Dict[str, Any] = Field(default_factory=dict)


class ChartExportRequest(BaseModel):
    chartTitle: str = "ATLAS Chart"
    chartType: str = "bar"
    datasetName: str = "selected dataset"
    datasetId: str = ""
    versionId: str = ""
    businessQuestion: str = ""
    roles: Dict[str, Any] = Field(default_factory=dict)
    chartData: list[Dict[str, Any]] = Field(default_factory=list)
    interpretation: str = ""
    generatedPython: str = ""
    generatedSql: str = ""
    generatedR: str = ""
    generatedSas: str = ""
    metadata: Dict[str, Any] = Field(default_factory=dict)


@router.get("/charts/insights")
def get_chart_insights(datasetId: Optional[str] = None, projectId: Optional[str] = None) -> Dict[str, Any]:
    return {"ok": True, "insights": list_chart_insights(dataset_id=datasetId, project_id=projectId)}


@router.post("/charts/insights")
def post_chart_insight(payload: ChartInsightRequest) -> Dict[str, Any]:
    insight = create_chart_insight(payload.model_dump())
    audit_event(
        "chart.insight.saved",
        "dataset",
        payload.datasetId,
        "success",
        metadata={"insightId": insight.get("insightId"), "chartType": payload.chartType, "projectId": payload.projectId or payload.linkedProjectId},
    )
    return {"ok": True, "insight": insight}


@router.get("/charts/diagnostics")
def get_chart_diagnostics() -> Dict[str, Any]:
    return chart_diagnostics()


@router.post("/charts/export/excel")
def export_chart_excel(payload: ChartExportRequest) -> StreamingResponse:
    try:
        content = build_chart_excel(payload.model_dump())
    except ImportError as exc:
        raise HTTPException(status_code=503, detail="Excel export requires openpyxl. Install server requirements and retry.") from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"ATLAS could not create the chart Excel export: {exc}") from exc
    filename = safe_export_filename(payload.chartTitle or payload.datasetName or "atlas-chart", "xlsx")
    audit_event(
        "chart.export.created",
        "dataset",
        payload.datasetId,
        "success",
        metadata={"format": "excel", "chartType": payload.chartType, "versionId": payload.versionId},
    )
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def build_chart_excel(payload: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    write_pairs(
        summary,
        [
            ("Chart Title", payload.get("chartTitle")),
            ("Chart Type", payload.get("chartType")),
            ("Dataset", payload.get("datasetName")),
            ("Dataset ID", payload.get("datasetId")),
            ("Version", payload.get("versionId")),
            ("Business Question", payload.get("businessQuestion")),
            ("Interpretation", payload.get("interpretation")),
            ("Preview Limitation", "Chart Studio v1 exports chart data computed from Dataset Library preview/sample rows."),
        ],
    )

    write_table(workbook.create_sheet("Chart Data"), payload.get("chartData") or [])
    write_pairs(workbook.create_sheet("Role Assignments"), sorted((payload.get("roles") or {}).items()))
    code = workbook.create_sheet("Generated Code")
    write_pairs(
        code,
        [
            ("Python", payload.get("generatedPython")),
            ("SQL", payload.get("generatedSql")),
            ("R", payload.get("generatedR")),
            ("SAS-compatible", payload.get("generatedSas")),
        ],
    )
    write_pairs(workbook.create_sheet("Interpretation"), [("ATLAS Interpretation", payload.get("interpretation"))])
    write_pairs(workbook.create_sheet("Metadata"), sorted((payload.get("metadata") or {}).items()))

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="0F172A")
            cell.fill = PatternFill("solid", fgColor="E2E8F0")
        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column_cells[:100])
            sheet.column_dimensions[column_letter].width = min(max(14, max_length + 2), 60)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def write_pairs(sheet: Any, pairs: list[tuple[Any, Any]]) -> None:
    sheet.append(["Field", "Value"])
    for key, value in pairs:
        sheet.append([str(key), serialize_cell(value)])


def write_table(sheet: Any, rows: list[Dict[str, Any]]) -> None:
    columns = sorted({key for row in rows for key in row.keys()}) or ["No chart data"]
    sheet.append(columns)
    if not rows:
        sheet.append(["No chart data was available."])
        return
    for row in rows:
        sheet.append([serialize_cell(row.get(column)) for column in columns])


def serialize_cell(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=True)


def safe_export_filename(value: str, extension: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", value or "atlas-chart").strip("-._")
    if not safe:
        safe = "atlas-chart"
    return f"{safe[:80]}-chart-data.{extension}"
