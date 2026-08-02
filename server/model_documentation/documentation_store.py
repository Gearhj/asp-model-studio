from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List

from ..analytics.analytics_store import list_insights as list_analytics_insights
from ..datasets.dataset_registry import load_registry, project_dataset_links, write_json
from ..persistence import get_persistence_service
from ..projects.project_registry import list_backend_projects
from .mdd_framework import DOCUMENT_STATUSES, family_for_project, template_sections


DOCUMENTATION_ROOT = Path(__file__).resolve().parents[1] / "artifacts" / "model_documentation"
DOCUMENT_INDEX_PATH = DOCUMENTATION_ROOT / "document_index.json"

TEMPLATE_TYPES = [
    {"templateType": "validator_aligned_mdd", "label": "Validator-Aligned Model Development Document", "status": "available", "description": "Evidence-linked common MDD backbone with method-specific validator requirements and governed document lifecycle."},
    {
        "templateType": "model_development_document",
        "label": "Model Development Document",
        "status": "available",
        "description": "Governed model development document with workflow evidence, diagnostics, and implementation readiness.",
    },
    {
        "templateType": "optimization_decision_report",
        "label": "Optimization Decision Report",
        "status": "available",
        "description": "Audit-ready optimization report covering inputs, formulation, solver evidence, scenarios, sensitivity, validation, recommendation, implementation, monitoring, and reproducibility.",
    },
    {
        "templateType": "validation_evidence_pack",
        "label": "Model Validation Evidence Pack",
        "status": "coming_soon",
        "description": "Validation-focused evidence package for independent model review.",
    },
    {
        "templateType": "monitoring_handoff",
        "label": "Model Monitoring Handoff",
        "status": "coming_soon",
        "description": "Monitoring setup, metrics, cadence, and escalation evidence.",
    },
    {
        "templateType": "executive_summary",
        "label": "Executive Summary",
        "status": "coming_soon",
        "description": "Senior leadership summary deck and one-page narrative.",
    },
    {
        "templateType": "technical_appendix",
        "label": "Technical Appendix",
        "status": "coming_soon",
        "description": "Detailed methods, code evidence, and technical diagnostics.",
    },
]

MODEL_DEVELOPMENT_OUTLINE = [
    ("executive_summary", "Executive Summary", True),
    ("business_objective", "Business Objective", True),
    ("model_use_case", "Model Use Case", True),
    ("target_definition", "Target Definition", True),
    ("data_sources", "Data Sources", True),
    ("population_definition", "Population Definition", True),
    ("inclusion_exclusion", "Inclusion and Exclusion Criteria", True),
    ("development_test_oot", "Development, Test, and OOT Data", True),
    ("eda", "Exploratory Data Analysis", True),
    ("missing_value_treatment", "Missing Value Treatment", True),
    ("outlier_treatment", "Outlier Treatment", True),
    ("feature_engineering", "Feature Engineering", True),
    ("correlation_vif", "Correlation and VIF Review", True),
    ("model_training_strategy", "Model Training Strategy", True),
    ("candidate_models", "Candidate Models", True),
    ("champion_model_selection", "Champion Model Selection", True),
    ("performance_summary", "Model Performance Summary", True),
    ("evaluation_diagnostics", "Evaluation Diagnostics", True),
    ("model_interpretation", "Model Interpretation", True),
    ("scoring_external_validation", "Scoring and External Validation", True),
    ("model_diagnostics", "Model Diagnostics", True),
    ("fair_lending_governance", "Fair Lending and Governance Considerations", True),
    ("implementation_plan", "Implementation Plan", True),
    ("monitoring_plan", "Monitoring Plan", True),
    ("limitations", "Limitations", True),
    ("appendices", "Appendices", False),
]

OPTIMIZATION_DECISION_OUTLINE = [
    ("optimization_executive_summary", "Executive Summary", True),
    ("optimization_business_decision", "Business Decision", True),
    ("optimization_data_inputs", "Data Inputs", True),
    ("optimization_formulation_readiness", "Formulation Readiness", True),
    ("optimization_objective_function", "Objective Function", True),
    ("optimization_decision_variables", "Decision Variables", True),
    ("optimization_constraints", "Constraints", True),
    ("optimization_solver_configuration", "Solver Configuration", True),
    ("optimization_baseline_scenarios", "Baseline and Scenarios", True),
    ("optimization_results", "Optimization Results", True),
    ("optimization_sensitivity_tradeoffs", "Sensitivity and Tradeoffs", True),
    ("optimization_validation", "Validation", True),
    ("optimization_recommendation", "Recommendation", True),
    ("optimization_implementation_plan", "Implementation Plan", True),
    ("optimization_monitoring_plan", "Monitoring Plan", True),
    ("optimization_limitations", "Limitations", True),
    ("optimization_approvals", "Approvals", True),
    ("optimization_reproducibility_appendix", "Reproducibility Appendix", False),
    ("optimization_python_code", "Python Code", False),
    ("optimization_sas_code", "SAS-Compatible Code", False),
]


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def ensure_store() -> None:
    DOCUMENTATION_ROOT.mkdir(parents=True, exist_ok=True)
    if not DOCUMENT_INDEX_PATH.exists():
        write_json(DOCUMENT_INDEX_PATH, {"documents": []})


def read_json(path: Path, default: Any) -> Any:
    try:
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default
    return default


def document_path(document_id: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", document_id).strip("-") or "document"
    return DOCUMENTATION_ROOT / f"{safe}.json"


def load_document(document_id: str) -> Dict[str, Any]:
    document = read_json(document_path(document_id), {})
    if not document:
        raise KeyError(document_id)
    return document


def save_document(document: Dict[str, Any]) -> Dict[str, Any]:
    ensure_store()
    document["updatedAt"] = now_iso()
    write_json(document_path(str(document["documentId"])), document)
    index = read_json(DOCUMENT_INDEX_PATH, {"documents": []})
    records = [item for item in index.get("documents", []) if item.get("documentId") != document.get("documentId")]
    records.insert(
        0,
        {
            "documentId": document.get("documentId"),
            "projectId": document.get("projectId"),
            "projectName": document.get("projectName"),
            "templateType": document.get("templateType"),
            "createdAt": document.get("createdAt"),
            "updatedAt": document.get("updatedAt"),
            "quality": document.get("quality", {}),
        },
    )
    write_json(DOCUMENT_INDEX_PATH, {"documents": records[:250]})
    return document


def list_templates() -> List[Dict[str, Any]]:
    return TEMPLATE_TYPES


def list_documents(project_id: str | None = None) -> List[Dict[str, Any]]:
    ensure_store()
    documents = read_json(DOCUMENT_INDEX_PATH, {"documents": []}).get("documents", [])
    if project_id:
        documents = [item for item in documents if item.get("projectId") == project_id]
    return documents


def _mapping(value: Any) -> Dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _first_nonempty(
    sources: List[Dict[str, Any]],
    *keys: str,
    default: Any = "",
) -> Any:
    for source in sources:
        for key in keys:
            value = source.get(key)
            if value not in (None, "", [], {}):
                return value
    return default


def _relational_project_context(
    relational_project: Dict[str, Any],
    project_id: str,
) -> Dict[str, Any]:
    metadata = _mapping(relational_project.get("metadata_json"))
    snapshot = _mapping(metadata.get("snapshot"))
    snapshot_project = _mapping(snapshot.get("project"))
    project_context = _mapping(metadata.get("projectContext"))
    snapshot_context = _mapping(snapshot.get("projectContext"))
    fraud_details = _mapping(project_context.get("fraudDetails"))
    snapshot_fraud_details = _mapping(snapshot_context.get("fraudDetails"))
    sources = [
        metadata,
        snapshot,
        snapshot_project,
        project_context,
        snapshot_context,
    ]

    linked_datasets = _first_nonempty(
        sources,
        "linkedDatasets",
        default=[],
    )
    if not isinstance(linked_datasets, list):
        linked_datasets = []

    return {
        "projectId": project_id,
        "name": relational_project.get("project_name")
        or "Untitled ATLAS Project",
        "projectType": relational_project.get("project_type")
        or relational_project.get("workflow_type")
        or "unknown",
        "taskType": _first_nonempty(
            sources,
            "taskType",
            "task",
            default="unknown",
        ),
        "target": _first_nonempty(
            sources,
            "target",
            "targetColumn",
            "targetVariable",
            default=(
                fraud_details.get("targetColumnName")
                or snapshot_fraud_details.get("targetColumnName")
                or ""
            ),
        ),
        "championModel": _first_nonempty(
            sources,
            "championModel",
            "champion",
            default="",
        ),
        "linkedDatasets": linked_datasets,
        "createdAt": relational_project.get("created_at") or "",
        "lastSavedAt": relational_project.get("updated_at") or "",
        "storageLocation": "persistence_database",
        "recoveryStatus": "No exact legacy enrichment record found",
    }


def _enrich_exact_legacy_project(
    project: Dict[str, Any],
    legacy_project: Dict[str, Any] | None,
) -> Dict[str, Any]:
    if not legacy_project:
        return project

    enriched = {**project}
    if enriched.get("taskType") in (None, "", "unknown"):
        enriched["taskType"] = (
            legacy_project.get("taskType")
            or legacy_project.get("task")
            or "unknown"
        )
    if not enriched.get("target"):
        enriched["target"] = legacy_project.get("target") or ""
    if not enriched.get("championModel"):
        enriched["championModel"] = (
            legacy_project.get("championModel")
            or legacy_project.get("champion")
            or ""
        )
    if not enriched.get("linkedDatasets"):
        legacy_datasets = legacy_project.get("linkedDatasets")
        enriched["linkedDatasets"] = (
            legacy_datasets if isinstance(legacy_datasets, list) else []
        )
    enriched["storageLocation"] = (
        legacy_project.get("storageLocation")
        or enriched.get("storageLocation")
        or "persistence_database"
    )
    enriched["recoveryStatus"] = (
        legacy_project.get("recoveryStatus")
        or "Exact legacy enrichment record applied"
    )
    return enriched


def build_project_context(project_id: str) -> Dict[str, Any]:
    requested_project_id = str(project_id or "").strip()
    relational_project = get_persistence_service().projects.get(
        requested_project_id
    )
    if not relational_project:
        raise LookupError(
            "Model Documentation project context requires an exact "
            "relational project."
        )

    projects = list_backend_projects(include_recovery_candidates=True)
    legacy_project = next(
        (
            item
            for item in projects
            if str(item.get("projectId") or "") == requested_project_id
        ),
        None,
    )
    project = _enrich_exact_legacy_project(
        _relational_project_context(
            relational_project,
            requested_project_id,
        ),
        legacy_project,
    )

    linked = project_dataset_links(requested_project_id)
    datasets = enrich_linked_datasets(project, linked)
    analytics = list_analytics_insights(project_id=requested_project_id)
    diagnostics = [
        item
        for item in analytics
        if item.get("analysisType") == "model_diagnostics"
        or str(item.get("createdIn") or "").lower() == "atlas model diagnostics"
    ]
    normalized_project = normalize_project(project)
    if str(normalized_project.get("projectId") or "") != requested_project_id:
        raise RuntimeError(
            "Model Documentation project context identity mismatch."
        )
    return {
        "project": normalized_project,
        "datasets": datasets,
        "analyticsInsights": analytics,
        "chartInsights": [item for item in analytics if item.get("analysisType") == "chart_studio"],
        "diagnosticInsights": diagnostics,
        "recoveredShell": "recovered" in str(project.get("storageLocation") or project.get("recoveryStatus") or "").lower(),
        "contextStatus": (
            "exact_legacy_enrichment"
            if legacy_project
            else "relational_project_only"
        ),
        "generatedAt": now_iso(),
    }


def normalize_project(project: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "projectId": project.get("projectId") or "",
        "projectName": project.get("name") or project.get("projectName") or "Untitled ATLAS Project",
        "projectType": project.get("projectType") or "unknown",
        "taskType": project.get("taskType") or project.get("task") or "unknown",
        "target": project.get("target") or "",
        "championModel": project.get("championModel") or project.get("champion") or infer_champion_from_text(project),
        "createdAt": project.get("createdAt") or "",
        "lastModified": project.get("lastSavedAt") or project.get("lastModified") or project.get("updatedAt") or "",
        "recoveryStatus": project.get("recoveryStatus") or "",
        "storageLocation": project.get("storageLocation") or "",
    }


def enrich_linked_datasets(project: Dict[str, Any], linked: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    registry = load_registry().get("datasets", [])
    by_id = {item.get("datasetId"): item for item in registry}
    rows: List[Dict[str, Any]] = []
    for link in linked:
        dataset = link.get("dataset") or by_id.get(link.get("datasetId")) or {}
        rows.append(summarize_dataset(dataset, link.get("datasetRole"), link.get("versionId")))
    for dataset in project.get("linkedDatasets") or []:
        if not any(item.get("datasetId") == dataset.get("datasetId") for item in rows):
            rows.append(summarize_dataset(dataset, dataset.get("datasetRole"), dataset.get("versionId")))
    return rows


def summarize_dataset(dataset: Dict[str, Any], role: str | None = None, version_id: str | None = None) -> Dict[str, Any]:
    versions = dataset.get("versions") or []
    latest = versions[-1] if versions else dataset.get("latestVersion") or dataset
    source = dataset.get("sourceInfo") or {}
    return {
        "datasetId": dataset.get("datasetId") or "",
        "dataset": dataset.get("name") or dataset.get("dataset") or dataset.get("datasetId") or "Dataset",
        "role": role or dataset.get("datasetRole") or source.get("datasetRole") or infer_dataset_role(dataset),
        "rows": latest.get("rowCount") or dataset.get("rowCount") or "",
        "columns": latest.get("columnCount") or dataset.get("columnCount") or "",
        "target": source.get("targetColumn") or source.get("target") or "",
        "eventRateOrTargetMean": source.get("eventRate") or source.get("targetMean") or "",
        "version": version_id or latest.get("versionId") or "",
        "sourceType": dataset.get("sourceType") or "registered",
        "owner": dataset.get("ownerDisplayName") or dataset.get("ownerUserId") or "",
        "createdAt": dataset.get("createdAt") or "",
        "updatedAt": dataset.get("updatedAt") or "",
    }


def infer_dataset_role(dataset: Dict[str, Any]) -> str:
    text = f"{dataset.get('name', '')} {dataset.get('description', '')}".lower()
    if "oot" in text:
        return "OOT"
    if "test" in text:
        return "TEST"
    if "train" in text or "development" in text:
        return "TRAIN"
    if "score" in text or "prediction" in text:
        return "Scored Output"
    return "Modeling"


def infer_champion_from_text(project: Dict[str, Any]) -> str:
    text = json.dumps(project).lower()
    if "random forest" in text:
        return "Random Forest"
    if "linear regression" in text:
        return "Linear Regression"
    if "logistic regression" in text:
        return "Logistic Regression"
    return "Not selected"


def create_document(project_id: str, template_type: str = "model_development_document", created_by: str = "ATLAS local user") -> Dict[str, Any]:
    context = build_project_context(project_id)
    project = context["project"]
    document_id = f"doc-{project_id}-{int(datetime.now(timezone.utc).timestamp())}"
    family = family_for_project(str(project.get("projectType") or ""), str(project.get("taskType") or ""))
    outline = OPTIMIZATION_DECISION_OUTLINE if template_type == "optimization_decision_report" else MODEL_DEVELOPMENT_OUTLINE
    if template_type == "validator_aligned_mdd":
        outline = [(section_id, title, required) for section_id, title, required in template_sections(family)]
    sections = [
        build_section(section_id, title, index + 1, required, context)
        for index, (section_id, title, required) in enumerate(outline)
    ]
    document = {
        "documentId": document_id,
        "projectId": project.get("projectId") or project_id,
        "projectName": project.get("projectName"),
        "templateType": template_type,
        "templateFamily": family,
        "documentVersion": "0.1",
        "documentStatus": "working_draft",
        "currentValidatedVersion": None,
        "versionControlEntries": [],
        "sectionChangeRecords": [],
        "createdAt": now_iso(),
        "updatedAt": now_iso(),
        "createdBy": created_by,
        "sections": sections,
        "quality": quality_check({"sections": sections}, context),
        "exports": [],
        "contextSummary": context,
    }
    return save_document(document)


def build_section(section_id: str, title: str, order: int, required: bool, context: Dict[str, Any]) -> Dict[str, Any]:
    evidence = default_evidence_for_section(section_id, context)
    status = "draft"
    if evidence:
        status = "ready_for_review"
    return {
        "sectionId": section_id,
        "title": title,
        "order": order,
        "required": required,
        "status": status,
        "autoDraft": draft_for_section(section_id, context),
        "userNarrative": "",
        "attachedEvidence": evidence,
        "includeInExport": True,
        "missingEvidenceWarning": missing_warning(section_id, evidence),
    }


def draft_for_section(section_id: str, context: Dict[str, Any]) -> str:
    project = context["project"]
    datasets = context["datasets"]
    analytics = context.get("analyticsInsights") or []
    charts = context.get("chartInsights") or []
    diagnostics = context.get("diagnosticInsights") or []
    task = project.get("taskType") or "unknown task"
    target = project.get("target") or "the selected target"
    champion = project.get("championModel") or "the champion model"
    dataset_count = len(datasets)
    diagnostic_count = len(diagnostics)
    chart_count = len(charts)
    insight_count = len(analytics)
    if section_id.startswith("optimization_"):
        topic = section_id.removeprefix("optimization_").replace("_", " ")
        dataset_names = ", ".join(item.get("dataset") for item in datasets[:8]) or "No governed optimization inputs are linked"
        return (
            f"ATLAS optimization evidence topic: {topic}. Project: {project.get('projectName')}. "
            f"Governed input inventory: {dataset_names}. "
            "Populate this section from captured optimization evidence for the exact formulation version, solver run, decision vector, scenarios, sensitivity analysis, validation, and approval record. "
            "Do not represent the section as complete when the required evidence is missing, stale, infeasible, or business-invalid."
        )
    if section_id == "executive_summary":
        return (
            f"What ATLAS found: {project.get('projectName')} is documented as a {task} model development effort targeting {target}. "
            f"The current champion model is {champion}. ATLAS found {dataset_count} linked dataset record(s), "
            f"{diagnostic_count} model diagnostic insight(s), {chart_count} chart insight(s), and {insight_count} total analytics insight(s).\n\n"
            "Evidence support: The draft is assembled from the backend project registry, Dataset Library metadata, saved diagnostics, "
            "Chart Studio outputs, Analytics Workbench insights, and governance notes where available.\n\n"
            "Needs review: Confirm the champion model, final target definition, development/test/OOT evidence, and validation-ready performance tables.\n\n"
            "Governance caveat: Missing evidence should remain visible in the quality checklist until the MDD is ready for model risk review."
        )
    if section_id == "business_objective":
        return f"The business objective is to support the documented decision use case for {project.get('projectName')} using a governed, reproducible model development process."
    if section_id == "model_use_case":
        return f"The model is intended for {task} decision support. Intended-use language should be reviewed by business, compliance, and model risk stakeholders."
    if section_id == "target_definition":
        return f"The modeled target is {target}. The definition should state the event/outcome, observation window, performance window, exclusions, and treatment of ambiguous records."
    if section_id == "data_sources":
        names = ", ".join(item.get("dataset") for item in datasets[:5]) or "No linked datasets were found"
        return (
            f"What ATLAS found: Data sources currently linked to the project include: {names}.\n\n"
            "Evidence support: Dataset Library records provide available lineage, owner, version, row count, column count, target metadata, and dataset role evidence.\n\n"
            "Needs review: Confirm that source systems, extraction dates, exclusions, and refresh cadence are complete.\n\n"
            "Governance caveat: The final MDD should preserve dataset version identifiers so validation can reproduce the development record."
        )
    if section_id == "development_test_oot":
        return "Development, test, and OOT data should be reviewed for separation, leakage risk, temporal ordering, population consistency, and target availability."
    if section_id == "eda":
        return "EDA evidence should summarize target distribution, missingness, data quality, distribution shape, key relationships, and potential leakage or stability concerns."
    if section_id == "missing_value_treatment":
        return "Missing value treatment should identify variable-level strategies, rationale, imputation statistics, and any variables excluded due to missingness."
    if section_id == "outlier_treatment":
        return "Outlier treatment should document detection thresholds, caps or no-treatment decisions, and whether extreme values represent valid signal."
    if section_id == "feature_engineering":
        return "Feature engineering should document transformations, encodings, scaling, generated features, excluded fields, and reproducibility considerations."
    if section_id == "correlation_vif":
        return "Correlation and VIF review should document multicollinearity risks, high-VIF predictors, and the modeling decision for correlated fields."
    if section_id == "model_training_strategy":
        return "Training strategy should document split or cross-validation design, candidate algorithms, hyperparameter search, calibration, PCA settings, and optimization metric."
    if section_id == "candidate_models":
        return "Candidate models should be compared using task-appropriate metrics, runtime, interpretability, stability, and governance constraints."
    if section_id == "champion_model_selection":
        return f"The champion selection rationale should explain why {champion} was selected over challengers, including performance, interpretability, stability, and implementation constraints."
    if section_id == "performance_summary":
        return (
            f"What ATLAS found: {diagnostic_count} diagnostic insight(s) are currently available for performance support.\n\n"
            "Evidence support: Attach training results, evaluation metrics, model diagnostics fit summaries, and scored split comparison tables.\n\n"
            "Needs review: Confirm the primary metric, secondary metrics, train/test/OOT comparison, and whether performance is sufficient for intended use.\n\n"
            "Governance caveat: If performance evidence is missing or stale, keep the section in needs-review status until final model approval."
        )
    if section_id == "evaluation_diagnostics":
        return "Evaluation diagnostics should include residual diagnostics for regression or lift, calibration, ROC/PR, threshold analysis, and decile review for classification."
    if section_id == "model_interpretation":
        return "Interpretation should summarize feature importance, coefficients, SHAP-style summaries, PDP/sensitivity, RFE, and model behavior caveats."
    if section_id == "scoring_external_validation":
        return "Scoring and external validation should document TEST and OOT scoring outputs, stability metrics, prediction distribution, and any model drift or validation warnings."
    if section_id == "model_diagnostics":
        return (
            f"What ATLAS found: {diagnostic_count} saved model diagnostic insight(s) are available for this project.\n\n"
            "Evidence support: Relevant evidence can include prediction summary, regression fit summary, detailed lift table, population waterfall, VIF analysis, KPI aggregation, and Pozzolo calibration.\n\n"
            "Needs review: Confirm field mapping, actual-vs-predicted alignment, prediction distribution quality, segment performance, and any drift/stability warnings.\n\n"
            "Governance caveat: Diagnostics should be interpreted as model behavior evidence and not as causal proof."
        )
    if section_id == "fair_lending_governance":
        return "Governance considerations should address fair lending, FCRA/ECOA applicability, data lineage, model risk controls, limitations, and monitoring expectations."
    if section_id == "implementation_plan":
        return "Implementation plan should identify the deployment path, scoring cadence, controls, owners, dependencies, rollback plan, and production readiness decision."
    if section_id == "monitoring_plan":
        return "Monitoring plan should include performance metrics, PSI/VSI, input drift, prediction drift, outcome drift, residual drift for regression, calibration/lift drift for classification, review cadence, and redevelopment triggers."
    if section_id == "limitations":
        return "Known limitations should include sample constraints, data quality issues, external validation gaps, interpretability caveats, monitoring dependencies, and assumptions requiring stakeholder approval."
    return "Appendices should include generated code, detailed diagnostics, evidence inventories, governance notes, and decision summaries as appropriate."


def default_evidence_for_section(section_id: str, context: Dict[str, Any]) -> List[Dict[str, Any]]:
    project = context.get("project") or {}
    datasets = context.get("datasets") or []
    analytics = context.get("analyticsInsights") or []
    charts = context.get("chartInsights") or []
    diagnostics = context.get("diagnosticInsights") or []
    evidence: List[Dict[str, Any]] = []
    project_metadata = {
        "type": "Project Metadata",
        "title": "Project Metadata and Workflow Context",
        "source": "Project Registry",
        "preview": project,
    }
    if section_id.startswith("optimization_"):
        evidence.append(project_metadata)
        if section_id in {"optimization_data_inputs", "optimization_formulation_readiness", "optimization_reproducibility_appendix"} and datasets:
            evidence.append({"type": "Dataset Inventory", "title": "Optimization Input Inventory", "source": "Dataset Library", "rows": datasets})
        return evidence
    target_metadata = {
        "type": "Target Metadata",
        "title": "Target Definition Evidence",
        "source": "Project Registry",
        "preview": {
            "target": project.get("target") or "",
            "taskType": project.get("taskType") or "",
            "projectType": project.get("projectType") or "",
            "championModel": project.get("championModel") or "",
        },
    }
    if section_id in {
        "executive_summary",
        "business_objective",
        "model_use_case",
        "target_definition",
        "population_definition",
        "inclusion_exclusion",
        "model_training_strategy",
        "champion_model_selection",
        "implementation_plan",
        "monitoring_plan",
        "limitations",
        "appendices",
    }:
        evidence.append(project_metadata)
    if section_id == "target_definition":
        evidence.append(target_metadata)
    if section_id in {"data_sources", "development_test_oot", "target_definition", "population_definition"} and datasets:
        evidence.append({"type": "Dataset Inventory", "title": "Dataset Inventory", "source": "Dataset Library", "rows": datasets})
    if section_id in {"eda", "missing_value_treatment", "outlier_treatment", "feature_engineering", "correlation_vif", "model_training_strategy", "candidate_models", "appendices"} and analytics:
        evidence.append({"type": "Workflow Evidence", "title": "Analytics and Workflow Insights", "source": "Analytics Workbench", "rows": analytics[:12]})
    if section_id in {"evaluation_diagnostics", "model_interpretation", "scoring_external_validation", "appendices"} and charts:
        evidence.append({"type": "Chart Evidence", "title": "Chart Studio Insights", "source": "Chart Studio", "rows": charts[:12]})
    if section_id in {"model_diagnostics", "scoring_external_validation", "performance_summary", "evaluation_diagnostics", "model_interpretation", "champion_model_selection"} and diagnostics:
        evidence.append({"type": "Diagnostic Evidence", "title": "Model Diagnostics Insights", "source": "Model Diagnostics", "rows": diagnostics[:12]})
    if section_id == "fair_lending_governance":
        evidence.append({"type": "Governance Notes", "title": "Governance and Fair Lending Notes", "source": "ATLAS Guide", "rows": governance_warnings(context)})
    return dedupe_evidence(evidence)


def dedupe_evidence(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    seen = set()
    deduped = []
    for item in items:
        key = (str(item.get("source")), str(item.get("type")), str(item.get("title")))
        if key in seen:
            continue
        seen.add(key)
        deduped.append(item)
    return deduped


def governance_warnings(context: Dict[str, Any]) -> List[Dict[str, Any]]:
    project = context.get("project") or {}
    warnings = [
        {"topic": "Model Risk", "detail": "Document validation readiness, limitations, monitoring, and final model approval."},
        {"topic": "Fair Lending", "detail": "Review protected-class proxy risk and adverse-action explainability where applicable."},
    ]
    if context.get("recoveredShell"):
        warnings.append({"topic": "Recovered Project", "detail": "This project was reconstructed from artifacts; review section evidence before finalization."})
    if not project.get("target"):
        warnings.append({"topic": "Target Definition", "detail": "Target definition is missing from backend project metadata."})
    return warnings


def missing_warning(section_id: str, evidence: List[Dict[str, Any]]) -> str:
    if evidence:
        return ""
    if section_id in {"limitations", "implementation_plan", "monitoring_plan", "appendices"}:
        return "Drafted from standard ATLAS guidance; attach project-specific evidence before final governance review."
    return "No direct ATLAS evidence is attached yet. Review and attach evidence before marking complete."


def save_section(document_id: str, section_payload: Dict[str, Any]) -> Dict[str, Any]:
    document = load_document(document_id)
    sections = document.get("sections") or []
    for section in sections:
        if section.get("sectionId") == section_payload.get("sectionId"):
            prior_text = section.get("userNarrative", "")
            status = section_payload.get("status", section.get("status", "draft"))
            rationale = section_payload.get("notApplicableRationale", section.get("notApplicableRationale", ""))
            if status == "not_applicable" and not rationale.strip():
                raise ValueError("Not applicable requires a documented rationale.")
            section.update(
                {
                    "status": section_payload.get("status", section.get("status", "draft")),
                    "userNarrative": section_payload.get("userNarrative", section.get("userNarrative", "")),
                    "includeInExport": bool(section_payload.get("includeInExport", section.get("includeInExport", True))),
                    "attachedEvidence": section_payload.get("attachedEvidence", section.get("attachedEvidence", [])),
                    "notApplicableRationale": rationale,
                }
            )
            if prior_text != section.get("userNarrative", ""):
                document.setdefault("sectionChangeRecords", []).append({"section_id": section.get("sectionId"), "section_number": section.get("order"), "section_title": section.get("title"), "change_type": "modification", "prior_text_hash": str(hash(prior_text)), "new_text_hash": str(hash(section.get("userNarrative", ""))), "created_at": now_iso()})
            break
    document["quality"] = quality_check(document, document.get("contextSummary") or build_project_context(document.get("projectId", "")))
    return save_document(document)


def attach_evidence(document_id: str, section_id: str, evidence: Dict[str, Any]) -> Dict[str, Any]:
    document = load_document(document_id)
    matched = False
    for section in document.get("sections") or []:
        if section.get("sectionId") == section_id:
            matched = True
            evidence_id = str(
                evidence.get("evidenceId")
                or evidence.get("evidence_id")
                or evidence.get("sourceRef")
                or ""
            ).strip()
            attached = list(section.get("attachedEvidence") or [])
            if evidence_id:
                replaced = False
                for index, existing in enumerate(attached):
                    existing_id = str(
                        existing.get("evidenceId")
                        or existing.get("evidence_id")
                        or existing.get("sourceRef")
                        or ""
                    ).strip()
                    if existing_id == evidence_id:
                        attached[index] = evidence
                        replaced = True
                        break
                if not replaced:
                    attached.append(evidence)
            elif evidence not in attached:
                attached.append(evidence)
            section["attachedEvidence"] = attached
            if section.get("status") in {"missing_evidence", "needs_evidence", "draft"}:
                section["status"] = "ready_for_review"
            break
    if not matched:
        raise ValueError(
            f"Model documentation section '{section_id}' was not found."
        )
    document["quality"] = quality_check(document, document.get("contextSummary") or build_project_context(document.get("projectId", "")))
    return save_document(document)


def quality_check(document: Dict[str, Any], context: Dict[str, Any] | None = None) -> Dict[str, Any]:
    sections = document.get("sections") or []
    required = [section for section in sections if section.get("required")]
    complete = [section for section in required if section.get("status") in {"complete", "not_applicable"}]
    missing = [
        section for section in required
        if section.get("status") in {"missing_evidence", "needs_evidence"}
        or (not section.get("attachedEvidence") and not section.get("userNarrative") and not section.get("autoDraft"))
    ]
    needs_review = [section for section in sections if section.get("status") in {"draft", "needs_review", "ready_for_review"}]
    governance_warnings_count = len(governance_warnings(context or document.get("contextSummary") or {}))
    completion = round((len(complete) / max(1, len(required))) * 100)
    checks = [
        {"id": "required_sections", "label": "Required sections complete", "status": "pass" if not missing else "warning", "detail": f"{len(complete)} of {len(required)} required sections complete."},
        {"id": "dataset_lineage", "label": "Dataset lineage documented", "status": "pass" if has_evidence(sections, "data_sources") else "warning", "detail": "Dataset inventory evidence is attached." if has_evidence(sections, "data_sources") else "Attach Dataset Library evidence."},
        {"id": "target_definition", "label": "Target definition documented", "status": "pass" if section_has_text(sections, "target_definition") else "warning", "detail": "Target definition draft is available."},
        {"id": "performance_table", "label": "Model performance table", "status": "pass" if has_evidence(sections, "performance_summary") else "warning", "detail": "Attach evaluation or diagnostic evidence."},
        {"id": "model_interpretation", "label": "Model interpretation", "status": "pass" if section_has_text(sections, "model_interpretation") else "warning", "detail": "Interpretability narrative is drafted."},
        {"id": "monitoring_plan", "label": "Monitoring plan", "status": "pass" if section_has_text(sections, "monitoring_plan") else "warning", "detail": "Monitoring plan draft is available."},
        {"id": "governance_notes", "label": "Fair Lending / governance notes", "status": "pass" if section_has_text(sections, "fair_lending_governance") else "warning", "detail": f"{governance_warnings_count} governance note(s) detected."},
    ]
    return {
        "completionPct": completion,
        "requiredSectionsComplete": len(complete),
        "requiredSectionsTotal": len(required),
        "sectionsNeedingReview": len(needs_review),
        "missingEvidenceCount": len(missing),
        "governanceWarningsCount": governance_warnings_count,
        "exportReady": len(missing) == 0 and len(needs_review) <= 3,
        "checks": checks,
    }


def has_evidence(sections: List[Dict[str, Any]], section_id: str) -> bool:
    return any(section.get("sectionId") == section_id and section.get("attachedEvidence") for section in sections)


def section_has_text(sections: List[Dict[str, Any]], section_id: str) -> bool:
    return any(section.get("sectionId") == section_id and (section.get("userNarrative") or section.get("autoDraft")) for section in sections)


def evidence_library(project_id: str) -> List[Dict[str, Any]]:
    context = build_project_context(project_id)
    items: List[Dict[str, Any]] = []
    project = context.get("project") or {}
    items.append({"type": "Project Metadata", "title": "Project Metadata and Workflow Context", "source": "Project Registry", "createdAt": context.get("generatedAt"), "preview": project})
    items.append({
        "type": "Target Metadata",
        "title": "Target Definition Evidence",
        "source": "Project Registry",
        "createdAt": context.get("generatedAt"),
        "preview": {
            "target": project.get("target") or "",
            "taskType": project.get("taskType") or "",
            "projectType": project.get("projectType") or "",
            "championModel": project.get("championModel") or "",
        },
    })
    for dataset in context.get("datasets") or []:
        items.append({"type": "Dataset", "title": dataset.get("dataset"), "source": "Dataset Library", "createdAt": dataset.get("createdAt"), "preview": dataset})
    for insight in context.get("analyticsInsights") or []:
        source = "Chart Studio" if insight.get("analysisType") == "chart_studio" else "Analytics Workbench"
        items.append({"type": "Chart" if source == "Chart Studio" else "Narrative", "title": insight.get("title"), "source": source, "createdAt": insight.get("createdAt"), "preview": insight})
    for insight in context.get("diagnosticInsights") or []:
        items.append({"type": "Diagnostic", "title": insight.get("title"), "source": "Model Diagnostics", "createdAt": insight.get("createdAt"), "preview": insight})
    try:
        from ..notebooks.notebook_store import list_notebook_evidence

        for evidence in list_notebook_evidence(project_id):
            items.append(
                {
                    "type": "Notebook Output",
                    "title": evidence.get("title") or "Notebook Output",
                    "source": "Notebook Studio",
                    "createdAt": evidence.get("createdAt"),
                    "preview": evidence,
                }
            )
    except Exception:
        pass
    for warning in governance_warnings(context):
        items.append({"type": "Warning", "title": warning["topic"], "source": "ATLAS Guide", "createdAt": context.get("generatedAt"), "preview": warning})
    return dedupe_evidence(items)


def export_word(document: Dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("ATLAS", 0)
    doc.add_paragraph("Enterprise AI Model Development Platform")
    doc.add_heading(document.get("projectName") or "Model Development Document", level=1)
    doc.add_paragraph(f"Template: {document.get('templateType')}")
    doc.add_paragraph(f"Generated: {now_iso()}")
    doc.add_page_break()
    doc.add_heading("Table of Contents", level=1)
    doc.add_paragraph("Update fields in Microsoft Word to generate the final table of contents.")
    for section in included_sections(document):
        doc.add_heading(f"{section.get('order')}. {section.get('title')}", level=1)
        status = str(section.get("status") or "draft").replace("_", " ").title()
        doc.add_paragraph(f"Status: {status}")
        text = section.get("userNarrative") or section.get("autoDraft") or ""
        doc.add_paragraph(text)
        warning = section.get("missingEvidenceWarning")
        if warning:
            p = doc.add_paragraph()
            run = p.add_run(f"Evidence note: {warning}")
            run.bold = True
            run.font.size = Pt(9)
        for evidence in section.get("attachedEvidence") or []:
            add_word_evidence(doc, evidence)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def add_word_evidence(doc: Any, evidence: Dict[str, Any]) -> None:
    doc.add_heading(str(evidence.get("title") or "Evidence"), level=2)
    rows = evidence.get("rows") or []
    if isinstance(rows, list) and rows and isinstance(rows[0], dict):
        columns = list(rows[0].keys())[:8]
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for index, column in enumerate(columns):
            table.rows[0].cells[index].text = str(column)
        for row in rows[:20]:
            cells = table.add_row().cells
            for index, column in enumerate(columns):
                cells[index].text = cell_text(row.get(column))
    else:
        doc.add_paragraph(cell_text(evidence))


def export_powerpoint(document: Dict[str, Any]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = document.get("projectName") or "ATLAS Model Development Document"
    title_slide.placeholders[1].text = "Enterprise AI Model Development Platform"
    slide_specs = [
        ("Business Objective", "business_objective"),
        ("Data Summary", "data_sources"),
        ("Population / Target Summary", "target_definition"),
        ("Candidate Model Summary", "candidate_models"),
        ("Champion Model Performance", "performance_summary"),
        ("Key Diagnostics", "model_diagnostics"),
        ("Interpretation / Key Drivers", "model_interpretation"),
        ("Governance and Risk Notes", "fair_lending_governance"),
        ("Implementation and Monitoring Plan", "monitoring_plan"),
    ]
    by_id = {section.get("sectionId"): section for section in document.get("sections") or []}
    for title, section_id in slide_specs:
        section = by_id.get(section_id, {})
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        text = section.get("userNarrative") or section.get("autoDraft") or "No section narrative available."
        slide.placeholders[1].text = bulletize(text)
    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()


def export_evidence_excel(document: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_pairs(summary, [
        ("Product", "ATLAS"),
        ("Subtitle", "Enterprise AI Model Development Platform"),
        ("Project", document.get("projectName")),
        ("Document ID", document.get("documentId")),
        ("Template", document.get("templateType")),
        ("Completion %", (document.get("quality") or {}).get("completionPct")),
    ])
    write_table(wb.create_sheet("Document Checklist"), (document.get("quality") or {}).get("checks") or [])
    context = document.get("contextSummary") or {}
    write_table(wb.create_sheet("Dataset Inventory"), context.get("datasets") or [])
    for title, section_id in [
        ("Development Test OOT Data", "development_test_oot"),
        ("EDA Evidence", "eda"),
        ("Feature Engineering", "feature_engineering"),
        ("VIF Analysis", "correlation_vif"),
        ("Training Results", "candidate_models"),
        ("Evaluation Metrics", "performance_summary"),
        ("Model Diagnostics", "model_diagnostics"),
        ("Scored Outputs", "scoring_external_validation"),
        ("Charts and Insights", "evaluation_diagnostics"),
        ("Generated Code", "appendices"),
        ("Governance Notes", "fair_lending_governance"),
        ("Limitations", "limitations"),
    ]:
        rows = evidence_rows_for_section(document, section_id)
        write_table(wb.create_sheet(title[:31]), rows or [{"Section": section_id, "Evidence": "No attached evidence"}])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="0F172A")
            cell.fill = PatternFill("solid", fgColor="E2E8F0")
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            width = min(max(14, max(len(str(cell.value or "")) for cell in column_cells[:120]) + 2), 60)
            sheet.column_dimensions[letter].width = width
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def evidence_rows_for_section(document: Dict[str, Any], section_id: str) -> List[Dict[str, Any]]:
    section = next((item for item in document.get("sections") or [] if item.get("sectionId") == section_id), {})
    rows: List[Dict[str, Any]] = []
    for evidence in section.get("attachedEvidence") or []:
        if isinstance(evidence.get("rows"), list):
            for row in evidence.get("rows")[:250]:
                rows.append(row if isinstance(row, dict) else {"Evidence": cell_text(row)})
        else:
            rows.append({"Title": evidence.get("title"), "Source": evidence.get("source"), "Type": evidence.get("type")})
    return rows


def included_sections(document: Dict[str, Any]) -> List[Dict[str, Any]]:
    return sorted([section for section in document.get("sections") or [] if section.get("includeInExport")], key=lambda item: int(item.get("order") or 0))


def write_pairs(sheet: Any, pairs: List[tuple[Any, Any]]) -> None:
    sheet.append(["Field", "Value"])
    for key, value in pairs:
        sheet.append([cell_text(key), cell_text(value)])


def write_table(sheet: Any, rows: List[Dict[str, Any]]) -> None:
    columns = sorted({key for row in rows for key in row.keys()}) or ["Message"]
    sheet.append(columns)
    if not rows:
        sheet.append(["No rows available."])
        return
    for row in rows:
        sheet.append([cell_text(row.get(column)) for column in columns])


def bulletize(text: str) -> str:
    sentences = re.split(r"(?<=[.!?])\s+", text.strip())
    return "\n".join(f"- {sentence.strip()}" for sentence in sentences[:5] if sentence.strip()) or "- No summary available."


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return str(value)
    return json.dumps(value, ensure_ascii=True)[:5000]


def safe_filename(value: str, extension: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", value or "atlas-model-document").strip("-._")
    return f"{(safe or 'atlas-model-document')[:80]}.{extension}"
