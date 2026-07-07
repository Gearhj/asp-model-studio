from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, Iterable, List


class ExportDependencyError(RuntimeError):
    pass


def build_excel_export(payload: Dict[str, Any]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:  # pragma: no cover - environment dependent
        raise ExportDependencyError("Excel export requires openpyxl. Install backend requirements.") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)

    header_fill = PatternFill("solid", fgColor="E2E8F0")
    title_font = Font(bold=True, size=13, color="0F172A")
    header_font = Font(bold=True, color="0F172A")

    def add_sheet(name: str, rows: List[List[Any]]) -> None:
        sheet = workbook.create_sheet(safe_sheet_name(name))
        for row in rows:
            sheet.append([normalize_cell(value) for value in row])
        if sheet.max_row:
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        autosize_sheet(sheet)

    add_sheet("Summary", key_value_rows(summary_fields(payload)))
    add_sheet("Analysis Plan", [["Step", "Recommendation"]] + [[index + 1, item] for index, item in enumerate(payload.get("analysisPlan") or [])])
    add_table_sheet(add_sheet, "Pivot Results", payload.get("pivotRows") or [])
    add_table_sheet(add_sheet, "SQL Results", result_rows(payload.get("sqlResult")))
    add_sas_sheet(add_sheet, payload.get("sasResult"))
    add_output_sheet(add_sheet, "Python Output", payload.get("pythonResult"))
    add_output_sheet(add_sheet, "R Output", payload.get("rResult"))
    add_sheet(
        "Generated Code",
        [
            ["Language", "Generated Code", "User Edited Code"],
            ["Python", payload.get("generatedPython", ""), payload.get("userEditedPython", "")],
            ["SQL", payload.get("generatedSql", ""), payload.get("userEditedSql", "")],
            ["R", payload.get("generatedR", ""), payload.get("userEditedR", "")],
            ["SAS-compatible", payload.get("generatedSas", ""), payload.get("userEditedSas", "")],
        ],
    )
    add_sheet("Dataset Metadata", key_value_rows(dataset_metadata(payload)))
    add_sheet("Governance Notes", key_value_rows(governance_notes(payload)))

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_result_excel_export(payload: Dict[str, Any]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:  # pragma: no cover - environment dependent
        raise ExportDependencyError("Excel result export requires openpyxl. Install backend requirements.") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    result_type = str(payload.get("resultType") or "analysis_result")
    result_payload = payload.get("resultPayload") or {}

    def add_sheet(name: str, rows: List[List[Any]]) -> None:
        sheet = workbook.create_sheet(safe_sheet_name(name))
        for row in rows:
            sheet.append([normalize_cell(value) for value in row])
        if sheet.max_row:
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="E2E8F0")
                cell.font = Font(bold=True, color="0F172A")
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        autosize_sheet(sheet)

    add_sheet("Summary", key_value_rows(result_summary_fields(payload)))
    add_result_sheets(add_sheet, result_type, result_payload)
    add_sheet("Metadata", key_value_rows(result_metadata(payload)))

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_word_export(payload: Dict[str, Any]) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as exc:  # pragma: no cover - environment dependent
        raise ExportDependencyError("Word export requires python-docx. Install backend requirements.") from exc

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    document.add_heading(payload.get("exportTitle") or "ATLAS Analytics Workbench Report", 0)
    document.add_paragraph(f"Created: {payload.get('createdAt') or datetime.utcnow().isoformat()}Z")
    document.add_paragraph(f"Dataset: {dataset_name(payload)}")

    add_doc_section(document, "Business Question", payload.get("businessQuestion") or "No business question provided.")
    add_doc_section(document, "Executive Summary", executive_summary(payload))
    add_doc_table(document, "Dataset Context", summary_fields(payload))
    add_doc_list(document, "Analysis Plan", payload.get("analysisPlan") or ["No analysis plan generated."])
    add_doc_section(document, "Key Findings", key_findings(payload))
    add_doc_rows(document, "Pivot Results", payload.get("pivotRows") or [])
    add_doc_rows(document, "SQL Results", result_rows(payload.get("sqlResult")))
    add_doc_sas(document, payload.get("sasResult"))
    add_doc_section(document, "Python Output Summary", output_summary(payload.get("pythonResult")))
    add_doc_section(document, "R Output Summary", output_summary(payload.get("rResult")))
    add_doc_code(document, "Generated Code Appendix", payload)
    add_doc_table(document, "Governance and Limitations", governance_notes(payload))

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def build_powerpoint_export(payload: Dict[str, Any]) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception as exc:  # pragma: no cover - environment dependent
        raise ExportDependencyError("PowerPoint export requires python-pptx. Install backend requirements.") from exc

    presentation = Presentation()

    def add_bullet_slide(title: str, bullets: Iterable[str]) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.font.size = Pt(18 if index < 5 else 14)

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = payload.get("exportTitle") or "ATLAS Analytics Workbench Report"
    title_slide.placeholders[1].text = f"{dataset_name(payload)}\n{payload.get('createdAt') or datetime.utcnow().isoformat()}Z"

    add_bullet_slide("Business Question", [payload.get("businessQuestion") or "No business question provided.", payload.get("cleanedQuestion") or "No cleaned question provided."])
    add_bullet_slide("Dataset Context", [f"{key}: {value}" for key, value in summary_fields(payload)])
    add_bullet_slide("Executive Summary", executive_summary(payload).split("\n"))
    add_bullet_slide("Analysis Plan", payload.get("analysisPlan") or ["No analysis plan generated."])
    add_bullet_slide("Pivot Summary", summarize_rows(payload.get("pivotRows") or [], "Pivot"))
    add_bullet_slide("SAS-Compatible Summary", sas_summary(payload.get("sasResult")))
    add_bullet_slide("SQL Summary", result_summary(payload.get("sqlResult"), "SQL"))
    add_bullet_slide("Code Appendix Overview", code_overview(payload))
    add_bullet_slide("Governance and Limitations", [f"{key}: {value}" for key, value in governance_notes(payload)])

    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def key_value_rows(items: List[tuple[str, Any]]) -> List[List[Any]]:
    return [["Field", "Value"]] + [[key, value] for key, value in items]


def result_summary_fields(payload: Dict[str, Any]) -> List[tuple[str, Any]]:
    return [
        ("Export title", payload.get("exportTitle") or "ATLAS Analysis Result"),
        ("Result type", payload.get("resultType") or "-"),
        ("Dataset", payload.get("datasetName") or "-"),
        ("Version", payload.get("versionId") or "-"),
        ("Business question", payload.get("businessQuestion") or "-"),
        ("Created at", payload.get("createdAt") or datetime.utcnow().isoformat() + "Z"),
    ]


def result_metadata(payload: Dict[str, Any]) -> List[tuple[str, Any]]:
    metadata = payload.get("metadata") or {}
    rows = result_summary_fields(payload)
    rows.extend((str(key), value) for key, value in metadata.items())
    return rows


def add_result_sheets(add_sheet, result_type: str, result_payload: Dict[str, Any]) -> None:
    if result_type == "column_explorer":
        add_sheet("Result Data", key_value_rows(result_payload.get("statistics") or []))
        add_table_sheet(add_sheet, "Chart Data", result_payload.get("histogramBins") or [])
        add_sheet("Warnings", [["Warning"], *[[item] for item in result_payload.get("warnings", [])]])
        return
    if result_type == "target_relationship":
        add_table_sheet(add_sheet, "Result Data", result_payload.get("ranking") or [])
        add_sheet("Target Context", key_value_rows(result_payload.get("summary") or []))
        return
    if result_type == "group_comparison":
        add_table_sheet(add_sheet, "Result Data", result_payload.get("rows") or [])
        add_sheet("Configuration", key_value_rows(result_payload.get("configuration") or []))
        return
    if result_type == "data_quality":
        add_table_sheet(add_sheet, "Result Data", result_payload.get("qualityChecks") or [])
        add_table_sheet(add_sheet, "Chart Data", result_payload.get("correlations") or [])
        add_sheet("Warnings", [["Warning"], *[[item] for item in result_payload.get("warnings", [])]])
        return
    if result_type == "pivot":
        add_sheet("Configuration", key_value_rows(result_payload.get("configuration") or []))
        add_table_sheet(add_sheet, "Result Data", result_payload.get("rows") or [])
        add_sheet("Chart Data", [["Readout"], [result_payload.get("readout") or ""]])
        return
    if result_type == "sql_result":
        add_table_sheet(add_sheet, "Result Data", result_payload.get("rows") or [])
        add_sheet("Execution Log", key_value_rows(result_payload.get("execution") or []))
        return
    if result_type == "sas_compatible_result":
        tables = result_payload.get("tables") or []
        if tables:
            for index, table in enumerate(tables):
                add_table_sheet(add_sheet, table.get("title") or table.get("name") or f"SAS Table {index + 1}", table.get("rows") or [])
        else:
            add_sheet("Result Data", [["Status"], ["No SAS-compatible result tables available."]])
        add_sheet("STDOUT", [["STDOUT"], [result_payload.get("stdout") or ""]])
        add_sheet("Translated Python", [["Translated Python"], [result_payload.get("translatedPython") or ""]])
        add_sheet("Warnings", [["Unsupported Statement"], *[[item] for item in result_payload.get("unsupportedStatements", [])]])
        return
    if result_type in {"python_result", "r_result"}:
        add_sheet("Result Data", key_value_rows(result_payload.get("execution") or []))
        add_table_sheet(add_sheet, "Tables", result_payload.get("tables") or [])
        add_sheet("STDOUT", [["STDOUT"], [result_payload.get("stdout") or ""]])
        add_sheet("STDERR", [["STDERR / Errors"], [result_payload.get("stderr") or result_payload.get("error") or ""]])
        return
    add_sheet("Result Data", key_value_rows([("Payload", result_payload)]))


def summary_fields(payload: Dict[str, Any]) -> List[tuple[str, Any]]:
    dataset = payload.get("dataset") or {}
    return [
        ("Business question", payload.get("businessQuestion") or "-"),
        ("Cleaned question", payload.get("cleanedQuestion") or "-"),
        ("Dataset", dataset.get("name") or payload.get("datasetName") or "-"),
        ("Dataset ID", dataset.get("datasetId") or payload.get("datasetId") or "-"),
        ("Version", dataset.get("versionId") or payload.get("versionId") or "-"),
        ("Target", payload.get("targetVariable") or "-"),
        ("Preview rows", payload.get("previewRowCount") or 0),
        ("Linked project", payload.get("linkedProjectId") or payload.get("projectId") or "-"),
        ("Created by", payload.get("createdBy") or "ATLAS local user"),
        ("Created at", payload.get("createdAt") or datetime.utcnow().isoformat() + "Z"),
    ]


def dataset_metadata(payload: Dict[str, Any]) -> List[tuple[str, Any]]:
    metadata = payload.get("metadata") or {}
    pivot_config = payload.get("pivotConfig") or {}
    rows = summary_fields(payload)
    rows.extend((f"Metadata: {key}", value) for key, value in metadata.items())
    rows.extend((f"Pivot: {key}", value) for key, value in pivot_config.items())
    return rows


def governance_notes(payload: Dict[str, Any]) -> List[tuple[str, Any]]:
    return [
        ("Governance note", "Analytics Workbench exports are generated from preview results and saved ATLAS findings unless connected to a warehouse execution backend."),
        ("Limitations", payload.get("limitations") or "Preview rows may not represent the full governed dataset. Validate conclusions before production or governance use."),
        ("SAS mode", "SAS-compatible output uses ATLAS translation for supported PROC statements and is not a licensed SAS runtime."),
        ("Auditability", "Generated code and user-edited code are included so another analyst can reproduce the analysis path."),
    ]


def add_table_sheet(add_sheet, name: str, rows: List[Dict[str, Any]]) -> None:
    columns = collect_columns(rows)
    add_sheet(name, [columns] + [[row.get(column, "") for column in columns] for row in rows] if columns else [["Status"], ["No rows available."]])


def add_sas_sheet(add_sheet, sas_result: Dict[str, Any] | None) -> None:
    rows: List[List[Any]] = [["Section", "Field", "Value"]]
    for table in (sas_result or {}).get("tables") or []:
        title = table.get("title") or table.get("name") or "SAS-compatible table"
        table_rows = table.get("rows") or []
        columns = table.get("columns") or collect_columns(table_rows)
        rows.append([title, "Columns", ", ".join(columns)])
        for item in table_rows[:200]:
            rows.append([title, "Row", json.dumps(item, default=str)])
    if len(rows) == 1:
        rows.append(["SAS-compatible Results", "Status", "No SAS-compatible tables available."])
    add_sheet("SAS-Compatible Results", rows)


def add_output_sheet(add_sheet, name: str, result: Dict[str, Any] | None) -> None:
    result = result or {}
    add_sheet(
        name,
        [
            ["Field", "Value"],
            ["Status", result.get("status") or ("success" if result.get("ok") else "not run")],
            ["Runtime ms", result.get("runtimeMs") or result.get("runtime_ms") or 0],
            ["Stdout", result.get("stdout") or ""],
            ["Stderr", result.get("stderr") or result.get("error") or ""],
            ["Return code", result.get("returnCode") if result.get("returnCode") is not None else result.get("return_code", "")],
        ],
    )


def result_rows(result: Dict[str, Any] | None) -> List[Dict[str, Any]]:
    if not result:
        return []
    rows = result.get("rows")
    return rows if isinstance(rows, list) else []


def collect_columns(rows: List[Dict[str, Any]]) -> List[str]:
    columns: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in columns:
                columns.append(key)
    return columns


def add_doc_section(document, title: str, text: str) -> None:
    document.add_heading(title, level=1)
    document.add_paragraph(text or "-")


def add_doc_list(document, title: str, items: Iterable[Any]) -> None:
    document.add_heading(title, level=1)
    for item in items:
        document.add_paragraph(str(item), style="List Bullet")


def add_doc_table(document, title: str, items: List[tuple[str, Any]]) -> None:
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"
    for key, value in items:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = normalize_cell(value)


def add_doc_rows(document, title: str, rows: List[Dict[str, Any]]) -> None:
    document.add_heading(title, level=1)
    if not rows:
        document.add_paragraph("No rows available.")
        return
    columns = collect_columns(rows)[:8]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, column in enumerate(columns):
        table.rows[0].cells[index].text = column
    for row in rows[:20]:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            cells[index].text = normalize_cell(row.get(column, ""))


def add_doc_sas(document, sas_result: Dict[str, Any] | None) -> None:
    document.add_heading("SAS-Compatible Results", level=1)
    tables = (sas_result or {}).get("tables") or []
    if not tables:
        document.add_paragraph("No SAS-compatible tables available.")
        return
    for table in tables[:3]:
        document.add_paragraph(str(table.get("title") or table.get("name") or "SAS-compatible table"), style="List Bullet")
        add_doc_rows(document, "Table Preview", table.get("rows") or [])


def add_doc_code(document, title: str, payload: Dict[str, Any]) -> None:
    document.add_heading(title, level=1)
    for label, key in [
        ("Python", "generatedPython"),
        ("SQL", "generatedSql"),
        ("R", "generatedR"),
        ("SAS-compatible", "generatedSas"),
    ]:
        document.add_heading(label, level=2)
        document.add_paragraph(payload.get(key) or "No generated code available.")


def executive_summary(payload: Dict[str, Any]) -> str:
    plan_count = len(payload.get("analysisPlan") or [])
    pivot_count = len(payload.get("pivotRows") or [])
    sql_count = len(result_rows(payload.get("sqlResult")))
    return (
        f"ATLAS generated an analytics package for {dataset_name(payload)}. "
        f"The package includes {plan_count} analysis-plan step(s), {pivot_count} pivot row(s), "
        f"and {sql_count} SQL preview row(s). Generated Python, SQL, R, and SAS-compatible code are included for reproducibility."
    )


def key_findings(payload: Dict[str, Any]) -> str:
    insight = payload.get("savedInsight") or {}
    if insight.get("summary"):
        return str(insight["summary"])
    pivot_readout = payload.get("pivotReadout")
    if pivot_readout:
        return str(pivot_readout)
    return "No saved finding summary was provided for this export."


def output_summary(result: Dict[str, Any] | None) -> str:
    if not result:
        return "No output available."
    status = result.get("status") or ("success" if result.get("ok") else "not run")
    stdout = (result.get("stdout") or "").strip()
    stderr = (result.get("stderr") or result.get("error") or "").strip()
    return f"Status: {status}\nStdout: {stdout[:1200] or '-'}\nStderr/errors: {stderr[:1200] or '-'}"


def summarize_rows(rows: List[Dict[str, Any]], label: str) -> List[str]:
    if not rows:
        return [f"No {label.lower()} rows available."]
    return [f"{label} rows available: {len(rows)}", f"Columns: {', '.join(collect_columns(rows)[:8])}"]


def sas_summary(result: Dict[str, Any] | None) -> List[str]:
    tables = (result or {}).get("tables") or []
    unsupported = (result or {}).get("unsupportedStatements") or []
    return [f"SAS-compatible tables: {len(tables)}", f"Unsupported statements: {len(unsupported)}", "No SAS license is required for ATLAS SAS-compatible mode."]


def result_summary(result: Dict[str, Any] | None, label: str) -> List[str]:
    result = result or {}
    return [
        f"{label} status: {result.get('status') or ('success' if result.get('ok') else 'not run')}",
        f"Rows: {result.get('rowCount') or len(result.get('rows') or [])}",
        f"Runtime: {result.get('runtimeMs') or 0} ms",
    ]


def code_overview(payload: Dict[str, Any]) -> List[str]:
    return [
        f"Python code lines: {line_count(payload.get('generatedPython'))}",
        f"SQL code lines: {line_count(payload.get('generatedSql'))}",
        f"R code lines: {line_count(payload.get('generatedR'))}",
        f"SAS-compatible code lines: {line_count(payload.get('generatedSas'))}",
    ]


def line_count(value: Any) -> int:
    return len(str(value or "").splitlines())


def dataset_name(payload: Dict[str, Any]) -> str:
    dataset = payload.get("dataset") or {}
    return str(dataset.get("name") or payload.get("datasetName") or "selected dataset")


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return str(value)
    return json.dumps(value, default=str)


def safe_sheet_name(name: str) -> str:
    return "".join(ch for ch in name if ch not in r"[]:*?/\\")[:31] or "Sheet"


def autosize_sheet(sheet) -> None:
    for column_cells in sheet.columns:
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
        sheet.column_dimensions[column_cells[0].column_letter].width = width
