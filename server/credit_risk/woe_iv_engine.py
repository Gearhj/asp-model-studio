from __future__ import annotations

import hashlib
import json
import math
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from fastapi import HTTPException

from ..datasets.dataset_registry import latest_version_id, require_dataset, version_dir
from ..evidence.evidence_store import upsert_evidence
from .leakage_guard import default_predictors, validate_predictors


WOE_IV_ROOT = Path(__file__).resolve().parents[1] / "artifacts" / "credit_risk" / "woe_iv"
ANALYSES_ROOT = WOE_IV_ROOT / "analyses"
REGISTRY_PATH = WOE_IV_ROOT / "woe_iv_registry.json"


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def ensure_store() -> None:
    ANALYSES_ROOT.mkdir(parents=True, exist_ok=True)
    if not REGISTRY_PATH.exists():
        write_json(REGISTRY_PATH, {"analyses": []})


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")


def read_json(path: Path, default: Any) -> Any:
    try:
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default
    return default


def analysis_path(analysis_id: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", analysis_id)
    return ANALYSES_ROOT / f"{safe}.json"


def compact_analysis(analysis: Dict[str, Any]) -> Dict[str, Any]:
    top = sorted(analysis.get("variableSummaries") or [], key=lambda row: float(row.get("iv") or 0), reverse=True)
    return {
        "analysisId": analysis.get("analysisId"),
        "projectId": analysis.get("projectId"),
        "datasetId": analysis.get("datasetId"),
        "datasetVersionId": analysis.get("datasetVersionId"),
        "targetColumn": analysis.get("targetColumn"),
        "eventValue": analysis.get("eventValue"),
        "eventRate": analysis.get("eventRate"),
        "variableCount": len(analysis.get("variableSummaries") or []),
        "variablesNeedingReview": len([row for row in analysis.get("variableSummaries") or [] if row.get("status") == "needs_review"]),
        "topIvVariable": top[0].get("variable") if top else None,
        "topIv": top[0].get("iv") if top else None,
        "status": "Needs Review" if any(row.get("status") == "needs_review" for row in analysis.get("variableSummaries") or []) else "Complete",
        "generatedAt": analysis.get("generatedAt"),
        "updatedAt": analysis.get("updatedAt") or analysis.get("generatedAt"),
    }


def load_registry() -> List[Dict[str, Any]]:
    ensure_store()
    payload = read_json(REGISTRY_PATH, {"analyses": []})
    return payload.get("analyses", []) if isinstance(payload, dict) else []


def save_registry(rows: List[Dict[str, Any]]) -> None:
    write_json(REGISTRY_PATH, {"analyses": sorted(rows, key=lambda row: str(row.get("updatedAt") or ""), reverse=True)})


def save_analysis(analysis: Dict[str, Any]) -> Dict[str, Any]:
    ensure_store()
    timestamp = now_iso()
    analysis["updatedAt"] = timestamp
    analysis["createdAt"] = analysis.get("createdAt") or analysis.get("generatedAt") or timestamp
    write_json(analysis_path(str(analysis["analysisId"])), analysis)
    registry = [row for row in load_registry() if row.get("analysisId") != analysis.get("analysisId")]
    registry.append(compact_analysis(analysis))
    save_registry(registry)
    return analysis


def get_analysis(analysis_id: str) -> Optional[Dict[str, Any]]:
    found = read_json(analysis_path(analysis_id), None)
    return found if isinstance(found, dict) else None


def get_project_analysis(project_id: str) -> Optional[Dict[str, Any]]:
    match = next((row for row in load_registry() if row.get("projectId") == project_id), None)
    return get_analysis(str(match.get("analysisId"))) if match else None


def load_dataset_frame(dataset_id: str, version_id: Optional[str] = None) -> tuple[Dict[str, Any], str, pd.DataFrame]:
    dataset = require_dataset(dataset_id)
    resolved_version = version_id or latest_version_id(dataset)
    if not resolved_version:
        raise HTTPException(status_code=400, detail="Dataset has no version available for WOE / IV analysis.")
    full_path = version_dir(dataset_id, resolved_version) / "data_full.csv"
    if not full_path.exists():
        raise HTTPException(status_code=409, detail="Full dataset artifact is required for governed WOE / IV analysis. Re-register the dataset so ATLAS can materialize data_full.csv; sample artifacts are not valid for reporting.")
    df = pd.read_csv(full_path)
    if df.empty:
        raise HTTPException(status_code=400, detail="Dataset full artifact has no rows.")
    return dataset, resolved_version, df


def run_woe_iv(payload: Dict[str, Any]) -> Dict[str, Any]:
    dataset, resolved_version, df = load_dataset_frame(str(payload.get("datasetId")), payload.get("datasetVersionId"))
    target = str(payload.get("targetColumn") or "")
    if target not in df.columns:
        raise HTTPException(status_code=400, detail=f"Target column {target!r} was not found in the dataset.")
    event_value = payload.get("eventValue")
    requested_predictors = [str(col) for col in payload.get("predictorColumns") or []]
    predictors = validate_predictors(
        requested_predictors,
        target=target,
        dataset_meta=dataset,
        columns=df.columns,
        explicit=bool(requested_predictors),
    )
    if not predictors:
        predictors = default_predictors(target=target, dataset_meta=dataset, columns=df.columns, limit=10)
    max_bins = max(2, min(int(payload.get("maxBins") or 5), 20))
    min_bin_pct = max(0.0, min(float(payload.get("minBinPct") or 0.05), 0.5))
    smoothing = max(0.0, float(payload.get("smoothing") if payload.get("smoothing") is not None else 0.5))
    method = str(payload.get("binningMethod") or "quantile")
    warnings: List[str] = []
    y = event_mask(df[target], event_value)
    event_count = int(y.sum())
    row_count = int(len(df))
    non_event_count = int(row_count - event_count)
    if event_count == 0 or non_event_count == 0:
        warnings.append("Target has one class only.")
    summaries = []
    for variable in predictors:
        try:
            summaries.append(analyze_variable(df, y, variable, method, max_bins, min_bin_pct, smoothing, bool(payload.get("excludeMissingFromBins")), warnings))
        except Exception as exc:
            summaries.append(
                {
                    "variable": variable,
                    "variableType": "unknown",
                    "iv": 0.0,
                    "predictiveStrength": "not_useful",
                    "monotonicity": "not_applicable",
                    "missingRate": 0.0,
                    "distinctCount": 0,
                    "binCount": 0,
                    "status": "error",
                    "recommendation": f"Analysis failed: {exc}",
                    "bins": [],
                }
            )
    analysis_id = analysis_id_for(payload, resolved_version)
    analysis = {
        "analysisId": analysis_id,
        "projectId": payload.get("projectId"),
        "datasetId": dataset.get("datasetId"),
        "datasetName": dataset.get("name"),
        "datasetVersionId": resolved_version,
        "targetColumn": target,
        "eventValue": event_value,
        "rowCount": row_count,
        "eventCount": event_count,
        "nonEventCount": non_event_count,
        "eventRate": round(event_count / max(row_count, 1), 6),
        "settings": {
            "binningMethod": method,
            "maxBins": max_bins,
            "minBinPct": min_bin_pct,
            "smoothing": smoothing,
            "excludeMissingFromBins": bool(payload.get("excludeMissingFromBins")),
        },
        "variableSummaries": sorted(summaries, key=lambda row: float(row.get("iv") or 0), reverse=True),
        "generatedCode": generate_code(
            {
                **payload,
                "datasetVersionId": resolved_version,
                "predictorColumns": predictors,
                "binningMethod": method,
                "maxBins": max_bins,
                "minBinPct": min_bin_pct,
                "smoothing": smoothing,
            }
        ),
        "generatedAt": now_iso(),
        "warnings": dedupe(warnings),
    }
    return save_analysis(analysis)


def analysis_id_for(payload: Dict[str, Any], version_id: str) -> str:
    basis = json.dumps(
        {
            "projectId": payload.get("projectId"),
            "datasetId": payload.get("datasetId"),
            "version": version_id,
            "target": payload.get("targetColumn"),
            "event": payload.get("eventValue"),
            "predictors": payload.get("predictorColumns"),
            "method": payload.get("binningMethod"),
        },
        sort_keys=True,
        default=str,
    )
    return "woeiv-" + hashlib.sha256(basis.encode("utf-8")).hexdigest()[:14]


def event_mask(series: pd.Series, event_value: Any) -> pd.Series:
    event_text = str(event_value)
    return series.astype(str).str.strip().eq(event_text)


def analyze_variable(df: pd.DataFrame, y: pd.Series, variable: str, method: str, max_bins: int, min_bin_pct: float, smoothing: float, exclude_missing: bool, global_warnings: List[str]) -> Dict[str, Any]:
    series = df[variable]
    missing_rate = float(series.isna().mean())
    distinct_count = int(series.nunique(dropna=True))
    numeric = pd.to_numeric(series, errors="coerce")
    variable_type = "numeric" if numeric.notna().mean() >= 0.8 and distinct_count > 2 else "categorical"
    warnings: List[str] = []
    if missing_rate > 0.4:
        warnings.append("Variable has too many missing values.")
    if distinct_count < 2:
        warnings.append("Variable has too few distinct values.")
    if method == "tree":
        warnings.append("Tree binning unavailable in this environment. Quantile binning used as fallback.")
        global_warnings.append("Tree binning unavailable in this environment. Quantile binning used as fallback.")
        method = "quantile"
    if variable_type == "numeric":
        assignments, definitions, method_warnings = numeric_bins(numeric, y, method, max_bins, min_bin_pct)
    else:
        assignments, definitions, method_warnings = categorical_bins(series, max_bins)
    warnings.extend(method_warnings)
    if series.isna().any() and not exclude_missing:
        assignments = assignments.astype(object)
        assignments[series.isna()] = "missing"
        definitions.append({"binId": "missing", "label": "Missing", "isMissingBin": True})
    bins = compute_bin_stats(assignments, definitions, y, smoothing, warnings)
    iv = round(sum(float(item.get("ivContribution") or 0) for item in bins), 6)
    monotonicity = monotonicity_status([item for item in bins if not item.get("isMissingBin")])
    status = "needs_review" if warnings or predictive_strength(iv) == "suspicious" or monotonicity == "not_monotonic" else "ready"
    return {
        "variable": variable,
        "variableType": variable_type,
        "iv": iv,
        "predictiveStrength": predictive_strength(iv),
        "monotonicity": monotonicity,
        "missingRate": round(missing_rate, 6),
        "distinctCount": distinct_count,
        "binCount": len(bins),
        "status": status,
        "recommendation": recommendation(iv, monotonicity, warnings),
        "warnings": dedupe(warnings),
        "bins": bins,
    }


def numeric_bins(series: pd.Series, y: pd.Series, method: str, max_bins: int, min_bin_pct: float) -> tuple[pd.Series, List[Dict[str, Any]], List[str]]:
    valid = series.dropna()
    warnings: List[str] = []
    if valid.nunique() < 2:
        return pd.Series(["all"] * len(series), index=series.index), [{"binId": "all", "label": "All values"}], ["Variable has too few distinct values."]
    try:
        if method in {"quantile", "monotonic"}:
            labels = pd.qcut(valid.rank(method="first"), q=min(max_bins, valid.nunique()), duplicates="drop")
            cuts = sorted(set(float(valid.loc[labels.index][labels == interval].min()) for interval in labels.cat.categories))
            edges = sorted(set([float(valid.min())] + [float(valid.quantile(i / min(max_bins, valid.nunique()))) for i in range(1, min(max_bins, valid.nunique()))] + [float(valid.max())]))
        elif method == "equal_width":
            edges = list(pd.interval_range(float(valid.min()), float(valid.max()), periods=max_bins).left) + [float(valid.max())]
        else:
            edges = [float(valid.min()), float(valid.max())]
    except Exception:
        edges = [float(valid.min()), float(valid.max())]
        warnings.append("Auto-binning produced duplicate cut points. A single interval was used.")
    edges = sorted(set(round(float(edge), 10) for edge in edges if math.isfinite(float(edge))))
    if len(edges) < 2:
        edges = [float(valid.min()), float(valid.max())]
    assignments = pd.Series([None] * len(series), index=series.index, dtype=object)
    definitions: List[Dict[str, Any]] = []
    for index in range(len(edges) - 1):
        lower = edges[index]
        upper = edges[index + 1]
        bin_id = f"bin_{index + 1}"
        if index == len(edges) - 2:
            mask = series.ge(lower) & series.le(upper)
        else:
            mask = series.ge(lower) & series.lt(upper)
        assignments[mask] = bin_id
        definitions.append({"binId": bin_id, "label": f"{lower:g} to {upper:g}", "lowerBound": lower, "upperBound": upper})
    assignments = assignments.fillna("missing")
    definitions = [definition for definition in definitions if (assignments == definition["binId"]).sum() > 0]
    if method == "monotonic":
        assignments, definitions, merge_warnings = merge_to_monotonic(assignments, definitions, y, min_bin_pct)
        warnings.extend(merge_warnings)
    return assignments, definitions, warnings


def categorical_bins(series: pd.Series, max_bins: int) -> tuple[pd.Series, List[Dict[str, Any]], List[str]]:
    clean = series.fillna("__MISSING__").astype(str)
    counts = clean[clean != "__MISSING__"].value_counts()
    top = list(counts.head(max_bins - 1).index)
    assignments = clean.where(clean.isin(top), "Other")
    definitions = [{"binId": safe_bin_id(value), "label": str(value), "categoryValues": [str(value)]} for value in top]
    if (assignments == "Other").any():
        definitions.append({"binId": "other", "label": "Other", "categoryValues": [value for value in clean.unique() if value not in top and value != "__MISSING__"]})
    assignments = assignments.map(lambda value: "missing" if value == "__MISSING__" else safe_bin_id(value) if value in top else "other")
    return assignments, definitions, []


def safe_bin_id(value: Any) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", str(value)).strip("_").lower()[:40] or "bin"


def merge_to_monotonic(assignments: pd.Series, definitions: List[Dict[str, Any]], y: pd.Series, min_bin_pct: float) -> tuple[pd.Series, List[Dict[str, Any]], List[str]]:
    warnings: List[str] = []
    current = list(definitions)
    while len(current) > 2:
        rates = []
        for definition in current:
            mask = assignments == definition["binId"]
            rates.append(float(y[mask].mean()) if mask.any() else 0.0)
        if is_monotonic(rates):
            return assignments, current, warnings
        diffs = [abs(rates[i + 1] - rates[i]) for i in range(len(rates) - 1)]
        idx = diffs.index(min(diffs))
        left = current[idx]
        right = current[idx + 1]
        merged_id = f"{left['binId']}_{right['binId']}"
        assignments = assignments.replace({left["binId"]: merged_id, right["binId"]: merged_id})
        current[idx : idx + 2] = [
            {
                "binId": merged_id,
                "label": f"{left.get('label')} + {right.get('label')}",
                "lowerBound": left.get("lowerBound"),
                "upperBound": right.get("upperBound"),
            }
        ]
        if min((assignments == item["binId"]).mean() for item in current) < min_bin_pct:
            warnings.append("Monotonic binning could not produce a stable bin structure. Review manual bins.")
            return assignments, current, warnings
    if not is_monotonic([float(y[assignments == item["binId"]].mean()) for item in current]):
        warnings.append("Monotonic binning could not produce a stable bin structure. Review manual bins.")
    return assignments, current, warnings


def is_monotonic(values: List[float]) -> bool:
    return all(values[i] <= values[i + 1] for i in range(len(values) - 1)) or all(values[i] >= values[i + 1] for i in range(len(values) - 1))


def compute_bin_stats(assignments: pd.Series, definitions: List[Dict[str, Any]], y: pd.Series, smoothing: float, warnings: List[str]) -> List[Dict[str, Any]]:
    total_bad = int(y.sum())
    total_good = int(len(y) - total_bad)
    bin_count = max(len(definitions), 1)
    bins = []
    for order, definition in enumerate(definitions):
        mask = assignments == definition["binId"]
        total = int(mask.sum())
        bad = int(y[mask].sum())
        good = int(total - bad)
        smoothing_applied = False
        if good == 0:
            warnings.append("Zero non-event count in bin, smoothing applied.")
            smoothing_applied = True
        if bad == 0:
            warnings.append("Zero event count in bin, smoothing applied.")
            smoothing_applied = True
        dist_good = (good + smoothing) / max(total_good + smoothing * bin_count, 1e-12)
        dist_bad = (bad + smoothing) / max(total_bad + smoothing * bin_count, 1e-12)
        woe = math.log(max(dist_good, 1e-12) / max(dist_bad, 1e-12))
        iv_contribution = (dist_good - dist_bad) * woe
        bins.append(
            {
                **definition,
                "binId": definition.get("binId") or f"bin_{order + 1}",
                "totalCount": total,
                "goodCount": good,
                "badCount": bad,
                "badRate": round(bad / max(total, 1), 6),
                "distGood": round(dist_good, 6),
                "distBad": round(dist_bad, 6),
                "woe": round(woe, 6),
                "ivContribution": round(iv_contribution, 6),
                "smoothingApplied": smoothing_applied,
                "manualOverride": bool(definition.get("manualOverride")),
                "reviewerNote": definition.get("reviewerNote"),
            }
        )
    return bins


def monotonicity_status(bins: List[Dict[str, Any]]) -> str:
    if len(bins) < 3:
        return "not_applicable"
    rates = [float(item.get("badRate") or 0) for item in bins]
    if all(rates[i] <= rates[i + 1] for i in range(len(rates) - 1)):
        return "monotonic_increasing"
    if all(rates[i] >= rates[i + 1] for i in range(len(rates) - 1)):
        return "monotonic_decreasing"
    return "not_monotonic"


def predictive_strength(iv: float) -> str:
    if iv < 0.02:
        return "not_useful"
    if iv < 0.10:
        return "weak"
    if iv < 0.30:
        return "medium"
    if iv < 0.50:
        return "strong"
    return "suspicious"


def recommendation(iv: float, monotonicity: str, warnings: List[str]) -> str:
    strength = predictive_strength(iv)
    if warnings:
        return "Review warnings before approving characteristic bins."
    if strength == "suspicious":
        return "High IV may indicate leakage or instability. Review before use."
    if monotonicity == "not_monotonic":
        return "Non-monotonic pattern requires business review or manual binning."
    if strength in {"medium", "strong"}:
        return "Candidate underwriting characteristic."
    return "Weak characteristic. Retain only with business rationale."


def apply_manual_bins(payload: Dict[str, Any]) -> Dict[str, Any]:
    analysis = get_analysis(str(payload.get("analysisId")))
    if not analysis:
        raise HTTPException(status_code=404, detail="WOE / IV analysis not found.")
    variable = str(payload.get("variable") or "")
    dataset_id = str(analysis.get("datasetId"))
    version_id = str(analysis.get("datasetVersionId") or "")
    _, _, df = load_dataset_frame(dataset_id, version_id)
    y = event_mask(df[str(analysis.get("targetColumn"))], analysis.get("eventValue"))
    bins = payload.get("bins") or []
    warnings = validate_manual_bins(bins, bool(payload.get("allowNumericGaps")))
    series = df[variable]
    assignments = pd.Series([None] * len(series), index=series.index, dtype=object)
    definitions = []
    numeric = pd.to_numeric(series, errors="coerce")
    for index, item in enumerate(bins):
        bin_id = item.get("binId") or f"manual_{index + 1}"
        definition = {**item, "binId": bin_id, "manualOverride": True}
        definitions.append(definition)
        if item.get("isMissingBin"):
            assignments[series.isna()] = bin_id
        elif item.get("categoryValues"):
            assignments[series.astype(str).isin([str(value) for value in item.get("categoryValues") or []])] = bin_id
        else:
            lower = item.get("lowerBound")
            upper = item.get("upperBound")
            mask = pd.Series([True] * len(series), index=series.index)
            if lower is not None:
                mask &= numeric.ge(float(lower))
            if upper is not None:
                mask &= numeric.lt(float(upper))
            assignments[mask] = bin_id
    if assignments.isna().any():
        warnings.append("Manual bins have gaps.")
        assignments = assignments.fillna("unassigned")
        definitions.append({"binId": "unassigned", "label": "Unassigned / gap", "manualOverride": True})
    bin_stats = compute_bin_stats(assignments, definitions, y, float(analysis.get("settings", {}).get("smoothing") or 0.5), warnings)
    updated = None
    for summary in analysis.get("variableSummaries") or []:
        if summary.get("variable") == variable:
            summary["bins"] = bin_stats
            summary["iv"] = round(sum(float(item.get("ivContribution") or 0) for item in bin_stats), 6)
            summary["predictiveStrength"] = predictive_strength(float(summary["iv"]))
            summary["monotonicity"] = monotonicity_status([item for item in bin_stats if not item.get("isMissingBin")])
            summary["status"] = "needs_review" if warnings or summary["monotonicity"] == "not_monotonic" else "ready"
            summary["recommendation"] = recommendation(float(summary["iv"]), str(summary["monotonicity"]), warnings)
            summary["warnings"] = dedupe(warnings)
            summary["manualBinsLocked"] = True
            updated = summary
            break
    if not updated:
        raise HTTPException(status_code=404, detail="Variable not found in WOE / IV analysis.")
    analysis["manualOverrides"] = {**(analysis.get("manualOverrides") or {}), variable: bins}
    analysis["generatedCode"] = generate_code({**analysis.get("settings", {}), **analysis, "manualBins": analysis.get("manualOverrides")})
    return save_analysis(analysis)


def validate_manual_bins(bins: List[Dict[str, Any]], allow_gaps: bool) -> List[str]:
    warnings: List[str] = []
    numeric_ranges = []
    category_seen = set()
    missing_bins = 0
    for item in bins:
        if item.get("isMissingBin"):
            missing_bins += 1
        if item.get("categoryValues"):
            for value in item.get("categoryValues") or []:
                if value in category_seen:
                    warnings.append("Category values cannot be assigned to multiple bins.")
                category_seen.add(value)
        if item.get("lowerBound") is not None or item.get("upperBound") is not None:
            lower = float(item.get("lowerBound") if item.get("lowerBound") is not None else -math.inf)
            upper = float(item.get("upperBound") if item.get("upperBound") is not None else math.inf)
            if lower >= upper:
                warnings.append("Manual bins overlap or have invalid bounds.")
            numeric_ranges.append((lower, upper))
    numeric_ranges.sort()
    for index in range(len(numeric_ranges) - 1):
        if numeric_ranges[index][1] > numeric_ranges[index + 1][0]:
            warnings.append("Manual bins overlap.")
        if not allow_gaps and numeric_ranges[index][1] < numeric_ranges[index + 1][0]:
            warnings.append("Manual bins have gaps.")
    if missing_bins > 1:
        warnings.append("Missing value treatment must be explicit and unique.")
    if len(bins) < 2:
        warnings.append("At least two bins required unless variable is excluded.")
    return dedupe(warnings)


def create_woe_iv_evidence(analysis_id: str, technique: str = "woe_iv", label: str = "WOE / IV analysis") -> Dict[str, Any]:
    analysis = get_analysis(analysis_id)
    if not analysis:
        raise HTTPException(status_code=404, detail="WOE / IV analysis not found.")
    top = sorted(analysis.get("variableSummaries") or [], key=lambda row: float(row.get("iv") or 0), reverse=True)[:10]
    needs_review = [row.get("variable") for row in analysis.get("variableSummaries") or [] if row.get("status") == "needs_review"]
    supports = ["Feature Engineering", "Correlation and VIF Review", "Model Interpretation", "Appendices"]
    if technique == "manual_binning":
        supports = ["Feature Engineering", "Model Interpretation", "Fair Lending and Governance Considerations", "Appendices"]
    item = {
        "projectId": analysis.get("projectId") or "",
        "projectName": "",
        "evidenceType": "feature_engineering",
        "financialDomain": "credit_risk",
        "financialTechnique": technique,
        "title": label,
        "summary": f"WOE / IV analysis reviewed {len(analysis.get('variableSummaries') or [])} candidate underwriting characteristics for target {analysis.get('targetColumn')}.",
        "sourcePage": "Credit Risk Studio",
        "sourceRoute": "/credit-risk-studio",
        "artifactType": "woe_iv_analysis",
        "artifactId": analysis_id,
        "datasetIds": [analysis.get("datasetId")] if analysis.get("datasetId") else [],
        "datasetVersionIds": [analysis.get("datasetVersionId")] if analysis.get("datasetVersionId") else [],
        "supportsSections": supports,
        "tags": ["credit_risk", "woe", "iv", "manual_binning", "underwriting"],
        "creditRiskContext": {
            "targetColumn": analysis.get("targetColumn"),
            "eventValue": analysis.get("eventValue"),
            "eventRate": analysis.get("eventRate"),
        },
        "governance": {"requiresReview": True, "fairLendingRelevant": True, "modelRiskRelevant": True, "explainability": True},
        "payload": {
            "targetColumn": analysis.get("targetColumn"),
            "eventValue": analysis.get("eventValue"),
            "eventRate": analysis.get("eventRate"),
            "variableCount": len(analysis.get("variableSummaries") or []),
            "topVariablesByIv": [{"variable": row.get("variable"), "iv": row.get("iv")} for row in top],
            "variablesNeedingReview": needs_review,
        },
        "status": "ready",
        "importance": "high",
    }
    evidence = upsert_evidence(item)
    analysis["evidenceIds"] = list(set((analysis.get("evidenceIds") or []) + [evidence.get("evidenceId")]))
    save_analysis(analysis)
    return evidence


def export_woe_iv_excel(analysis_id: str) -> bytes:
    analysis = get_analysis(analysis_id)
    if not analysis:
        raise HTTPException(status_code=404, detail="WOE / IV analysis not found.")
    try:
        from openpyxl import Workbook
    except Exception:
        return json.dumps(analysis, indent=2, default=str).encode("utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for row in [
        ("Analysis ID", analysis.get("analysisId")),
        ("Dataset", analysis.get("datasetName") or analysis.get("datasetId")),
        ("Target", analysis.get("targetColumn")),
        ("Event Value", analysis.get("eventValue")),
        ("Rows", analysis.get("rowCount")),
        ("Event Rate", analysis.get("eventRate")),
    ]:
        ws.append(row)
    ranking = wb.create_sheet("IV Ranking")
    ranking.append(["Variable", "IV", "Strength", "Monotonicity", "Missing %", "Distinct", "Bins", "Status", "Recommendation"])
    for summary in analysis.get("variableSummaries") or []:
        ranking.append([summary.get("variable"), summary.get("iv"), summary.get("predictiveStrength"), summary.get("monotonicity"), summary.get("missingRate"), summary.get("distinctCount"), summary.get("binCount"), summary.get("status"), summary.get("recommendation")])
    tables = wb.create_sheet("WOE Tables")
    tables.append(["Variable", "Bin", "Label", "Total", "Good", "Bad", "Bad Rate", "Dist Good", "Dist Bad", "WOE", "IV Contribution", "Smoothing"])
    for summary in analysis.get("variableSummaries") or []:
        for bin_row in summary.get("bins") or []:
            tables.append([summary.get("variable"), bin_row.get("binId"), bin_row.get("label"), bin_row.get("totalCount"), bin_row.get("goodCount"), bin_row.get("badCount"), bin_row.get("badRate"), bin_row.get("distGood"), bin_row.get("distBad"), bin_row.get("woe"), bin_row.get("ivContribution"), bin_row.get("smoothingApplied")])
    for sheet_name in ["Manual Bins", "Variables Needing Review", "Monotonicity Review", "Evidence Mapping", "Code References"]:
        sheet = wb.create_sheet(sheet_name)
        sheet.append(["Item", "Value"])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def generate_code(settings: Dict[str, Any]) -> Dict[str, str]:
    target = settings.get("targetColumn") or "target"
    event = settings.get("eventValue", 1)
    predictors = settings.get("predictorColumns") or []
    first_var = predictors[0] if predictors else "variable_name"
    method = settings.get("binningMethod") or "quantile"
    max_bins = settings.get("maxBins") or 5
    return {
        "python": f"""import numpy as np\nimport pandas as pd\n\n# ATLAS WOE / IV analysis\n# Dataset: {settings.get('datasetId')}\ntarget = {target!r}\nevent_value = {event!r}\npredictors = {predictors!r}\n\n# Bad = event, Good = non-event. WOE = ln(dist_good / dist_bad).\ndf['_bad'] = (df[target].astype(str) == str(event_value)).astype(int)\nfor variable in predictors:\n    df['_bin'] = pd.qcut(df[variable].rank(method='first'), q={max_bins}, duplicates='drop')\n    grouped = df.groupby('_bin').agg(total=(target, 'size'), bad=('_bad', 'sum'))\n    grouped['good'] = grouped['total'] - grouped['bad']\n    grouped['dist_good'] = (grouped['good'] + 0.5) / (grouped['good'].sum() + 0.5 * len(grouped))\n    grouped['dist_bad'] = (grouped['bad'] + 0.5) / (grouped['bad'].sum() + 0.5 * len(grouped))\n    grouped['woe'] = np.log(grouped['dist_good'] / grouped['dist_bad'])\n    grouped['iv_contribution'] = (grouped['dist_good'] - grouped['dist_bad']) * grouped['woe']\n    print(variable, grouped['iv_contribution'].sum())\n""",
        "sql": f"""-- ATLAS WOE / IV skeleton for {first_var}\nwith binned as (\n  select *, ntile({max_bins}) over (order by {first_var}) as {first_var}_bin\n  from input_table\n), counts as (\n  select {first_var}_bin,\n         count(*) as total_count,\n         sum(case when {target} = {json.dumps(event)} then 1 else 0 end) as bad_count,\n         sum(case when {target} <> {json.dumps(event)} then 1 else 0 end) as good_count\n  from binned\n  group by {first_var}_bin\n)\nselect * from counts;\n""",
        "r": f"""library(dplyr)\n# ATLAS WOE / IV analysis\ninput <- input %>% mutate(.bad = as.integer(as.character({target}) == as.character({json.dumps(event)})))\n# Use ntile for {method} binning, then calculate WOE = log(dist_good / dist_bad).\n""",
        "sas": f"""/* ATLAS SAS-compatible WOE / IV code. No SAS license is required to generate this code. */\nproc rank data=work.input groups={max_bins} out=work.binned;\n  var {first_var};\n  ranks {first_var}_bin;\nrun;\n\nproc sql;\n  create table work.woe_iv as\n  select\n    {first_var}_bin,\n    count(*) as total_count,\n    sum(case when {target} = {event} then 1 else 0 end) as bad_count,\n    sum(case when {target} ne {event} then 1 else 0 end) as good_count\n  from work.binned\n  group by {first_var}_bin;\nquit;\n""",
    }


def dedupe(values: List[str]) -> List[str]:
    output = []
    for value in values:
        if value and value not in output:
            output.append(value)
    return output
