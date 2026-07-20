from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Iterable, List

from .evidence_types import DOCUMENTATION_SECTIONS, EVIDENCE_IMPORTANCE, EVIDENCE_STATUSES, EVIDENCE_TYPES
from .financial_services_taxonomy import (
    CREDIT_RISK_DOCUMENTATION_MAPPING,
    credit_risk_checklist,
    infer_financial_domain,
)

ARTIFACT_ROOT = Path(__file__).resolve().parents[1] / "artifacts"
EVIDENCE_ROOT = ARTIFACT_ROOT / "evidence"
ITEMS_ROOT = EVIDENCE_ROOT / "items"
REGISTRY_PATH = EVIDENCE_ROOT / "evidence_registry.json"

COVERAGE_CATEGORIES = [
    "dataset",
    "target",
    "eda",
    "feature_engineering",
    "training",
    "evaluation",
    "diagnostics",
    "interpretation",
    "scoring",
    "governance",
    "documentation",
]


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def ensure_store() -> None:
    ITEMS_ROOT.mkdir(parents=True, exist_ok=True)
    if not REGISTRY_PATH.exists():
        REGISTRY_PATH.write_text(json.dumps({"items": []}, indent=2), encoding="utf-8")


def read_json(path: Path, default: Any) -> Any:
    try:
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default
    return default


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")


def evidence_id_for(payload: Dict[str, Any]) -> str:
    basis = "|".join(
        str(payload.get(key, ""))
        for key in ["projectId", "evidenceType", "sourcePage", "artifactType", "artifactId", "title"]
    )
    digest = hashlib.sha256(basis.encode("utf-8")).hexdigest()[:16]
    slug = re.sub(r"[^a-z0-9]+", "-", str(payload.get("title") or payload.get("evidenceType") or "evidence").lower()).strip("-")[:42]
    return f"ev-{slug or 'item'}-{digest}"


def compact_item(item: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "evidenceId": item["evidenceId"],
        "projectId": item.get("projectId", ""),
        "projectName": item.get("projectName", ""),
        "evidenceType": item.get("evidenceType", "documentation"),
        "title": item.get("title", "Untitled evidence"),
        "summary": item.get("summary", ""),
        "sourcePage": item.get("sourcePage", ""),
        "sourceRoute": item.get("sourceRoute", ""),
        "artifactType": item.get("artifactType", ""),
        "artifactId": item.get("artifactId", ""),
        "datasetIds": item.get("datasetIds", []),
        "modelIds": item.get("modelIds", []),
        "documentIds": item.get("documentIds", []),
        "notebookIds": item.get("notebookIds", []),
        "supportsSections": item.get("supportsSections", []),
        "tags": item.get("tags", []),
        "financialDomain": item.get("financialDomain", ""),
        "financialTechnique": item.get("financialTechnique", ""),
        "regulatoryRelevance": item.get("regulatoryRelevance", {}),
        "creditRiskContext": item.get("creditRiskContext", {}),
        "status": item.get("status", "draft"),
        "importance": item.get("importance", "medium"),
        "qualityScore": item.get("qualityScore", quality_score(item)),
        "createdAt": item.get("createdAt", now_iso()),
        "capturedAt": item.get("capturedAt", ""),
        "updatedAt": item.get("updatedAt", item.get("createdAt", now_iso())),
        "createdBy": item.get("createdBy", "ATLAS"),
        "workflowStep": item.get("workflowStep", ""),
        "sourceRunReference": item.get("sourceRunReference", ""),
    }


def normalize_item(payload: Dict[str, Any], existing: Dict[str, Any] | None = None) -> Dict[str, Any]:
    item = {**(existing or {}), **payload}
    item["evidenceType"] = item.get("evidenceType") if item.get("evidenceType") in EVIDENCE_TYPES else "documentation"
    item["status"] = item.get("status") if item.get("status") in EVIDENCE_STATUSES else "draft"
    item["importance"] = item.get("importance") if item.get("importance") in EVIDENCE_IMPORTANCE else "medium"
    item["evidenceId"] = item.get("evidenceId") or evidence_id_for(item)
    item["title"] = item.get("title") or item["evidenceType"].replace("_", " ").title()
    item["summary"] = item.get("summary") or "Evidence captured by ATLAS."
    item["supportsSections"] = item.get("supportsSections") or DOCUMENTATION_SECTIONS.get(item["evidenceType"], [])
    for key in ["datasetIds", "modelIds", "documentIds", "notebookIds", "tags"]:
        value = item.get(key, [])
        item[key] = value if isinstance(value, list) else [value]
    item["lineage"] = item.get("lineage") if isinstance(item.get("lineage"), dict) else {}
    item["payload"] = item.get("payload") if isinstance(item.get("payload"), dict) else {}
    item["governance"] = item.get("governance") if isinstance(item.get("governance"), dict) else {}
    item["regulatoryRelevance"] = item.get("regulatoryRelevance") if isinstance(item.get("regulatoryRelevance"), dict) else {}
    item["creditRiskContext"] = item.get("creditRiskContext") if isinstance(item.get("creditRiskContext"), dict) else {}
    if item.get("financialTechnique") and not payload.get("supportsSections"):
        item["supportsSections"] = CREDIT_RISK_DOCUMENTATION_MAPPING.get(str(item.get("financialTechnique")), item["supportsSections"])
    item["createdAt"] = item.get("createdAt") or now_iso()
    if item["status"] in {"ready", "accepted"}:
        item["capturedAt"] = item.get("capturedAt") or now_iso()
    item["updatedAt"] = now_iso()
    item["createdBy"] = item.get("createdBy") or "ATLAS"
    item["qualityScore"] = quality_score(item)
    if item["qualityScore"] >= 85 and item["status"] == "draft":
        item["status"] = "ready"
    elif item["qualityScore"] < 55 and item["status"] == "ready":
        item["status"] = "needs_review"
    return item


def load_registry() -> List[Dict[str, Any]]:
    ensure_store()
    payload = read_json(REGISTRY_PATH, {"items": []})
    return payload.get("items", []) if isinstance(payload, dict) else []


def save_registry(items: List[Dict[str, Any]]) -> None:
    write_json(REGISTRY_PATH, {"items": sorted(items, key=lambda row: row.get("updatedAt", ""), reverse=True)})


def item_path(evidence_id: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", evidence_id)
    return ITEMS_ROOT / f"{safe}.json"


def upsert_evidence(payload: Dict[str, Any]) -> Dict[str, Any]:
    ensure_store()
    proposed_id = payload.get("evidenceId") or evidence_id_for(payload)
    existing = get_evidence(proposed_id)
    item = normalize_item({**payload, "evidenceId": proposed_id}, existing)
    write_json(item_path(item["evidenceId"]), item)
    registry = [row for row in load_registry() if row.get("evidenceId") != item["evidenceId"]]
    registry.append(compact_item(item))
    save_registry(registry)
    return item


def get_evidence(evidence_id: str) -> Dict[str, Any] | None:
    item = read_json(item_path(evidence_id), None)
    return item if isinstance(item, dict) else None


def delete_evidence(evidence_id: str) -> bool:
    path = item_path(evidence_id)
    existed = path.exists()
    if existed:
        path.unlink()
    save_registry([row for row in load_registry() if row.get("evidenceId") != evidence_id])
    return existed


def list_evidence(filters: Dict[str, str | None]) -> List[Dict[str, Any]]:
    rows = load_registry()
    search = (filters.get("search") or "").lower()
    for key in [
        "projectId",
        "datasetId",
        "evidenceType",
        "sourcePage",
        "status",
        "importance",
        "tag",
        "financialDomain",
        "financialTechnique",
        "productType",
        "modelPurpose",
        "regulatoryRelevance",
    ]:
        value = filters.get(key)
        if not value:
            continue
        if key == "datasetId":
            rows = [row for row in rows if value in row.get("datasetIds", [])]
        elif key == "tag":
            rows = [row for row in rows if value in row.get("tags", [])]
        elif key == "productType":
            rows = [row for row in rows if str(row.get("creditRiskContext", {}).get("productType", "")) == value]
        elif key == "modelPurpose":
            rows = [row for row in rows if str(row.get("creditRiskContext", {}).get("modelPurpose", "")) == value]
        elif key == "regulatoryRelevance":
            rows = [row for row in rows if bool(row.get("regulatoryRelevance", {}).get(str(value), False))]
        else:
            rows = [row for row in rows if str(row.get(key, "")) == value]
    if search:
        rows = [
            row
            for row in rows
            if search in " ".join(str(row.get(key, "")) for key in ["title", "summary", "sourcePage", "projectName"]).lower()
        ]
    return rows


def quality_score(item: Dict[str, Any]) -> int:
    checks = [
        bool(item.get("sourcePage") or item.get("sourceRoute")),
        bool(item.get("projectId")),
        bool(item.get("datasetIds")),
        bool(item.get("summary")),
        bool(item.get("lineage")),
        bool(item.get("supportsSections")),
        bool(item.get("governance")),
    ]
    return round(sum(1 for check in checks if check) / len(checks) * 100)


def graph(project_id: str | None = None) -> Dict[str, Any]:
    rows = list_evidence({"projectId": project_id} if project_id else {})
    nodes: Dict[str, Dict[str, Any]] = {}
    edges: List[Dict[str, Any]] = []
    inferred_project_id = project_id or next((str(row.get("projectId") or "") for row in rows if row.get("projectId")), "")
    inferred_project_name = next((str(row.get("projectName") or "") for row in rows if row.get("projectName")), "") or inferred_project_id or "Current ATLAS Project"
    project_node_id = f"project:{inferred_project_id or 'current'}"
    nodes[project_node_id] = {
        "id": project_node_id,
        "label": inferred_project_name,
        "kind": "project",
        "type": "Project",
        "status": "ready" if rows else "draft",
        "importance": "critical",
        "metadata": {"projectId": inferred_project_id, "projectName": inferred_project_name},
    }

    export_node_id = f"export:{inferred_project_id or 'current'}"
    nodes[export_node_id] = {
        "id": export_node_id,
        "label": "Evidence Pack Export",
        "kind": "export_package",
        "type": "Export Package",
        "status": "ready" if rows else "draft",
        "importance": "high",
        "metadata": {"description": "Governance-ready evidence package"},
    }
    edges.append(_edge(project_node_id, export_node_id, "produces", "produces"))

    for row in rows:
        eid = row["evidenceId"]
        node_type = _graph_node_type(row)
        nodes[eid] = {
            "id": eid,
            "label": row["title"],
            "kind": row["evidenceType"],
            "type": node_type,
            "status": row["status"],
            "importance": row.get("importance", "medium"),
            "sourcePage": row.get("sourcePage", ""),
            "evidenceId": eid,
            "metadata": {
                "summary": row.get("summary", ""),
                "datasetIds": row.get("datasetIds", []),
                "supportsSections": row.get("supportsSections", []),
                "financialDomain": row.get("financialDomain", ""),
                "financialTechnique": row.get("financialTechnique", ""),
                "sourceRoute": row.get("sourceRoute", ""),
            },
        }
        edges.append(_edge(project_node_id, eid, "linked to project", "linked_to_project"))
        for dataset_id in row.get("datasetIds", []):
            did = f"dataset:{dataset_id}"
            nodes.setdefault(
                did,
                {
                    "id": did,
                    "label": dataset_id,
                    "kind": "dataset",
                    "type": "Dataset",
                    "status": "ready",
                    "importance": "high",
                    "metadata": {"datasetId": dataset_id},
                },
            )
            edges.append(_edge(project_node_id, did, "uses dataset", "uses_dataset"))
            edges.append(_edge(did, eid, "generated from", "generated_from"))
        for section in row.get("supportsSections", []):
            sid = f"section:{section}"
            nodes.setdefault(
                sid,
                {
                    "id": sid,
                    "label": section,
                    "kind": "documentation_section",
                    "type": "Documentation Section",
                    "status": "ready",
                    "importance": "high",
                    "metadata": {"section": section},
                },
            )
            edges.append(_edge(eid, sid, "supports section", "supports_section"))
            edges.append(_edge(sid, export_node_id, "documents", "documents"))
        if row.get("status") == "needs_review" or row.get("importance") == "critical" or row.get("regulatoryRelevance"):
            gid = f"governance:{eid}"
            nodes[gid] = {
                "id": gid,
                "label": "Governance Review",
                "kind": "governance_warning",
                "type": "Governance Warning",
                "status": "needs_review" if row.get("status") == "needs_review" else "ready",
                "importance": row.get("importance", "high"),
                "evidenceId": eid,
                "metadata": {"sourceEvidenceId": eid},
            }
            edges.append(_edge(eid, gid, "governance flag", "governance_flag"))
            edges.append(_edge(gid, export_node_id, "requires review", "requires_review"))
    return {"nodes": list(nodes.values()), "edges": edges}


def _edge(source: str, target: str, label: str, edge_type: str) -> Dict[str, Any]:
    edge_id = hashlib.sha256(f"{source}|{target}|{edge_type}".encode("utf-8")).hexdigest()[:16]
    return {"id": f"edge:{edge_id}", "source": source, "target": target, "from": source, "to": target, "label": label, "type": edge_type}


def _graph_node_type(row: Dict[str, Any]) -> str:
    if row.get("financialDomain") == "credit_risk" or "credit_risk" in row.get("tags", []):
        return "Credit Risk Evidence"
    evidence_type = row.get("evidenceType", "")
    if str(evidence_type).startswith("optimization_"):
        return "Optimization Evidence"
    return {
        "dataset": "Dataset",
        "sql_query": "SQL Query",
        "sql_result": "SQL Result",
        "eda": "Analytics Insight",
        "chart": "Chart",
        "notebook": "Notebook",
        "model_training": "Model Training",
        "model_selection": "Model Training",
        "evaluation": "Evaluation",
        "model_diagnostics": "Model Diagnostics",
        "documentation": "Documentation Section",
        "governance": "Governance Warning",
        "decision": "Governance Warning",
        "export": "Export Package",
    }.get(evidence_type, "Analytics Insight")


def coverage(project_id: str) -> Dict[str, Any]:
    rows = list_evidence({"projectId": project_id})
    optimization_rows = [row for row in rows if str(row.get("evidenceType", "")).startswith("optimization_")]
    if optimization_rows:
        optimization_types = [
            "optimization_decision_charter",
            "optimization_input_bundle",
            "optimization_formulation_readiness",
            "optimization_objective",
            "optimization_decision_variables",
            "optimization_constraints",
            "optimization_solver_configuration",
            "optimization_scenarios",
            "optimization_solution",
            "optimization_sensitivity",
            "optimization_validation",
            "optimization_recommendation",
            "optimization_implementation_plan",
            "optimization_monitoring_plan",
            "optimization_documentation_package",
            "optimization_reproducibility_code",
        ]
        categories = []
        for evidence_type in optimization_types:
            aliases = {evidence_type}
            if evidence_type == "optimization_reproducibility_code":
                aliases.add("optimization_code_reproducibility")
            matches = [row for row in optimization_rows if row.get("evidenceType") in aliases]
            ready = [row for row in matches if row.get("status") in {"ready", "accepted"}]
            categories.append(
                {
                    "category": evidence_type,
                    "count": len(matches),
                    "ready": len(ready),
                    "status": "ready" if ready else "missing" if not matches else "needs_review",
                }
            )
        ready_categories = sum(1 for item in categories if item["status"] == "ready")
        missing = [item["category"] for item in categories if item["status"] == "missing"]
        warnings = [] if not missing else [f"{len(missing)} optimization evidence artifact(s) remain uncaptured or unsupported by the documentation package."]
        return {
            "projectId": project_id,
            "workflowType": "optimization",
            "totalEvidence": len(optimization_rows),
            "readyEvidence": sum(1 for row in optimization_rows if row.get("status") in {"ready", "accepted"}),
            "needsReviewEvidence": sum(1 for row in optimization_rows if row.get("status") == "needs_review"),
            "coverageScore": round(ready_categories / len(categories) * 100),
            "categories": categories,
            "warnings": warnings,
            "missingRecommended": missing,
            "financialServices": {
                "suggestedDomain": infer_financial_domain(target=project_id),
                "creditRisk": credit_risk_checklist(rows, project_id),
                "placeholders": {},
            },
        }
    category_map = {
        "dataset": {"dataset", "sql_result"},
        "target": {"dataset", "decision"},
        "eda": {"eda", "chart"},
        "feature_engineering": {"feature_engineering", "vif"},
        "training": {"model_training", "model_selection"},
        "evaluation": {"evaluation"},
        "diagnostics": {"model_diagnostics", "vif"},
        "interpretation": {"model_diagnostics", "chart", "notebook"},
        "scoring": {"evaluation", "model_diagnostics"},
        "governance": {"governance", "decision"},
        "documentation": {"documentation", "export"},
    }
    categories = []
    for category in COVERAGE_CATEGORIES:
        matches = [row for row in rows if row.get("evidenceType") in category_map[category]]
        ready = [row for row in matches if row.get("status") in {"ready", "accepted"}]
        categories.append(
            {
                "category": category,
                "count": len(matches),
                "ready": len(ready),
                "status": "ready" if ready else "missing" if not matches else "needs_review",
            }
        )
    ready_categories = sum(1 for item in categories if item["status"] == "ready")
    warnings = []
    if not any("fair_lending" in row.get("tags", []) for row in rows):
        warnings.append("Fair Lending evidence has not been linked to this project.")
    if not any(row.get("evidenceType") == "governance" for row in rows):
        warnings.append("Governance rationale evidence is missing.")
    return {
        "projectId": project_id,
        "totalEvidence": len(rows),
        "readyEvidence": sum(1 for row in rows if row.get("status") in {"ready", "accepted"}),
        "needsReviewEvidence": sum(1 for row in rows if row.get("status") == "needs_review"),
        "coverageScore": round(ready_categories / len(categories) * 100),
        "categories": categories,
        "warnings": warnings,
        "missingRecommended": [item["category"] for item in categories if item["status"] == "missing"],
        "financialServices": {
            "suggestedDomain": infer_financial_domain(target=project_id),
            "creditRisk": credit_risk_checklist(rows, project_id),
            "placeholders": {
                "fraud_aml": "Fraud / AML taxonomy placeholder is registered for the next domain build.",
                "marketing_profitability": "Marketing / Profitability taxonomy placeholder is registered for the next domain build.",
            },
        },
    }


def auto_capture(project_id: str | None = None, project_name: str | None = None) -> List[Dict[str, Any]]:
    captured: List[Dict[str, Any]] = []
    for dataset in _iter_registry_records(ARTIFACT_ROOT / "datasets" / "registry.json"):
        dataset_id = str(dataset.get("datasetId") or dataset.get("id") or "")
        if not dataset_id:
            continue
        linked_projects = dataset.get("linkedProjects") if isinstance(dataset.get("linkedProjects"), list) else []
        if project_id and linked_projects and not any(link.get("projectId") == project_id for link in linked_projects):
            continue
        captured.append(
            upsert_evidence(
                {
                    "projectId": project_id or _first_linked_project(linked_projects),
                    "projectName": project_name or "",
                    "evidenceType": "dataset",
                    "title": f"Dataset: {dataset.get('name') or dataset_id}",
                    "summary": dataset.get("description") or "Governed dataset registered in the ATLAS Dataset Library.",
                    "sourcePage": "Dataset Library",
                    "sourceRoute": f"/datasets/{dataset_id}",
                    "artifactType": "dataset",
                    "artifactId": dataset_id,
                    "datasetIds": [dataset_id],
                    "tags": [tag for tag in ["model_risk", str(dataset.get("domain", "")).lower()] if tag],
                    "lineage": {"registry": "datasets/registry.json"},
                    "governance": {"retention": "Reference only; raw dataset is not duplicated in Evidence Graph."},
                }
            )
        )
    return captured


def _iter_registry_records(path: Path) -> Iterable[Dict[str, Any]]:
    payload = read_json(path, {})
    if isinstance(payload, list):
        yield from (row for row in payload if isinstance(row, dict))
    elif isinstance(payload, dict):
        for key in ["datasets", "items", "notebooks", "documents", "insights"]:
            value = payload.get(key)
            if isinstance(value, list):
                yield from (row for row in value if isinstance(row, dict))
        for value in payload.values():
            if isinstance(value, dict) and any(name in value for name in ["datasetId", "notebookId", "documentId", "insightId"]):
                yield value


def _first_linked_project(links: Any) -> str:
    if isinstance(links, list) and links and isinstance(links[0], dict):
        return str(links[0].get("projectId") or "")
    return ""


def export_excel(project_id: str | None = None) -> bytes:
    from openpyxl import Workbook

    rows = list_evidence({"projectId": project_id} if project_id else {})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evidence Registry"
    headers = ["Title", "Type", "Reference ID", "Status", "Project", "Workflow Step", "Created", "Captured", "Datasets", "Source Run", "Source", "Importance", "Quality", "Sections"]
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                row.get("title", ""),
                row.get("evidenceType", ""),
                row.get("evidenceId", ""),
                row.get("status", ""),
                row.get("projectName", "") or row.get("projectId", ""),
                row.get("workflowStep", ""),
                row.get("createdAt", ""),
                row.get("capturedAt", ""),
                ", ".join(row.get("datasetIds", [])),
                row.get("sourceRunReference", ""),
                row.get("sourcePage", ""),
                row.get("importance", ""),
                row.get("qualityScore", ""),
                ", ".join(row.get("supportsSections", [])),
            ]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
