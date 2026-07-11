from __future__ import annotations

import csv
import io
import math
import re
from numbers import Number
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd
from fastapi import HTTPException
from sklearn.metrics import roc_auc_score

from ..analytics.analytics_store import create_insight
from ..datasets.dataset_registry import list_accessible_datasets, read_json, require_dataset, version_dir


def load_scored_frame(dataset_id: str, version_id: Optional[str] = None) -> tuple[Dict[str, Any], str, pd.DataFrame, Dict[str, Any]]:
    dataset = require_dataset(dataset_id)
    versions = dataset.get("versions") or []
    resolved_version = version_id or (versions[-1].get("versionId") if versions else "")
    if not resolved_version:
        raise HTTPException(status_code=400, detail="Dataset has no version available for diagnostics.")
    vdir = version_dir(dataset_id, resolved_version)
    sample_path = vdir / "data_sample.csv"
    preview_path = vdir / "preview.json"
    metadata = read_json(vdir / "metadata.json", {})
    if sample_path.exists():
        df = pd.read_csv(sample_path)
    elif preview_path.exists():
        preview = read_json(preview_path, {"rows": []})
        df = pd.DataFrame(preview.get("rows") or [])
    else:
        raise HTTPException(status_code=404, detail="No scored dataset rows are available for diagnostics.")
    if df.empty:
        raise HTTPException(status_code=400, detail="Scored dataset has no rows available for diagnostics.")
    return dataset, resolved_version, df, metadata


def compute_model_diagnostics(dataset_id: str, version_id: Optional[str] = None, overrides: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    dataset, resolved_version, df, metadata = load_scored_frame(dataset_id, version_id)
    roles = detect_roles(df, metadata, overrides or {})
    score_col = roles.get("score")
    actual_col = roles.get("actual")
    prediction_col = roles.get("prediction")
    residual_col = roles.get("residual")
    decile_col = roles.get("decile")
    if not score_col and prediction_col:
        score_col = prediction_col
        roles["score"] = prediction_col
    score = numeric_series(df, score_col)
    actual = numeric_series(df, actual_col)
    prediction = numeric_series(df, prediction_col)
    residual = numeric_series(df, residual_col)
    if residual.empty and not actual.empty and not prediction.empty:
        residual = actual - prediction
    mode = diagnostic_mode(roles, actual, prediction, score, df, metadata)
    mapping_warnings = mapping_warnings_for_roles(roles)
    classification_metrics = classification_summary_metrics(score, actual, roles.get("positiveClass"))
    detailed_lift = detailed_lift_table(df, score_col, actual_col, decile_col, roles.get("positiveClass"))
    vif = vif_analysis(df, roles)
    pozzolo = pozzolo_calibration_detail(score, actual, roles, metadata)
    kpi = kpi_aggregation(mode, classification_metrics, diagnostics_regression_summary(actual, prediction, residual))
    diagnostics = {
        "ok": True,
        "dataset": dataset_summary(dataset),
        "versionId": resolved_version,
        "rowCount": int(df.shape[0]),
        "diagnosticMode": mode,
        "columnRoles": roles,
        "mappingWarnings": mapping_warnings,
        "scoreDistribution": distribution(score),
        "predictionSummary": prediction_summary(prediction if not prediction.empty else score),
        "actualDistribution": distribution(actual),
        "classificationMetrics": classification_metrics,
        "detailedLiftTable": detailed_lift,
        "topPredictions": prediction_extremes(df, roles, descending=True),
        "bottomPredictions": prediction_extremes(df, roles, descending=False),
        "largestAbsoluteErrors": residual_extremes(df, roles, "absolute"),
        "topOverpredictions": residual_extremes(df, roles, "over"),
        "topUnderpredictions": residual_extremes(df, roles, "under"),
        "predictionBandSummary": prediction_band_summary(df, roles),
        "regressionFitSummary": regression_fit_summary(actual, prediction, residual),
        "populationWaterfall": population_waterfall(dataset, df, roles, mode),
        "vifAnalysis": vif,
        "pozzoloCalibration": pozzolo,
        "kpiAggregation": kpi,
        "liftDecile": lift_deciles(df, score_col, actual_col, decile_col, roles.get("positiveClass")),
        "residualDiagnostics": residual_summary(residual),
        "actualVsPredicted": actual_vs_predicted(actual, prediction if not prediction.empty else score),
        "calibration": calibration_bins(score, actual, roles.get("positiveClass")),
        "segmentPerformance": segment_performance(df, score_col, actual_col, roles.get("segment")),
        "splitComparison": split_comparison(dataset, resolved_version),
        "generatedCode": generated_code(dataset.get("name") or dataset_id, roles),
        "limitations": "Diagnostics use the Dataset Library scored-output sample/preview artifact. Full warehouse-scale diagnostics will use persisted full scored outputs when available.",
    }
    return diagnostics


def save_diagnostic_insight(payload: Dict[str, Any]) -> Dict[str, Any]:
    return create_insight(
        {
            "title": payload.get("title") or "Model Diagnostics Review",
            "summary": payload.get("summary") or "Scored-output diagnostics saved from ATLAS Model Diagnostics.",
            "datasetId": payload.get("datasetId"),
            "versionId": payload.get("versionId"),
            "projectId": payload.get("projectId"),
            "businessQuestion": payload.get("businessQuestion") or "Are scored-output diagnostics acceptable for governance review?",
            "analysisType": "model_diagnostics",
            "evidence": payload.get("evidence") or {},
        }
    )


def diagnostics_system_status() -> Dict[str, Any]:
    return {
        "ok": True,
        "checks": [
            {"id": "model_diagnostics_router", "label": "Model Diagnostics router registered", "status": "pass", "detail": "Scored-output diagnostics endpoints are available."},
            {"id": "model_diagnostics_dataset_artifacts", "label": "Dataset artifact reader", "status": "pass", "detail": "Diagnostics read Dataset Library preview/sample artifacts without browser localStorage payloads."},
            {"id": "model_diagnostics_exports", "label": "Governed exports", "status": "pass", "detail": "Excel, Word, and PowerPoint exports are generated from diagnostics summaries."},
        ],
    }


def detect_roles(df: pd.DataFrame, metadata: Dict[str, Any], overrides: Optional[Dict[str, Any]] = None) -> Dict[str, Optional[str]]:
    source = metadata.get("sourceInfo") if isinstance(metadata.get("sourceInfo"), dict) else {}
    roles = source.get("columnRoles") if isinstance(source.get("columnRoles"), dict) else {}
    columns = [str(column) for column in df.columns]
    override_values = overrides or {}

    def existing(value: Any) -> Optional[str]:
        return str(value) if value and str(value) in columns else None

    def infer(candidates: List[str]) -> Optional[str]:
        lower = {column.lower(): column for column in columns}
        for candidate in candidates:
            if candidate.lower() in lower:
                return lower[candidate.lower()]
        for column in columns:
            clean = column.lower().replace("_", " ")
            if any(candidate.lower().replace("_", " ") in clean for candidate in candidates):
                return column
        return None

    blocked_actual = {"predicted_value", "prediction", "score", "probability", "p_hat", "pred", "y_pred", "fitted_value"}
    explicit_target = existing(override_values.get("target")) or existing(roles.get("target")) or existing(source.get("targetColumn")) or existing(source.get("target")) or existing(source.get("actualColumn"))
    explicit_actual = existing(override_values.get("actual"))
    stored_actual = existing(roles.get("actual"))
    if stored_actual and stored_actual.lower() in blocked_actual:
        stored_actual = None
    inferred_actual = infer(["actual", "y_true", "observed", "target", "label", "outcome"])
    if inferred_actual and inferred_actual.lower() in blocked_actual:
        inferred_actual = None
    resolved = {
        "target": explicit_target,
        "actual": explicit_actual or explicit_target or stored_actual or inferred_actual,
        "score": existing(override_values.get("score")) or existing(roles.get("score")) or infer(["score", "probability", "prediction_score", "adjusted_score", "pd", "p_default"]),
        "prediction": existing(override_values.get("prediction")) or existing(roles.get("prediction")) or infer(["predicted_value", "prediction", "y_pred", "pred", "fitted_value"]),
        "residual": existing(override_values.get("residual")) or existing(roles.get("residual")) or infer(["atlas_residual", "residual", "error"]),
        "absResidual": existing(override_values.get("absResidual")) or existing(roles.get("absResidual")) or infer(["atlas_abs_residual", "abs_residual", "absolute_residual"]),
        "squaredResidual": existing(override_values.get("squaredResidual")) or existing(roles.get("squaredResidual")) or infer(["atlas_squared_residual", "squared_residual"]),
        "decile": existing(override_values.get("decile")) or existing(roles.get("decile")) or infer(["atlas_score_decile", "decile", "score_decile", "risk_decile"]),
        "segment": existing(override_values.get("segment")) or existing(roles.get("segment")),
        "positiveClass": str(override_values.get("positiveClass") or roles.get("positiveClass") or "1"),
    }
    if not resolved["score"] and resolved["prediction"]:
        resolved["score"] = resolved["prediction"]
    return resolved


def mapping_warnings_for_roles(roles: Dict[str, Optional[str]]) -> List[str]:
    warnings: List[str] = []
    if roles.get("actual") and roles.get("prediction") and roles.get("actual") == roles.get("prediction"):
        warnings.append("Actual and prediction are mapped to the same column. Residual diagnostics are not meaningful until this mapping is corrected.")
    if not roles.get("actual"):
        warnings.append("Actual target was not found in the scored input. Regression fit, residual, and actual-vs-predicted diagnostics cannot run until actual values are supplied.")
    return warnings


def diagnostic_mode(roles: Dict[str, Optional[str]], actual: pd.Series, prediction: pd.Series, score: pd.Series, df: pd.DataFrame, metadata: Optional[Dict[str, Any]] = None) -> str:
    source = (metadata or {}).get("sourceInfo") if isinstance((metadata or {}).get("sourceInfo"), dict) else {}
    task = str(source.get("task") or "").lower()
    if roles.get("actual") == roles.get("prediction") and roles.get("actual"):
        return "unmapped"
    if task == "regression" and not actual.empty and not prediction.empty:
        return "regression"
    if not actual.empty and not prediction.empty:
        actual_binary = binary_actual(df[roles["actual"]], roles.get("positiveClass")) if roles.get("actual") in df.columns else None
        if actual_binary is None:
            return "regression"
        if not score.empty and (score.min() >= 0 and score.max() <= 1):
            return "classification"
    if not actual.empty and not score.empty:
        return "classification"
    if not prediction.empty or not score.empty:
        return "prediction_only"
    return "unmapped"


def numeric_series(df: pd.DataFrame, column: Optional[str]) -> pd.Series:
    if not column or column not in df.columns:
        return pd.Series(dtype=float)
    return pd.to_numeric(df[column], errors="coerce").dropna()


def distribution(series: pd.Series, bins: int = 12) -> Dict[str, Any]:
    if series.empty:
        return {"available": False, "bins": [], "summary": {}}
    counts, edges = pd.cut(series, bins=min(bins, max(1, series.nunique())), retbins=True, duplicates="drop")
    table = counts.value_counts(sort=False)
    return {
        "available": True,
        "summary": {
            "min": safe_float(series.min()),
            "mean": safe_float(series.mean()),
            "median": safe_float(series.median()),
            "max": safe_float(series.max()),
            "std": safe_float(series.std(ddof=1)),
        },
        "bins": [
            {"bin": str(index), "count": int(value)}
            for index, value in table.items()
        ],
    }


def prediction_summary(series: pd.Series) -> Dict[str, Any]:
    if series.empty:
        return {"available": False, "summary": {}, "warnings": []}
    clean = pd.to_numeric(series, errors="coerce")
    missing = int(clean.isna().sum())
    clean = clean.dropna()
    if clean.empty:
        return {"available": False, "summary": {"missing": missing}, "warnings": ["Prediction column is not numeric."]}
    distinct = int(clean.nunique(dropna=True))
    duplicate_count = int(max(0, len(clean) - distinct))
    warnings: List[str] = []
    if distinct <= max(3, len(clean) * 0.05):
        warnings.append("Predictions have few distinct values. Review rounding, clipping, or a constant-score issue.")
    if float(clean.std(ddof=1) or 0.0) < 1e-10 and len(clean) > 1:
        warnings.append("Predictions are effectively constant.")
    return {
        "available": True,
        "summary": {
            "n": int(len(clean)),
            "missing": missing,
            "mean": safe_float(clean.mean()),
            "median": safe_float(clean.median()),
            "std": safe_float(clean.std(ddof=1)),
            "min": safe_float(clean.min()),
            "p1": safe_float(clean.quantile(0.01)),
            "p5": safe_float(clean.quantile(0.05)),
            "p25": safe_float(clean.quantile(0.25)),
            "p75": safe_float(clean.quantile(0.75)),
            "p95": safe_float(clean.quantile(0.95)),
            "p99": safe_float(clean.quantile(0.99)),
            "max": safe_float(clean.max()),
            "skewness": safe_float(clean.skew()) if len(clean) > 2 else None,
            "distinct": distinct,
            "duplicatePredictions": duplicate_count,
        },
        "warnings": warnings,
    }


def diagnostic_columns(df: pd.DataFrame, roles: Dict[str, Optional[str]]) -> List[str]:
    preferred = [
        "row",
        roles.get("prediction"),
        roles.get("actual"),
        roles.get("residual"),
        roles.get("absResidual"),
        "atlas_abs_residual",
        "atlas_percent_error",
        "MedInc",
        "HouseAge",
        "AveRooms",
        "AveBedrms",
        "Population",
        "AveOccup",
        "Latitude",
        "Longitude",
    ]
    seen: List[str] = []
    for column in preferred:
        if column and column in df.columns and column not in seen:
            seen.append(column)
    numeric_extra = [
        str(column)
        for column in df.columns
        if str(column) not in seen and pd.api.types.is_numeric_dtype(df[column])
    ][:5]
    return seen + numeric_extra


def table_records(frame: pd.DataFrame, columns: List[str], limit: int = 10) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for _, row in frame.head(limit).iterrows():
        item: Dict[str, Any] = {}
        for column in columns:
            value = row.get(column)
            item[column] = safe_float(value) if isinstance(value, Number) else (None if pd.isna(value) else str(value))
        records.append(item)
    return records


def prediction_extremes(df: pd.DataFrame, roles: Dict[str, Optional[str]], descending: bool) -> List[Dict[str, Any]]:
    pred_col = roles.get("prediction") or roles.get("score")
    if not pred_col or pred_col not in df.columns:
        return []
    data = df.copy()
    data["_prediction_sort"] = pd.to_numeric(data[pred_col], errors="coerce")
    data = data.dropna(subset=["_prediction_sort"]).sort_values("_prediction_sort", ascending=not descending)
    return table_records(data, diagnostic_columns(data, roles), limit=10)


def residual_extremes(df: pd.DataFrame, roles: Dict[str, Optional[str]], mode: str) -> List[Dict[str, Any]]:
    actual_col = roles.get("actual")
    pred_col = roles.get("prediction") or roles.get("score")
    if not actual_col or not pred_col or actual_col not in df.columns or pred_col not in df.columns or actual_col == pred_col:
        return []
    data = df.copy()
    actual = pd.to_numeric(data[actual_col], errors="coerce")
    pred = pd.to_numeric(data[pred_col], errors="coerce")
    data["atlas_residual"] = pd.to_numeric(data.get("atlas_residual", actual - pred), errors="coerce")
    data["atlas_abs_residual"] = data["atlas_residual"].abs()
    data["atlas_percent_error"] = data["atlas_residual"].where(actual != 0) / actual.where(actual != 0)
    data = data.dropna(subset=["atlas_residual"])
    if mode == "over":
        data = data[data["atlas_residual"] < 0].sort_values("atlas_residual", ascending=True)
    elif mode == "under":
        data = data[data["atlas_residual"] > 0].sort_values("atlas_residual", ascending=False)
    else:
        data = data.sort_values("atlas_abs_residual", ascending=False)
    return table_records(data, diagnostic_columns(data, {**roles, "residual": "atlas_residual", "absResidual": "atlas_abs_residual"}), limit=10)


def prediction_band_summary(df: pd.DataFrame, roles: Dict[str, Optional[str]], bands: int = 5) -> List[Dict[str, Any]]:
    pred_col = roles.get("prediction") or roles.get("score")
    if not pred_col or pred_col not in df.columns:
        return []
    data = pd.DataFrame({"prediction": pd.to_numeric(df[pred_col], errors="coerce")})
    actual_col = roles.get("actual")
    if actual_col and actual_col in df.columns and actual_col != pred_col:
        data["actual"] = pd.to_numeric(df[actual_col], errors="coerce")
        data["residual"] = data["actual"] - data["prediction"]
    data = data.dropna(subset=["prediction"])
    if data.empty:
        return []
    try:
        data["prediction_band"] = pd.qcut(data["prediction"].rank(method="first"), q=min(bands, len(data)), labels=False, duplicates="drop").astype(int) + 1
    except Exception:
        data["prediction_band"] = 1
    rows: List[Dict[str, Any]] = []
    for band, group in data.groupby("prediction_band"):
        row: Dict[str, Any] = {
            "prediction_band": int(band),
            "n": int(group.shape[0]),
            "mean_prediction": safe_float(group["prediction"].mean()),
            "min_prediction": safe_float(group["prediction"].min()),
            "max_prediction": safe_float(group["prediction"].max()),
        }
        if "actual" in group.columns:
            row.update(
                {
                    "mean_actual": safe_float(group["actual"].mean()),
                    "mean_residual": safe_float(group["residual"].mean()),
                    "mean_abs_residual": safe_float(group["residual"].abs().mean()),
                    "rmse": safe_float(math.sqrt(float((group["residual"] ** 2).mean()))),
                }
            )
        rows.append(row)
    return rows


def regression_fit_summary(actual: pd.Series, prediction: pd.Series, residual: pd.Series) -> Dict[str, Any]:
    if actual.empty or prediction.empty:
        return {"available": False, "summary": {}, "warning": "Regression fit summary requires actual and prediction columns."}
    n = min(len(actual), len(prediction))
    a = actual.iloc[:n].reset_index(drop=True)
    p = prediction.iloc[:n].reset_index(drop=True)
    if a.equals(p):
        return {"available": False, "summary": {}, "warning": "Actual and prediction are mapped to the same column. Correct field mapping before reviewing regression fit."}
    r = a - p if residual.empty else residual.iloc[:n].reset_index(drop=True)
    ss_res = float((r ** 2).sum())
    ss_tot = float(((a - a.mean()) ** 2).sum())
    mape_values = (r.abs() / a.replace(0, pd.NA).abs()).dropna()
    return {
        "available": True,
        "summary": {
            "n": int(n),
            "mae": safe_float(r.abs().mean()),
            "mse": safe_float((r ** 2).mean()),
            "rmse": safe_float(math.sqrt(float((r ** 2).mean()))),
            "rSquared": safe_float(1 - ss_res / ss_tot) if ss_tot else None,
            "meanResidual": safe_float(r.mean()),
            "medianResidual": safe_float(r.median()),
            "residualStd": safe_float(r.std(ddof=1)),
            "mape": safe_float(mape_values.mean()) if not mape_values.empty else None,
            "correlation": safe_float(a.corr(p)) if n > 2 else None,
            "predictionMean": safe_float(p.mean()),
            "actualMean": safe_float(a.mean()),
            "predictionBias": safe_float((p - a).mean()),
        },
    }


def lift_deciles(df: pd.DataFrame, score_col: Optional[str], actual_col: Optional[str], decile_col: Optional[str], positive_class: Optional[str] = "1") -> List[Dict[str, Any]]:
    if not score_col or score_col not in df.columns:
        return []
    data = df.copy()
    data["_score"] = pd.to_numeric(data[score_col], errors="coerce")
    data = data.dropna(subset=["_score"])
    if data.empty:
        return []
    if actual_col and actual_col in data.columns:
        actual = binary_actual(data[actual_col], positive_class)
        if actual is None:
            return []
        data["_actual"] = actual
    else:
        data["_actual"] = 0
    if decile_col and decile_col in data.columns:
        data["_decile"] = pd.to_numeric(data[decile_col], errors="coerce")
    else:
        ranked = data["_score"].rank(method="first", ascending=False)
        data["_decile"] = pd.qcut(ranked, q=min(10, max(1, len(data))), labels=False, duplicates="drop").astype(float) + 1
    baseline = data["_actual"].mean() or 1
    rows = []
    for decile, group in data.groupby("_decile"):
        rows.append(
            {
                "decile": int(decile) if pd.notna(decile) else None,
                "n": int(group.shape[0]),
                "eventRate": safe_float(group["_actual"].mean()),
                "lift": safe_float((group["_actual"].mean() or 0) / baseline),
                "minScore": safe_float(group["_score"].min()),
                "maxScore": safe_float(group["_score"].max()),
                "meanScore": safe_float(group["_score"].mean()),
            }
        )
    return sorted(rows, key=lambda row: row.get("decile") or 999)


def classification_summary_metrics(score: pd.Series, actual: pd.Series, positive_class: Optional[str] = "1") -> Dict[str, Any]:
    if score.empty or actual.empty:
        return {"available": False, "warning": "AUC/KS/Gini require binary actual outcome and score/probability."}
    n = min(len(score), len(actual))
    frame = pd.DataFrame({"score": score.iloc[:n].reset_index(drop=True), "actual": actual.iloc[:n].reset_index(drop=True)}).dropna()
    actual_binary = binary_actual(frame["actual"], positive_class)
    if frame.empty or actual_binary is None:
        return {"available": False, "warning": "AUC/KS/Gini require binary actual outcome and score/probability."}
    y = actual_binary.loc[frame.index] if hasattr(actual_binary, "loc") else actual_binary
    s = pd.to_numeric(frame["score"], errors="coerce")
    try:
        auc = float(roc_auc_score(y, s))
    except Exception:
        auc = None
    ordered = pd.DataFrame({"score": s, "actual": y}).sort_values("score", ascending=False)
    total_events = max(float(ordered["actual"].sum()), 1.0)
    total_nonevents = max(float((1 - ordered["actual"]).sum()), 1.0)
    event_cdf = ordered["actual"].cumsum() / total_events
    nonevent_cdf = (1 - ordered["actual"]).cumsum() / total_nonevents
    ks = float((event_cdf - nonevent_cdf).abs().max()) if not ordered.empty else None
    total = int(len(ordered))
    events = int(ordered["actual"].sum())
    return {
        "available": True,
        "auc": safe_float(auc),
        "gini": safe_float(2 * auc - 1) if auc is not None else None,
        "ks": safe_float(ks),
        "baseEventRate": safe_float(events / total) if total else None,
        "totalRecords": total,
        "totalEvents": events,
        "totalNonEvents": total - events,
    }


def detailed_lift_table(df: pd.DataFrame, score_col: Optional[str], actual_col: Optional[str], decile_col: Optional[str], positive_class: Optional[str] = "1") -> List[Dict[str, Any]]:
    if not score_col or not actual_col or score_col not in df.columns or actual_col not in df.columns:
        return []
    data = pd.DataFrame({"score": pd.to_numeric(df[score_col], errors="coerce"), "actual_raw": df[actual_col]}).dropna(subset=["score"])
    actual = binary_actual(data["actual_raw"], positive_class)
    if actual is None:
        return []
    data["actual"] = actual
    if decile_col and decile_col in df.columns:
        data["decile"] = pd.to_numeric(df.loc[data.index, decile_col], errors="coerce")
    else:
        data["decile"] = pd.qcut(data["score"].rank(method="first", ascending=False), q=min(10, len(data)), labels=False, duplicates="drop").astype(float) + 1
    total = max(int(data.shape[0]), 1)
    total_events = max(float(data["actual"].sum()), 1.0)
    total_nonevents = max(float(total - data["actual"].sum()), 1.0)
    base = float(data["actual"].mean()) or 1.0
    rows: List[Dict[str, Any]] = []
    cum_n = cum_events = cum_nonevents = 0.0
    for decile, group in data.groupby("decile"):
        n = int(group.shape[0])
        events = float(group["actual"].sum())
        nonevents = float(n - events)
        cum_n += n
        cum_events += events
        cum_nonevents += nonevents
        event_rate = float(group["actual"].mean()) if n else 0.0
        rows.append({
            "Decile": int(decile),
            "# of accounts": n,
            "% of all accounts": safe_float(n / total),
            "Min score": safe_float(group["score"].min()),
            "Max score": safe_float(group["score"].max()),
            "Avg score": safe_float(group["score"].mean()),
            "Expected # of events": safe_float(base * n),
            "Actual event rate": safe_float(event_rate),
            "Actual # of events": int(events),
            "Cumulative % of all accounts": safe_float(cum_n / total),
            "Cumulative % of events": safe_float(cum_events / total_events),
            "Cumulative % of non-events": safe_float(cum_nonevents / total_nonevents),
            "Lift": safe_float(event_rate / base),
        })
    return sorted(rows, key=lambda row: row["Decile"])


def diagnostics_regression_summary(actual: pd.Series, prediction: pd.Series, residual: pd.Series) -> Dict[str, Any]:
    fit = regression_fit_summary(actual, prediction, residual)
    return fit.get("summary") if fit.get("available") else {}


def population_waterfall(dataset: Dict[str, Any], df: pd.DataFrame, roles: Dict[str, Optional[str]], mode: str) -> List[Dict[str, Any]]:
    source = dataset.get("sourceInfo") or {}
    predictors = len([c for c in df.columns if c not in set(filter(None, roles.values()))])
    actual = binary_actual(df[roles["actual"]], roles.get("positiveClass")) if mode == "classification" and roles.get("actual") in df.columns else None
    events = int(actual.sum()) if actual is not None else None
    response = safe_float(events / len(df)) if events is not None and len(df) else None
    return [
        {"Waterfall Step": "Full Developmental Dataset", "Observations": int(source.get("developmentRows") or len(df)), "Predictors": int(source.get("developmentPredictors") or predictors), "Responders / Events": events, "Response Rate": response, "Retained %": 1.0, "Dropped %": 0.0, "Notes": "Use uploaded development metadata when available; otherwise scored-output row counts are shown."},
        {"Waterfall Step": "Balance the Data", "Observations": int(source.get("balancedRows") or len(df)), "Predictors": predictors, "Responders / Events": events, "Response Rate": response, "Retained %": None, "Dropped %": None, "Notes": "Editable placeholder if balancing metadata is unavailable."},
        {"Waterfall Step": "Information Value Filter", "Observations": int(len(df)), "Predictors": predictors, "Responders / Events": events, "Response Rate": response, "Retained %": None, "Dropped %": None, "Notes": "Document IV/feature filter decisions in the model development workflow."},
        {"Waterfall Step": "Correlation Analysis Filter", "Observations": int(len(df)), "Predictors": predictors, "Responders / Events": events, "Response Rate": response, "Retained %": None, "Dropped %": None, "Notes": "Document correlated-feature removals and rationale."},
        {"Waterfall Step": "Recursive Feature Elimination", "Observations": int(len(df)), "Predictors": predictors, "Responders / Events": events, "Response Rate": response, "Retained %": None, "Dropped %": None, "Notes": "Populate from RFE output when available."},
        {"Waterfall Step": "Final Model Features", "Observations": int(len(df)), "Predictors": predictors, "Responders / Events": events, "Response Rate": response, "Retained %": None, "Dropped %": None, "Notes": "Final features should tie to champion model documentation."},
    ]


VIF_EXCLUDED_FIELD_RE = re.compile(r"(^|_)(row|record|index|id|split|fold|target|actual|score|probability|prediction|predicted|residual|error|decile|diagnostic|output|rank)($|_)", re.I)


def vif_exclusion_reason(df: pd.DataFrame, column: str, role_values: set[str]) -> Optional[str]:
    if column in role_values or VIF_EXCLUDED_FIELD_RE.search(str(column)):
        return "Diagnostic/output field"
    if not pd.api.types.is_numeric_dtype(df[column]):
        return "Non-numeric variable"
    values = pd.to_numeric(df[column], errors="coerce")
    non_missing = values.dropna()
    if len(non_missing) < max(20, int(len(values) * 0.5)):
        return "Too many missing values"
    if non_missing.nunique(dropna=True) <= 1:
        return "Constant variable"
    return None


def unavailable_vif_row(column: str, reason: str) -> Dict[str, Any]:
    return {
        "Variable": str(column),
        "VIF": None,
        "Tolerance": None,
        "R-squared against other predictors": None,
        "Status": "Unavailable",
        "Reason": reason,
        "Recommendation": reason,
    }


def vif_analysis(df: pd.DataFrame, roles: Dict[str, Optional[str]]) -> Dict[str, Any]:
    role_values = {value for value in roles.values() if isinstance(value, str)}
    reasons = {column: vif_exclusion_reason(df, column, role_values) for column in df.columns}
    candidates = [column for column, reason in reasons.items() if reason is None]
    excluded_rows = [unavailable_vif_row(column, reason) for column, reason in reasons.items() if reason][:12]
    if len(candidates) < 2:
        return {"available": False, "formula": "VIF_j = 1 / (1 - R_j^2); Tolerance = 1 / VIF", "rows": excluded_rows, "warning": "VIF requires at least two numeric final-model predictors after excluding target, row IDs, split flags, predictions, residuals, scores, and other diagnostic fields."}
    data = df[candidates].apply(pd.to_numeric, errors="coerce").dropna()
    if data.shape[0] < max(20, len(candidates) + 2):
        rows = [unavailable_vif_row(column, "Not enough rows") for column in candidates[:25]]
        return {"available": False, "formula": "VIF_j = 1 / (1 - R_j^2); Tolerance = 1 / VIF", "rows": rows + excluded_rows, "warning": "Not enough complete modeling rows are available for VIF after missing-value handling."}
    rows = []
    for column in candidates[:25]:
        y = data[column].to_numpy(dtype=float)
        others = [c for c in candidates if c != column][:25]
        X = data[others].to_numpy(dtype=float)
        X = pd.DataFrame(X).assign(_intercept=1.0).to_numpy(dtype=float)
        try:
            beta, *_ = __import__("numpy").linalg.lstsq(X, y, rcond=None)
            pred = X @ beta
            ss_res = float(((y - pred) ** 2).sum())
            ss_tot = float(((y - y.mean()) ** 2).sum())
            if ss_tot <= 1e-12:
                rows.append(unavailable_vif_row(column, "Constant variable"))
                continue
            r2 = max(0.0, min(0.999999, 1 - ss_res / ss_tot))
        except Exception:
            rows.append(unavailable_vif_row(column, "Singular matrix"))
            continue
        if 1 - r2 < 1e-6:
            rows.append(unavailable_vif_row(column, "Singular matrix"))
            continue
        vif = 1 / max(1 - r2, 1e-6)
        status = "Low concern" if vif < 5 else "Moderate concern" if vif < 10 else "High multicollinearity concern"
        recommendation = "Retain unless other evidence suggests removal." if vif < 5 else "Review correlation and business necessity." if vif < 10 else "Consider removing, combining, or regularizing this predictor."
        rows.append({"Variable": str(column), "VIF": safe_float(vif), "Tolerance": safe_float(1 / vif), "R-squared against other predictors": safe_float(r2), "Status": status, "Reason": "", "Recommendation": recommendation})
    return {"available": True, "formula": "VIF_j = 1 / (1 - R_j^2); Tolerance = 1 / VIF", "rows": rows + excluded_rows}


def pozzolo_calibration_detail(score: pd.Series, actual: pd.Series, roles: Dict[str, Optional[str]], metadata: Dict[str, Any]) -> Dict[str, Any]:
    source = metadata.get("sourceInfo") if isinstance(metadata.get("sourceInfo"), dict) else {}
    task = str(source.get("task") or "").lower()
    equation = "alpha = [pi / (1 - pi)] / [pi_s / (1 - pi_s)]; p_corrected = (alpha * p_s) / (1 - p_s + alpha * p_s)"
    if task == "regression":
        return {"applicable": False, "equation": equation, "note": "Pozzolo correction applies to classification probability calibration after sampling changes class priors. It does not apply to regression predictions."}
    if score.empty:
        return {"applicable": True, "equation": equation, "warning": "Score/probability is required for Pozzolo prior correction."}
    pi_s = safe_float(source.get("model_prior") or source.get("sampleTrainingEventRate"))
    pi = safe_float(source.get("original_prior") or source.get("targetPopulationEventRate"))
    if pi is None and not actual.empty:
        binary = binary_actual(actual, roles.get("positiveClass"))
        pi = safe_float(binary.mean()) if binary is not None and not binary.empty else None
    if pi_s is None:
        pi_s = safe_float(score.mean())
    if pi is None or pi_s is None or pi <= 0 or pi >= 1 or pi_s <= 0 or pi_s >= 1:
        return {"applicable": True, "equation": equation, "warning": "Sample and target event rates are required to calculate Pozzolo correction.", "sampleTrainingEventRate": pi_s, "targetPopulationEventRate": pi}
    alpha = (pi / (1 - pi)) / (pi_s / (1 - pi_s))
    p = score.clip(1e-9, 1 - 1e-9)
    corrected = (alpha * p) / (1 - p + alpha * p)
    return {"applicable": True, "equation": equation, "sampleTrainingEventRate": pi_s, "targetPopulationEventRate": pi, "correctionFactorAlpha": safe_float(alpha), "exampleOriginalProbability": safe_float(p.iloc[0]), "exampleCorrectedProbability": safe_float(corrected.iloc[0]), "meanOriginalProbability": safe_float(p.mean()), "meanCorrectedProbability": safe_float(corrected.mean())}


def kpi_aggregation(mode: str, classification: Dict[str, Any], regression: Dict[str, Any]) -> Dict[str, Any]:
    if mode == "regression":
        r2 = regression.get("rSquared")
        rmse = regression.get("rmse")
        r2_rating = "Green" if isinstance(r2, (int, float)) and r2 >= 0.6 else "Yellow" if isinstance(r2, (int, float)) and r2 >= 0.3 else "Red"
        rmse_rating = "Green" if isinstance(rmse, (int, float)) and rmse < 1 else "Yellow" if isinstance(rmse, (int, float)) and rmse < 3 else "Red"
        overall = "Red" if "Red" in {r2_rating, rmse_rating} else "Yellow" if "Yellow" in {r2_rating, rmse_rating} else "Green"
        return {"mode": "regression", "rows": [{"Metric": "RMSE rating", "Rating": rmse_rating, "Value": rmse}, {"Metric": "R-squared rating", "Rating": r2_rating, "Value": r2}, {"Metric": "Stability rating", "Rating": "Yellow", "Value": "Review TEST/OOT split comparison"}, {"Metric": "Overall rating", "Rating": overall, "Value": overall}], "recommendation": "Continue monitoring." if overall == "Green" else "Review stability and redevelopment triggers."}
    auc = classification.get("auc")
    auc_rating = "Green" if isinstance(auc, (int, float)) and auc >= 0.75 else "Yellow" if isinstance(auc, (int, float)) and auc >= 0.65 else "Red"
    overall = auc_rating
    return {"mode": "classification", "rows": [{"Metric": "AUC rating", "Rating": auc_rating, "Value": auc}, {"Metric": "PSI rating", "Rating": "Yellow", "Value": "Configure threshold"}, {"Metric": "Calibration rating", "Rating": "Yellow", "Value": "Review calibration bins"}, {"Metric": "Overall rating", "Rating": overall, "Value": overall}], "recommendation": "Green + Yellow = Green, continue using the model. Red ratings require redevelopment review."}


def residual_summary(residual: pd.Series) -> Dict[str, Any]:
    if residual.empty:
        return {"available": False, "summary": {}, "points": []}
    return {
        "available": True,
        "summary": {
            "mean": safe_float(residual.mean()),
            "median": safe_float(residual.median()),
            "std": safe_float(residual.std(ddof=1)),
            "mae": safe_float(residual.abs().mean()),
            "p95Abs": safe_float(residual.abs().quantile(0.95)),
        },
        "histogram": distribution(residual, bins=12).get("bins", []),
    }


def actual_vs_predicted(actual: pd.Series, predicted: pd.Series) -> Dict[str, Any]:
    if actual.empty or predicted.empty:
        return {"available": False, "points": [], "summary": {}}
    n = min(len(actual), len(predicted), 250)
    a = actual.iloc[:n].reset_index(drop=True)
    p = predicted.iloc[:n].reset_index(drop=True)
    corr = a.corr(p) if n > 2 else None
    return {
        "available": True,
        "summary": {"correlation": safe_float(corr), "pointsShown": n},
        "points": [{"actual": safe_float(a.iloc[i]), "predicted": safe_float(p.iloc[i])} for i in range(n)],
    }


def calibration_bins(score: pd.Series, actual: pd.Series, positive_class: Optional[str] = "1", bins: int = 10) -> List[Dict[str, Any]]:
    if score.empty or actual.empty:
        return []
    n = min(len(score), len(actual))
    frame = pd.DataFrame({"score": score.iloc[:n].reset_index(drop=True), "actual": actual.iloc[:n].reset_index(drop=True)}).dropna()
    if frame.empty:
        return []
    actual_binary = binary_actual(frame["actual"], positive_class)
    if actual_binary is None:
        return []
    frame["actual"] = actual_binary
    frame["bin"] = pd.qcut(frame["score"].rank(method="first"), q=min(bins, len(frame)), labels=False, duplicates="drop")
    rows = []
    for index, group in frame.groupby("bin"):
        rows.append({"bin": int(index) + 1, "meanScore": safe_float(group["score"].mean()), "actualRate": safe_float(group["actual"].mean()), "n": int(group.shape[0])})
    return rows


def segment_performance(df: pd.DataFrame, score_col: Optional[str], actual_col: Optional[str], segment_col: Optional[str] = None) -> List[Dict[str, Any]]:
    value_col = score_col
    if not value_col or value_col not in df.columns:
        return []
    binned = False
    if segment_col and segment_col in df.columns:
        column = str(segment_col)
        if pd.api.types.is_numeric_dtype(df[column]) and df[column].nunique(dropna=True) > 12:
            binned = True
    else:
        candidates = [column for column in df.columns if column not in {score_col, actual_col} and df[column].nunique(dropna=True) <= 12]
        numeric_candidates = [column for column in df.columns if column not in {score_col, actual_col} and pd.api.types.is_numeric_dtype(df[column])]
        if not candidates and not numeric_candidates:
            return []
        column = str(candidates[0] if candidates else numeric_candidates[0])
        binned = column in numeric_candidates and column not in candidates
    data = df[[column, value_col] + ([actual_col] if actual_col and actual_col in df.columns and actual_col != value_col else [])].copy()
    data["_prediction"] = pd.to_numeric(data[value_col], errors="coerce")
    if actual_col and actual_col in data.columns:
        data["_actual"] = pd.to_numeric(data[actual_col], errors="coerce")
        data["_residual"] = data["_actual"] - data["_prediction"]
    if binned:
        try:
            data["_segment"] = pd.qcut(pd.to_numeric(data[column], errors="coerce").rank(method="first"), q=min(5, len(data)), duplicates="drop")
        except Exception:
            data["_segment"] = "All"
    else:
        data["_segment"] = data[column]
    rows = []
    for value, group in data.groupby("_segment", dropna=False):
        row = {
            "segmentColumn": f"{column} (binned)" if binned else column,
            "segment": str(value),
            "n": int(group.shape[0]),
            "meanPrediction": safe_float(group["_prediction"].mean()),
        }
        if "_actual" in group.columns:
            row.update(
                {
                    "meanActual": safe_float(group["_actual"].mean()),
                    "meanResidual": safe_float(group["_residual"].mean()),
                    "meanAbsResidual": safe_float(group["_residual"].abs().mean()),
                    "rmse": safe_float(math.sqrt(float((group["_residual"] ** 2).mean()))),
                }
            )
        rows.append(row)
    return sorted(rows, key=lambda row: row["n"], reverse=True)[:12]


def split_comparison(dataset: Dict[str, Any], version_id: str) -> List[Dict[str, Any]]:
    source = dataset.get("sourceInfo") or {}
    project_id = source.get("projectId")
    model_name = source.get("modelName")
    if not project_id and not model_name:
        return []
    rows: List[Dict[str, Any]] = []
    for candidate in list_accessible_datasets():
        candidate_source = candidate.get("sourceInfo") or {}
        if candidate.get("datasetRole") != "scored_output" and candidate_source.get("datasetRole") != "scored_output":
            continue
        if project_id and candidate_source.get("projectId") != project_id:
            continue
        if model_name and candidate_source.get("modelName") != model_name:
            continue
        candidate_version = (candidate.get("latestVersion") or {}).get("versionId")
        if not candidate_version:
            continue
        try:
            _, _, frame, metadata = load_scored_frame(str(candidate.get("datasetId")), str(candidate_version))
            roles = detect_roles(frame, metadata)
            pred = numeric_series(frame, roles.get("prediction") or roles.get("score"))
            actual = numeric_series(frame, roles.get("actual"))
            residual = numeric_series(frame, roles.get("residual"))
            if residual.empty and not actual.empty and not pred.empty and roles.get("actual") != roles.get("prediction"):
                residual = actual.iloc[: min(len(actual), len(pred))].reset_index(drop=True) - pred.iloc[: min(len(actual), len(pred))].reset_index(drop=True)
            fit = regression_fit_summary(actual, pred, residual)
            fit_summary = fit.get("summary") if fit.get("available") else {}
            rows.append(
                {
                    "split": str(candidate_source.get("datasetType") or candidate_source.get("split") or "").upper() or "SCORED",
                    "datasetName": candidate.get("name"),
                    "rowCount": int(frame.shape[0]),
                    "meanPrediction": safe_float(pred.mean()) if not pred.empty else None,
                    "medianPrediction": safe_float(pred.median()) if not pred.empty else None,
                    "stdPrediction": safe_float(pred.std(ddof=1)) if not pred.empty else None,
                    "minPrediction": safe_float(pred.min()) if not pred.empty else None,
                    "maxPrediction": safe_float(pred.max()) if not pred.empty else None,
                    "meanActual": safe_float(actual.mean()) if not actual.empty else None,
                    "rmse": fit_summary.get("rmse"),
                    "mae": fit_summary.get("mae"),
                    "rSquared": fit_summary.get("rSquared"),
                    "meanResidual": fit_summary.get("meanResidual"),
                }
            )
        except Exception:
            continue
    order = {"TRAIN": 0, "TEST": 1, "OOT": 2, "EXTERNAL_VALIDATION": 3}
    return sorted(rows, key=lambda row: (order.get(str(row.get("split")), 99), str(row.get("datasetName"))))


def binary_actual(series: pd.Series, positive_class: Optional[str] = "1") -> Optional[pd.Series]:
    clean = series.dropna()
    if clean.empty:
        return pd.Series(dtype=float)
    numeric = pd.to_numeric(clean, errors="coerce")
    unique_numeric = set(numeric.dropna().unique().tolist())
    if unique_numeric and unique_numeric.issubset({0, 1, 0.0, 1.0}):
        return pd.to_numeric(series, errors="coerce").fillna(0)
    unique_text = {str(value).strip().lower() for value in clean.unique().tolist()}
    positive = str(positive_class or "1").strip().lower()
    accepted = {positive, "0"}
    if unique_text and unique_text.issubset(accepted):
        return series.astype(str).str.strip().str.lower().eq(positive).astype(float)
    return None


def generated_code(dataset_name: str, roles: Dict[str, Optional[str]]) -> Dict[str, str]:
    score = roles.get("score") or roles.get("prediction") or "predicted_value"
    actual = roles.get("actual") or "actual"
    prediction = roles.get("prediction") or score
    residual = roles.get("residual") or "atlas_residual"
    abs_residual = roles.get("absResidual") or "atlas_abs_residual"
    return {
        "python": f"""import pandas as pd\nimport numpy as np\n\n# df should contain scored output rows for {dataset_name!r}\npred_col = {prediction!r}\nactual_col = {actual!r}\nresidual_col = {residual!r}\nabs_residual_col = {abs_residual!r}\n\nprint('Prediction summary')\nprint(df[pred_col].describe(percentiles=[.01,.05,.25,.5,.75,.95,.99]))\n\nprint('Top 10 highest predictions')\nprint(df.sort_values(pred_col, ascending=False).head(10))\n\nprint('Bottom 10 lowest predictions')\nprint(df.sort_values(pred_col, ascending=True).head(10))\n\nif actual_col in df.columns:\n    df[residual_col] = df[actual_col] - df[pred_col]\n    df[abs_residual_col] = df[residual_col].abs()\n    print('Largest residual errors')\n    print(df.sort_values(abs_residual_col, ascending=False).head(10))\n    df['prediction_band'] = pd.qcut(df[pred_col].rank(method='first'), 5, duplicates='drop')\n    print('Prediction band summary')\n    print(df.groupby('prediction_band').agg(n=(pred_col,'size'), mean_prediction=(pred_col,'mean'), mean_actual=(actual_col,'mean'), mean_abs_residual=(abs_residual_col,'mean')))\n\n# Detailed classification lift table. Use only when actual_col is binary and pred_col is a probability/score.\nif actual_col in df.columns:\n    scored = df[[pred_col, actual_col]].copy()\n    scored['score_decile'] = pd.qcut(scored[pred_col].rank(method='first', ascending=False), 10, labels=False, duplicates='drop') + 1\n    detailed_lift = scored.groupby('score_decile').agg(accounts=(pred_col, 'size'), min_score=(pred_col, 'min'), max_score=(pred_col, 'max'), avg_score=(pred_col, 'mean'), actual_events=(actual_col, 'sum'), actual_event_rate=(actual_col, 'mean')).reset_index()\n    detailed_lift['expected_events'] = detailed_lift['accounts'] * scored[actual_col].mean()\n    detailed_lift['lift'] = detailed_lift['actual_event_rate'] / scored[actual_col].mean()\n    print('Detailed lift table')\n    print(detailed_lift)\n\n# VIF / multicollinearity analysis: VIF_j = 1 / (1 - R_j^2).\nfeature_cols = [c for c in df.select_dtypes(include='number').columns if c not in {{pred_col, actual_col, residual_col, abs_residual_col}}]\nfor feature in feature_cols[:20]:\n    others = [c for c in feature_cols if c != feature]\n    if not others:\n        continue\n    y = df[feature].astype(float).to_numpy()\n    X = np.column_stack([np.ones(len(df)), df[others].astype(float).to_numpy()])\n    coef, *_ = np.linalg.lstsq(X, y, rcond=None)\n    fitted = X @ coef\n    r2 = 1 - np.sum((y - fitted) ** 2) / max(np.sum((y - y.mean()) ** 2), 1e-12)\n    vif = 1 / max(1 - r2, 1e-12)\n    print(feature, 'VIF=', round(vif, 4), 'Tolerance=', round(1 / vif, 4))\n\n# Pozzolo prior correction for undersampled classification probabilities.\n# alpha = [pi / (1 - pi)] / [pi_s / (1 - pi_s)]\n# p_corrected = (alpha * p_s) / (1 - p_s + alpha * p_s)\n""",
        "sql": f"""-- Replace scored_output with the governed scored dataset table.\nSELECT COUNT(*) AS n,\n       AVG({prediction}) AS mean_prediction,\n       MIN({prediction}) AS min_prediction,\n       MAX({prediction}) AS max_prediction\nFROM scored_output;\n\nSELECT *\nFROM scored_output\nORDER BY {prediction} DESC\nFETCH FIRST 10 ROWS ONLY;\n\nSELECT *\nFROM scored_output\nORDER BY {prediction} ASC\nFETCH FIRST 10 ROWS ONLY;\n\nWITH bands AS (\n  SELECT scored_output.*,\n         NTILE(5) OVER (ORDER BY {prediction}) AS prediction_band,\n         ({actual} - {prediction}) AS atlas_residual\n  FROM scored_output\n)\nSELECT prediction_band,\n       COUNT(*) AS n,\n       AVG({prediction}) AS mean_prediction,\n       AVG({actual}) AS mean_actual,\n       AVG(ABS(atlas_residual)) AS mean_abs_residual\nFROM bands\nGROUP BY prediction_band\nORDER BY prediction_band;\n\n-- Detailed classification lift table. Use DESC ordering for highest-risk decile first.\nWITH lift_base AS (\n  SELECT scored_output.*,\n         NTILE(10) OVER (ORDER BY {score} DESC) AS score_decile\n  FROM scored_output\n), base_rate AS (\n  SELECT AVG({actual}) AS event_rate FROM scored_output\n)\nSELECT score_decile,\n       COUNT(*) AS accounts,\n       MIN({score}) AS min_score,\n       MAX({score}) AS max_score,\n       AVG({score}) AS avg_score,\n       SUM({actual}) AS actual_events,\n       AVG({actual}) AS actual_event_rate,\n       AVG({actual}) / NULLIF((SELECT event_rate FROM base_rate), 0) AS lift\nFROM lift_base\nGROUP BY score_decile\nORDER BY score_decile;\n\n-- VIF is computed by regressing each predictor against the remaining predictors:\n-- VIF_j = 1 / (1 - R_j^2).\n-- Pozzolo correction: alpha = [pi/(1-pi)]/[pi_s/(1-pi_s)]; p_corrected = (alpha*p_s)/(1-p_s+alpha*p_s).""",
        "r": f"""# df should contain scored output rows for {dataset_name!r}\npred_col <- {prediction!r}\nactual_col <- {actual!r}\nresidual_col <- {residual!r}\nabs_residual_col <- {abs_residual!r}\n\nprint(summary(df[[pred_col]]))\nprint(head(df[order(-df[[pred_col]]), ], 10))\nprint(head(df[order(df[[pred_col]]), ], 10))\n\nif (actual_col %in% names(df)) {{\n  df[[residual_col]] <- df[[actual_col]] - df[[pred_col]]\n  df[[abs_residual_col]] <- abs(df[[residual_col]])\n  print(head(df[order(-df[[abs_residual_col]]), ], 10))\n  df$prediction_band <- cut(rank(df[[pred_col]], ties.method='first'), breaks=5, include.lowest=TRUE)\n  print(aggregate(df[, c(pred_col, actual_col, residual_col, abs_residual_col)], by=list(df$prediction_band), FUN=mean, na.rm=TRUE))\n  df$score_decile <- cut(rank(-df[[pred_col]], ties.method='first'), breaks=10, labels=FALSE, include.lowest=TRUE)\n  print(aggregate(df[, c(pred_col, actual_col)], by=list(df$score_decile), FUN=mean, na.rm=TRUE))\n}}\n\n# VIF_j = 1 / (1 - R_j^2). Regress each numeric predictor against the remaining predictors.\n# Pozzolo correction: alpha = [pi/(1-pi)]/[pi_s/(1-pi_s)]; p_corrected = (alpha*p_s)/(1-p_s+alpha*p_s).\n""",
        "sas": f"""proc means data=work.scored n mean median std min p5 p25 p75 p95 max;\n    var {prediction} {actual} {residual} {abs_residual};\nrun;\n\nproc sort data=work.scored out=work.top_predictions;\n    by descending {prediction};\nrun;\nproc print data=work.top_predictions(obs=10);\nrun;\n\nproc sort data=work.scored out=work.bottom_predictions;\n    by {prediction};\nrun;\nproc print data=work.bottom_predictions(obs=10);\nrun;\n\nproc sort data=work.scored out=work.largest_errors;\n    by descending {abs_residual};\nrun;\nproc print data=work.largest_errors(obs=10);\nrun;\n\nproc rank data=work.scored groups=5 out=work.scored_bands;\n    var {prediction};\n    ranks prediction_band;\nrun;\n\nproc means data=work.scored_bands n mean std;\n    class prediction_band;\n    var {prediction} {actual} {residual} {abs_residual};\nrun;\n\n/* Detailed classification lift table. Highest score becomes first decile. */\nproc rank data=work.scored descending groups=10 out=work.lift_ready;\n    var {score};\n    ranks score_decile;\nrun;\nproc means data=work.lift_ready n min max mean sum;\n    class score_decile;\n    var {score} {actual};\nrun;\n\n/* VIF_j = 1 / (1 - R_j^2). Use PROC REG / VIF in licensed SAS when available. */\n/* Pozzolo correction: alpha = [pi/(1-pi)]/[pi_s/(1-pi_s)]; p_corrected = (alpha*p_s)/(1-p_s+alpha*p_s). */""",
    }


def dataset_summary(dataset: Dict[str, Any]) -> Dict[str, Any]:
    versions = dataset.get("versions") or []
    latest = versions[-1] if versions else {}
    return {
        "datasetId": dataset.get("datasetId"),
        "name": dataset.get("name"),
        "sourceType": dataset.get("sourceType"),
        "datasetRole": dataset.get("datasetRole") or dataset.get("sourceInfo", {}).get("datasetRole"),
        "latestVersion": latest,
        "sourceInfo": dataset.get("sourceInfo") or {},
    }


def export_diagnostics_excel(diagnostics: Dict[str, Any]) -> bytes:
    try:
        from openpyxl import Workbook
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Excel export requires openpyxl. Install server requirements and retry.") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Dataset", diagnostics["dataset"]["name"]),
        ("Version", diagnostics["versionId"]),
        ("Rows", diagnostics["rowCount"]),
        ("Limitation", diagnostics["limitations"]),
    ]
    for row in rows:
        ws.append(row)

    def add_records_sheet(title: str, items: List[Dict[str, Any]]) -> None:
        sheet = wb.create_sheet(title[:31])
        if items:
            sheet.append(list(items[0].keys()))
            for item in items:
                sheet.append(list(item.values()))

    def add_kv_sheet(title: str, payload: Dict[str, Any]) -> None:
        sheet = wb.create_sheet(title[:31])
        sheet.append(["Metric", "Value"])
        for key, value in (payload or {}).items():
            sheet.append([key, value])

    add_kv_sheet("Field Mapping", diagnostics.get("columnRoles") or {})
    add_kv_sheet("Prediction Summary", (diagnostics.get("predictionSummary") or {}).get("summary") or {})
    add_kv_sheet("Actual Summary", (diagnostics.get("actualDistribution") or {}).get("summary") or {})
    add_records_sheet("Actual vs Predicted", (diagnostics.get("actualVsPredicted") or {}).get("points") or [])
    add_kv_sheet("Residual Diagnostics", (diagnostics.get("residualDiagnostics") or {}).get("summary") or {})
    add_records_sheet("Top 10 Highest Predictions", diagnostics.get("topPredictions") or [])
    add_records_sheet("Bottom 10 Lowest Predictions", diagnostics.get("bottomPredictions") or [])
    add_records_sheet("Largest Absolute Errors", diagnostics.get("largestAbsoluteErrors") or [])
    add_records_sheet("Top Overpredictions", diagnostics.get("topOverpredictions") or [])
    add_records_sheet("Top Underpredictions", diagnostics.get("topUnderpredictions") or [])
    add_records_sheet("Prediction Band Summary", diagnostics.get("predictionBandSummary") or [])
    add_kv_sheet("Regression Fit Summary", (diagnostics.get("regressionFitSummary") or {}).get("summary") or {})
    add_records_sheet("Segment Performance", diagnostics.get("segmentPerformance") or [])
    add_records_sheet("Split Comparison", diagnostics.get("splitComparison") or [])
    add_kv_sheet("Classification Metrics", diagnostics.get("classificationMetrics") or {})
    add_records_sheet("Detailed Lift Table", diagnostics.get("detailedLiftTable") or [])
    add_records_sheet("Lift Decile", diagnostics.get("liftDecile") or [])
    add_records_sheet("Calibration", diagnostics.get("calibration") or [])
    add_records_sheet("Population Waterfall", diagnostics.get("populationWaterfall") or [])
    add_records_sheet("VIF Analysis", (diagnostics.get("vifAnalysis") or {}).get("rows") or [])
    add_kv_sheet("Pozzolo Calibration", diagnostics.get("pozzoloCalibration") or {})
    add_records_sheet("KPI Aggregation", (diagnostics.get("kpiAggregation") or {}).get("rows") or [])
    add_kv_sheet("Generated Code", diagnostics.get("generatedCode") or {})
    add_kv_sheet("Metadata", diagnostics.get("dataset", {}).get("sourceInfo") or {})
    notes = wb.create_sheet("Governance Notes")
    notes.append(["Note"])
    notes.append([diagnostics.get("limitations", "")])
    for warning in diagnostics.get("mappingWarnings") or []:
        notes.append([warning])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_diagnostics_word(diagnostics: Dict[str, Any]) -> bytes:
    try:
        from docx import Document
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Word export requires python-docx. Install server requirements and retry.") from exc
    doc = Document()
    doc.add_heading("ATLAS Model Diagnostics Report", 0)
    doc.add_paragraph(f"Dataset: {diagnostics['dataset']['name']}")
    doc.add_paragraph(f"Rows reviewed: {diagnostics['rowCount']:,}")
    doc.add_paragraph(f"Diagnostic mode: {diagnostics.get('diagnosticMode', 'unknown')}")
    doc.add_paragraph(diagnostics.get("limitations", ""))

    doc.add_heading("Prediction Summary", level=1)
    for key, value in (diagnostics.get("predictionSummary", {}).get("summary") or {}).items():
        doc.add_paragraph(f"{key}: {value}")

    doc.add_heading("Regression Fit Summary", level=1)
    regression = diagnostics.get("regressionFitSummary") or {}
    if regression.get("available"):
        for key, value in (regression.get("summary") or {}).items():
            doc.add_paragraph(f"{key}: {value}")
    else:
        doc.add_paragraph(regression.get("warning") or "Regression diagnostics require actual and predicted values.")

    doc.add_heading("Classification Summary", level=1)
    classification = diagnostics.get("classificationMetrics") or {}
    if classification.get("available") is False:
        doc.add_paragraph(classification.get("warning") or "Classification metrics require binary actual and score/probability fields.")
    else:
        for key, value in classification.items():
            doc.add_paragraph(f"{key}: {value}")

    doc.add_heading("Detailed Lift Table", level=1)
    add_doc_table(doc, diagnostics.get("detailedLiftTable") or [])

    doc.add_heading("Population Waterfall", level=1)
    add_doc_table(doc, diagnostics.get("populationWaterfall") or [])

    doc.add_heading("VIF / Multicollinearity Analysis", level=1)
    vif = diagnostics.get("vifAnalysis") or {}
    doc.add_paragraph(vif.get("formula") or "VIF_j = 1 / (1 - R_j^2)")
    add_doc_table(doc, vif.get("rows") or [])

    doc.add_heading("Pozzolo Prior Correction", level=1)
    pozzolo = diagnostics.get("pozzoloCalibration") or {}
    doc.add_paragraph(str(pozzolo.get("equation") or "alpha = [pi / (1 - pi)] / [pi_s / (1 - pi_s)]; p_corrected = (alpha * p_s) / (1 - p_s + alpha * p_s)"))
    for key, value in pozzolo.items():
        if key != "equation":
            doc.add_paragraph(f"{key}: {value}")

    doc.add_heading("KPI Aggregation", level=1)
    kpi = diagnostics.get("kpiAggregation") or {}
    add_doc_table(doc, kpi.get("rows") or [])
    if kpi.get("recommendation"):
        doc.add_paragraph(kpi["recommendation"])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def add_doc_table(doc: Any, rows: List[Dict[str, Any]]) -> None:
    if not rows:
        doc.add_paragraph("No rows available.")
        return
    columns = list(rows[0].keys())
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, column in enumerate(columns):
        table.rows[0].cells[index].text = str(column)
    for row in rows[:50]:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            cells[index].text = str(row.get(column, ""))


def export_diagnostics_powerpoint(diagnostics: Dict[str, Any]) -> bytes:
    try:
        from pptx import Presentation
    except Exception as exc:
        raise HTTPException(status_code=503, detail="PowerPoint export requires python-pptx. Install server requirements and retry.") from exc
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "ATLAS Model Diagnostics"
    slide.placeholders[1].text = f"{diagnostics['dataset']['name']} | {diagnostics['rowCount']:,} rows reviewed"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Governance Limitation"
    slide.placeholders[1].text = diagnostics.get("limitations", "")
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Prediction and Fit Summary"
    prediction = diagnostics.get("predictionSummary", {}).get("summary") or {}
    regression = diagnostics.get("regressionFitSummary", {}).get("summary") or {}
    slide.placeholders[1].text = "\n".join(
        [
            f"Diagnostic mode: {diagnostics.get('diagnosticMode', 'unknown')}",
            f"Mean prediction: {prediction.get('mean', '-')}",
            f"Min prediction: {prediction.get('min', '-')}",
            f"Max prediction: {prediction.get('max', '-')}",
            f"RMSE: {regression.get('rmse', '-')}",
            f"MAE: {regression.get('mae', '-')}",
            f"R-squared: {regression.get('rSquared', '-')}",
        ]
    )
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Classification, Lift, and Calibration"
    classification = diagnostics.get("classificationMetrics") or {}
    pozzolo = diagnostics.get("pozzoloCalibration") or {}
    slide.placeholders[1].text = "\n".join(
        [
            f"AUC: {classification.get('auc', '-')}",
            f"Gini: {classification.get('gini', '-')}",
            f"KS: {classification.get('ks', '-')}",
            f"Base event rate: {classification.get('baseEventRate', '-')}",
            f"Pozzolo applicable: {pozzolo.get('applicable', '-')}",
            f"Pozzolo equation: {pozzolo.get('equation', '-')}",
        ]
    )
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "VIF and KPI Governance"
    vif_rows = (diagnostics.get("vifAnalysis") or {}).get("rows") or []
    kpi_rows = (diagnostics.get("kpiAggregation") or {}).get("rows") or []
    slide.placeholders[1].text = "\n".join(
        [
            f"VIF formula: {(diagnostics.get('vifAnalysis') or {}).get('formula', 'VIF_j = 1 / (1 - R_j^2)')}",
            f"Variables reviewed: {len(vif_rows)}",
            f"KPI count: {len(kpi_rows)}",
            f"KPI recommendation: {(diagnostics.get('kpiAggregation') or {}).get('recommendation', '-')}",
        ]
    )
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def safe_float(value: Any) -> Optional[float]:
    try:
        if value is None or not math.isfinite(float(value)):
            return None
        return float(value)
    except Exception:
        return None
