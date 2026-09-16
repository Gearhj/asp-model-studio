"""Compile the governed MLB treatment workbook into canonical source JSON.

The generated JSON is a source-recommendation artifact, not a project policy.
Project and Gold reconciliation happens in the Stage 3 service before a draft
may be persisted.  The source workbook is opened read-only and is never saved.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook


COMPILER_VERSION = "MLB_FEATURE_TREATMENT_COMPILER_V1"
SPECIFICATION_KEY = "MLB_FEATURE_TREATMENT_SPEC_V1"
PRIMARY_SHEET = "Feature Treatment Matrix"
SUPPORTING_SHEETS = (
    "Executive Summary",
    "Priority Exceptions",
    "Treatment Archetypes",
    "Raw Missing Export",
)
REQUIRED_COLUMNS = (
    "Feature",
    "Family",
    "Data Type",
    "Core Gold?",
    "Semantic Class",
    "Feature Meaning",
    "Missing %",
    "Primary Missing Reason",
    "No-Prior-History Rows",
    "Insufficient-History Rows",
    "Source-Unavailable Rows",
    "Current Platform Method",
    "Recommended Inclusion",
    "Outlier Detection",
    "Outlier Treatment",
    "Valid / Logical Domain",
    "Missing-Value Treatment",
    "Required Indicator / Cold-Start Handling",
    "Fold-Local / PIT Requirement",
    "Model Handling",
    "Priority",
    "Confidence",
    "Change Required?",
    "Scientific Rationale",
    "Reviewer Decision",
    "Reviewer Notes",
)
FIELD_NAMES = {
    "Feature": "feature_name",
    "Family": "feature_family",
    "Data Type": "data_type",
    "Core Gold?": "core_status",
    "Semantic Class": "semantic_class",
    "Feature Meaning": "feature_meaning",
    "Missing %": "missing_percent",
    "Primary Missing Reason": "primary_missing_reason",
    "No-Prior-History Rows": "no_prior_history_rows",
    "Insufficient-History Rows": "insufficient_history_rows",
    "Source-Unavailable Rows": "source_unavailable_rows",
    "Current Platform Method": "current_platform_method",
    "Recommended Inclusion": "recommended_inclusion",
    "Outlier Detection": "outlier_detection",
    "Outlier Treatment": "outlier_treatment",
    "Valid / Logical Domain": "logical_domain",
    "Missing-Value Treatment": "missing_treatment",
    "Required Indicator / Cold-Start Handling": "required_indicator_or_cold_start_rule",
    "Fold-Local / PIT Requirement": "fold_local_or_pit_requirement",
    "Model Handling": "model_handling",
    "Priority": "priority",
    "Confidence": "confidence",
    "Change Required?": "change_required",
    "Scientific Rationale": "scientific_rationale",
    "Reviewer Decision": "reviewer_decision",
    "Reviewer Notes": "reviewer_notes",
}
EXPECTED_DISPOSITION_COUNTS = {
    "EXCLUDE_UNTIL_SOURCE_REPAIRED": 4,
    "EXCLUDE_UNTIL_COVERAGE_REPAIRED": 2,
    "RECODE_AFTER_SEMANTIC_VALIDATION": 1,
    "RECODE_AFTER_CALENDAR_VALIDATION": 1,
    "HIERARCHICAL_PRIOR_PLUS_INDICATOR": 23,
    "COLD_START_EXPOSURE_ENCODING": 2,
    "ROLE_OR_LEAGUE_PRIOR_PLUS_INDICATOR": 1,
    "NO_PRIOR_GAME_STATE_PLUS_FOLD_LOCAL_FILL": 22,
    "RECOMPUTE_AFTER_PARENT_TREATMENT": 8,
    "STRUCTURAL_ZERO_PLUS_INDICATOR": 8,
    "INCLUDE_DERIVED_INDICATOR": 42,
    "SIMULATOR_REQUIRED": 5,
    "NO_MISSING_TREATMENT": 121,
}


class WorkbookSpecificationError(ValueError):
    """The workbook does not satisfy the governed source contract."""


def canonical_json_bytes(value: Mapping[str, Any]) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        allow_nan=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")


def _normalized_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise WorkbookSpecificationError("Workbook contains a non-finite numeric cell.")
        return int(value) if value.is_integer() else value
    return str(value)


def _sheet_dimensions(sheet: Any) -> dict[str, Any]:
    return {
        "name": sheet.title,
        "max_row": int(sheet.max_row),
        "max_column": int(sheet.max_column),
    }


def _rows_until_separator(values: Iterable[tuple[Any, ...]]) -> tuple[list[str], list[dict[str, Any]]]:
    iterator = iter(values)
    try:
        raw_header = next(iterator)
    except StopIteration as exc:
        raise WorkbookSpecificationError(f"{PRIMARY_SHEET} is empty.") from exc
    header = [str(value).strip() if value is not None else "" for value in raw_header]
    missing = [name for name in REQUIRED_COLUMNS if name not in header]
    if missing:
        raise WorkbookSpecificationError("Missing required columns: " + ", ".join(missing))
    if len(header) != len(REQUIRED_COLUMNS) or tuple(header) != REQUIRED_COLUMNS:
        raise WorkbookSpecificationError(
            f"{PRIMARY_SHEET} must expose the exact ordered {len(REQUIRED_COLUMNS)}-column contract."
        )
    records: list[dict[str, Any]] = []
    for row_number, raw_row in enumerate(iterator, start=2):
        source = dict(zip(header, raw_row, strict=False))
        feature_name = _normalized_cell(source.get("Feature"))
        if feature_name is None:
            break
        record = {
            FIELD_NAMES[column]: _normalized_cell(source.get(column))
            for column in REQUIRED_COLUMNS
        }
        raw_missing_percent = record["missing_percent"]
        if isinstance(raw_missing_percent, (int, float)):
            # The governed matrix stores this Excel percentage as a fraction
            # (1.0 renders as 100.00%).  Persist the human/API unit explicitly
            # so no consumer can misread a fully unavailable field as 1%.
            record["missing_percent"] = round(float(raw_missing_percent) * 100.0, 12)
        record["source_sheet"] = PRIMARY_SHEET
        record["source_row"] = row_number
        records.append(record)
    return header, records


def compile_workbook(path: str | Path) -> dict[str, Any]:
    workbook_path = Path(path).expanduser().resolve()
    workbook_bytes = workbook_path.read_bytes()
    workbook_sha = hashlib.sha256(workbook_bytes).hexdigest()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        required_sheets = (PRIMARY_SHEET, *SUPPORTING_SHEETS)
        absent_sheets = [name for name in required_sheets if name not in workbook.sheetnames]
        if absent_sheets:
            raise WorkbookSpecificationError("Missing required sheets: " + ", ".join(absent_sheets))
        matrix = workbook[PRIMARY_SHEET]
        header, records = _rows_until_separator(matrix.iter_rows(values_only=True))
        names = [str(record["feature_name"]) for record in records]
        duplicates = sorted(name for name, count in Counter(names).items() if count > 1)
        blank_names = [record["source_row"] for record in records if not record["feature_name"]]
        if len(records) != 240:
            raise WorkbookSpecificationError(
                f"Expected exactly 240 feature rows; parsed {len(records)}."
            )
        if duplicates:
            raise WorkbookSpecificationError("Duplicate feature names: " + ", ".join(duplicates))
        if blank_names:
            raise WorkbookSpecificationError(
                "Blank feature names at rows: " + ", ".join(map(str, blank_names))
            )
        dispositions = Counter(str(record["recommended_inclusion"] or "") for record in records)
        invalid = sorted(set(dispositions) - set(EXPECTED_DISPOSITION_COUNTS))
        if invalid:
            raise WorkbookSpecificationError("Unrecognized dispositions: " + ", ".join(invalid))
        if dict(dispositions) != EXPECTED_DISPOSITION_COUNTS:
            raise WorkbookSpecificationError(
                "Workbook disposition counts differ from the governed 240-row expectation."
            )
        nonblank_reviews = [
            record["feature_name"]
            for record in records
            if record["reviewer_decision"] is not None or record["reviewer_notes"] is not None
        ]
        if nonblank_reviews:
            raise WorkbookSpecificationError(
                "Reviewer Decision and Reviewer Notes must remain blank source fields: "
                + ", ".join(map(str, nonblank_reviews))
            )
        source_payload = {
            "schema_version": "MLB_FEATURE_TREATMENT_SOURCE_V1",
            "specification_key": SPECIFICATION_KEY,
            "compiler_version": COMPILER_VERSION,
            "source_workbook": {
                "filename": workbook_path.name,
                "size_bytes": len(workbook_bytes),
                "sha256": workbook_sha,
                "primary_sheet": PRIMARY_SHEET,
                "sheet_names": list(workbook.sheetnames),
                "sheets": [_sheet_dimensions(workbook[name]) for name in workbook.sheetnames],
                "required_columns": list(header),
            },
            "expected_disposition_counts": dict(EXPECTED_DISPOSITION_COUNTS),
            "features": records,
        }
        content_sha = hashlib.sha256(canonical_json_bytes(source_payload)).hexdigest()
        return {
            **source_payload,
            "source_content_sha256": content_sha,
            "source_artifact_id": f"{SPECIFICATION_KEY}:sha256:{content_sha}",
        }
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    compiled = compile_workbook(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(compiled, indent=2, ensure_ascii=False, allow_nan=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(compiled["source_artifact_id"])
    print(compiled["source_workbook"]["sha256"])
    print(len(compiled["features"]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
