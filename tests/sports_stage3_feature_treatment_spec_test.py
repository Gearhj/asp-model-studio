from __future__ import annotations

from collections import Counter
from copy import deepcopy
import hashlib
from pathlib import Path
import sqlite3
from types import SimpleNamespace
from typing import Any

from openpyxl import Workbook, load_workbook
import pytest

from scripts.compile_mlb_feature_treatment_spec import (
    FIELD_NAMES,
    PRIMARY_SHEET,
    REQUIRED_COLUMNS,
    SUPPORTING_SHEETS,
    WorkbookSpecificationError,
    compile_workbook,
)
from server.persistence.db import Database
from server.persistence.repositories.stage3_feature_set_repository import (
    NO_CURRENT_FEATURE_SET,
)
from server.persistence.repositories.stage3_treatment_repository import (
    Stage3TreatmentPersistenceError,
    Stage3TreatmentRepository,
)
from server.sports_current_statcast import stage3_feature_engineering as feature_engineering
from server.sports_current_statcast import stage3_feature_treatment_spec as specification
from server.sports_current_statcast import (
    stage3_interaction_recomputation_authority as interaction_authority,
)
from server.sports_current_statcast import stage3_player_prior_authority as player_prior
from server.sports_current_statcast import stage3_treatments as treatments
from server.sports_current_statcast.contextual_features import (
    ORDERED_GOLD_PLUS_CONTEXT_FEATURES,
)
from server.sports_gold.contracts import GOLD_FEATURES


PROJECT_ID = "asp-mtlozhdx"
WORKBOOK_SHA256 = "0e2d5f2b0d495c941ae97fb84c216f0b5ccdfe904519db0acb616d1543a66679"


def _source() -> dict[str, Any]:
    return specification.load_source_specification()


def _write_source_workbook(
    path: Path,
    source: dict[str, Any],
    *,
    duplicate_first_feature: bool = False,
) -> None:
    """Build a hermetic workbook with the production 240-row source contract."""

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet_names = [PRIMARY_SHEET, *SUPPORTING_SHEETS]
    for sheet_name in sheet_names:
        workbook.create_sheet(sheet_name)
    matrix = workbook[PRIMARY_SHEET]
    matrix.append(list(REQUIRED_COLUMNS))
    records = deepcopy(source["features"])
    if duplicate_first_feature:
        records[1]["feature_name"] = records[0]["feature_name"]
    for record in records:
        row: list[Any] = []
        for column in REQUIRED_COLUMNS:
            value = record[FIELD_NAMES[column]]
            if column == "Missing %" and value is not None:
                value = float(value) / 100.0
            row.append(value)
        matrix.append(row)
    matrix.append([None] * len(REQUIRED_COLUMNS))
    matrix.append(["Scientific footer; not a feature row"])
    for sheet_name in SUPPORTING_SHEETS:
        workbook[sheet_name].append([sheet_name])
    workbook.save(path)
    workbook.close()


def _extra_family(name: str) -> str:
    if name.endswith("_home_x_rivalry"):
        return "INTERACTION"
    if name.endswith("_rivalry_mapping_missing"):
        return "MISSINGNESS"
    return "DIVISION_RIVALRY"


def _backend_rows(source: dict[str, Any]) -> list[dict[str, Any]]:
    source_by_name = {str(row["feature_name"]): row for row in source["features"]}
    rows: list[dict[str, Any]] = []
    for order, name in enumerate(ORDERED_GOLD_PLUS_CONTEXT_FEATURES, start=1):
        scientific = source_by_name.get(name)
        if scientific is not None:
            family = str(scientific["feature_family"])
            data_type = str(scientific["data_type"])
            core = str(scientific["core_status"]).lower() == "yes"
            missing_percent = float(scientific["missing_percent"])
            disposition = str(scientific["recommended_inclusion"])
            role = (
                "SIMULATION_GENERATED_STATE"
                if disposition == "SIMULATOR_REQUIRED"
                else "MISSINGNESS_INDICATOR"
                if disposition == "INCLUDE_DERIVED_INDICATOR"
                else "PREGAME_FEATURE"
            )
            eligible = True
            governance_status = "ELIGIBLE"
        else:
            family = _extra_family(name)
            data_type = (
                "categorical"
                if name.endswith(("_rivalry_pair", "_rivalry_mapping_missing"))
                else "numeric"
            )
            core = False
            missing_percent = 0.0
            role = (
                "MISSINGNESS_INDICATOR"
                if name.endswith("_rivalry_mapping_missing")
                else "PREGAME_FEATURE"
            )
            # These two indicators are the important adversarial case: the
            # pre-spec inventory exposed them, so the combined authority must
            # explicitly hold them closed rather than silently omit them.
            eligible = name.endswith("_rivalry_mapping_missing")
            governance_status = (
                "ELIGIBLE" if eligible else "INACTIVE_PENDING_RIVALRY_APPROVAL"
            )
        rows.append(
            {
                "order": order,
                "name": name,
                "featureFamily": family,
                "dataType": data_type,
                "role": role,
                "core": core,
                "contextual": name not in GOLD_FEATURES,
                "binary": bool(
                    scientific
                    and scientific["semantic_class"]
                    in {"Binary indicator", "Missingness indicator"}
                ),
                "simulationGeneratedState": bool(
                    scientific
                    and scientific["recommended_inclusion"] == "SIMULATOR_REQUIRED"
                ),
                "governanceStatus": governance_status,
                "eligible": eligible,
                "availabilityStatus": "AVAILABLE_IN_EXACT_CONTEXTUAL_GOLD",
                "sourceLineage": f"accepted-current-gold::{name}",
                "pointInTimeTreatment": "source.game_date < target.game_date",
                "coverageDevelopment": 1.0 - missing_percent / 100.0,
                "missingPercentDevelopment": missing_percent,
            }
        )
    return rows


def _authority(tmp_path: Path) -> tuple[SimpleNamespace, Path]:
    gold_path = tmp_path / "immutable-current-gold.parquet"
    gold_path.write_bytes(b"immutable-current-gold-sentinel")
    public = {
        "projectId": PROJECT_ID,
        "sourceDatasetId": "asp.mlb.statcast.current-pa.v1",
        "sourceManifestSha": "1" * 64,
        "goldDatasetId": "gold-table:gold_mlb_current_pa_v1",
        "goldManifestSha": "2" * 64,
        "contextualDatasetId": "gold-table:gold_mlb_current_pa_context_v2",
        "contextualManifestSha": "3" * 64,
        "exclusionMembershipArtifactSha": "4" * 64,
        "reviewDigest": "5" * 64,
    }
    authority = SimpleNamespace(
        project_id=PROJECT_ID,
        registration_id="registration-current",
        dataset_version_id="current-pa-v1",
        activation_revision=7,
        source_dataset_id=public["sourceDatasetId"],
        source_manifest_sha=public["sourceManifestSha"],
        source_publication_sha="6" * 64,
        gold_manifest_sha=public["goldManifestSha"],
        contextual_manifest_sha=public["contextualManifestSha"],
        review_digest=public["reviewDigest"],
        paths=(gold_path,),
        schema=SimpleNamespace(
            names=tuple(ORDERED_GOLD_PLUS_CONTEXT_FEATURES)
            + ("game_date", "outcome", "canonicalPaId")
        ),
        public=lambda: dict(public),
    )
    return authority, gold_path


@pytest.fixture()
def governed(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> SimpleNamespace:
    source = _source()
    authority, gold_path = _authority(tmp_path)
    holder: dict[str, Any] = {
        "authority": authority,
        "rows": _backend_rows(source),
    }

    def resolve(_database: Database, project_id: str) -> SimpleNamespace:
        assert project_id == PROJECT_ID
        return holder["authority"]

    def feature_catalog(_authority_value: object):
        rows = deepcopy(holder["rows"])
        return rows, {str(row["name"]): row for row in rows}

    def treatment_catalog(_authority_value: object):
        return {
            str(row["name"]): {
                "name": row["name"],
                "dataType": row["dataType"],
                "role": row["role"],
                "category": row["featureFamily"],
                "featureFamily": row["featureFamily"],
                "selectable": bool(row["eligible"]),
            }
            for row in holder["rows"]
        }

    monkeypatch.setattr(specification.eda, "_resolve_authority", resolve)
    monkeypatch.setattr(feature_engineering, "_catalog", feature_catalog)
    monkeypatch.setattr(treatments, "_catalog", treatment_catalog)
    monkeypatch.setattr(
        specification.eda,
        "_scan",
        lambda *_args, **_kwargs: pytest.fail(
            "Treatment-spec compilation/drafting must not scan development or final-OOT rows."
        ),
    )
    database = Database(tmp_path / "stage3-feature-treatment-spec.db")
    database.initialize()
    return SimpleNamespace(
        database=database,
        authority=authority,
        authority_holder=holder,
        inventory_holder=holder,
        source=source,
        gold_path=gold_path,
        authority_key=treatments._authority_payload(authority)["authorityKey"],
    )


def _by_name(compiled: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {str(row["feature_name"]): row for row in compiled["features"]}


def _save(governed: SimpleNamespace, checksum: str | None = None) -> dict[str, Any]:
    compiled = specification.compile_project_specification(
        governed.database, PROJECT_ID
    )
    return specification.save_treatment_specification_draft(
        governed.database,
        PROJECT_ID,
        {
            "authorityKey": governed.authority_key,
            "expectedManifestChecksum": checksum or compiled["manifestChecksum"],
        },
        actor_user_id="scientist-1",
    )


def test_compiler_parses_exact_unique_240_row_contract_and_persists_workbook_hash(
    tmp_path: Path,
) -> None:
    """Backend properties 1, 2, and 6 (source artifact side)."""

    source = _source()
    assert len(source["features"]) == 240
    assert len({row["feature_name"] for row in source["features"]}) == 240
    assert source["source_workbook"]["sha256"] == WORKBOOK_SHA256
    assert source["source_workbook"]["required_columns"] == list(REQUIRED_COLUMNS)

    workbook_path = tmp_path / "feature-treatment.xlsx"
    _write_source_workbook(workbook_path, source)
    compiled = compile_workbook(workbook_path)
    assert len(compiled["features"]) == 240
    assert len({row["feature_name"] for row in compiled["features"]}) == 240
    assert compiled["source_workbook"]["sha256"] == hashlib.sha256(
        workbook_path.read_bytes()
    ).hexdigest()


def test_compiler_rejects_duplicate_keys_and_changed_workbook_gets_new_identity(
    tmp_path: Path,
) -> None:
    """Backend properties 2 and 28 (workbook identity side)."""

    source = _source()
    duplicate_path = tmp_path / "duplicate.xlsx"
    _write_source_workbook(duplicate_path, source, duplicate_first_feature=True)
    with pytest.raises(WorkbookSpecificationError, match="Duplicate feature names"):
        compile_workbook(duplicate_path)

    workbook_path = tmp_path / "changed.xlsx"
    _write_source_workbook(workbook_path, source)
    before = compile_workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    matrix = workbook[PRIMARY_SHEET]
    rationale_column = list(REQUIRED_COLUMNS).index("Scientific Rationale") + 1
    matrix.cell(row=2, column=rationale_column).value += " Reviewed source revision."
    workbook.save(workbook_path)
    workbook.close()
    after = compile_workbook(workbook_path)
    assert after["source_workbook"]["sha256"] != before["source_workbook"]["sha256"]
    assert after["source_artifact_id"] != before["source_artifact_id"]


def test_exact_reconciliation_has_no_fuzzy_path_and_holds_all_extras_closed(
    governed: SimpleNamespace,
) -> None:
    """Backend properties 3, 4, and 5."""

    compiled = specification.compile_project_specification(
        governed.database, PROJECT_ID
    )
    reconciliation = compiled["reconciliation"]
    assert reconciliation["workbookFeatureCount"] == 240
    assert reconciliation["backendFeatureCount"] == 254
    assert reconciliation["exactMatchCount"] == 240
    assert reconciliation["unmatchedWorkbookFeatures"] == []
    assert reconciliation["fuzzyMatchesUsed"] == 0
    assert reconciliation["fatal"] is False
    additional = reconciliation["additionalBackendFeatures"]
    assert [row["featureName"] for row in additional] == list(
        specification.EXPECTED_ADDITIONAL_FEATURES
    )
    assert len(additional) == 14
    assert all(row["executableSelectable"] is False for row in additional)
    exposed_mapping_indicators = [row for row in additional if row["backendEligibility"]]
    assert {row["featureName"] for row in exposed_mapping_indicators} == {
        "batting_rivalry_mapping_missing",
        "fielding_rivalry_mapping_missing",
    }
    assert all(
        row["executableResolution"] == "HELD_CLOSED_NO_TREATMENT_INFERRED"
        for row in additional
    )
    assert all(
        row["authorityAudit"]["recommendation"] == "KEEP_HELD_CLOSED"
        and row["authorityAudit"]["mappingActive"] is False
        and row["authorityAudit"]["currentlyActive"] is False
        and row["authorityAudit"]["selectable"] is False
        and row["authorityAudit"]["pointInTimeAvailability"]
        == "BLOCKED_NO_APPROVED_EFFECTIVE_DATED_MAPPING"
        and row["authorityAudit"]["sourceRights"]
        == "OPERATIONAL_USE_NOT_ESTABLISHED"
        for row in additional
    )
    audits = {row["featureName"]: row["authorityAudit"] for row in additional}
    formulas = {
        "rivalry_pair": (
            "None if mapping_missing else int({team_id, opponent_team_id} in approved_pairs)"
        ),
        "rivalry_game": "rivalry_pair",
        "home_rivalry_game": (
            "None if rivalry_pair is None else rivalry_pair * int(is_home)"
        ),
        "friday_night_home_rivalry": (
            "None if rivalry_pair is None else rivalry_pair * int(is_home) * "
            "int(local_day_of_week == FRIDAY) * int(local_start_hour >= 18)"
        ),
        "weekend_home_rivalry": (
            "None if rivalry_pair is None else rivalry_pair * int(is_home) * "
            "int(local_day_of_week in {SATURDAY, SUNDAY})"
        ),
        "home_x_rivalry": (
            "None if rivalry_pair is None else int(is_home) * rivalry_pair"
        ),
        "rivalry_mapping_missing": (
            "int(rivalry_mapping_authority_status != APPROVED)"
        ),
    }
    assert all(
        audit["mappingSource"] == "docs/MLB_RIVALRY_MAPPING_PROPOSAL_V1.json"
        and audit["mappingStatus"] == "PROPOSED_NOT_APPROVED"
        and audit["mappingSourceSha256"]
        == "5bd1ad26acacdc32a1ad9edb377c72559402e161a32278971e3d07748b588d74"
        and audit["source"]
        == "MLB_TEAM_GAME_CONTEXT_FEATURES_V1_RESEARCH_ONLY_EVIDENCE"
        and audit["formulaEvidenceOnly"] is True
        for audit in audits.values()
    )
    for side in ("batting", "fielding"):
        indicator = audits[f"{side}_rivalry_mapping_missing"]
        assert indicator["side"] == side.upper()
        assert indicator["formula"] == formulas["rivalry_mapping_missing"]
        assert indicator["constructObjectivity"] == "OBJECTIVE_GOVERNANCE_STATE_INDICATOR"
        assert indicator["physicalMissingPercentDevelopment"] == 0.0
        assert indicator["coveragePercentDevelopment"] == 100.0
        assert set(indicator["coveragePercentBySeason"]) == {
            "2023",
            "2024",
            "2025",
            "2026",
        }
        assert set(indicator["coveragePercentBySeason"].values()) == {100.0}
        assert indicator["pairedMissingnessIndicator"] is None
        assert indicator["exactDuplicateFeature"] is None
        assert indicator["overlapsObjectiveOpponentContext"] is False
        assert indicator["requiresCuratedRivalryTable"] is False
        for suffix in (
            "rivalry_pair",
            "rivalry_game",
            "home_rivalry_game",
            "friday_night_home_rivalry",
            "weekend_home_rivalry",
            "home_x_rivalry",
        ):
            substantive = audits[f"{side}_{suffix}"]
            assert substantive["side"] == side.upper()
            assert substantive["formula"] == formulas[suffix]
            assert substantive["constructObjectivity"] == "SUBJECTIVE_RIVALRY_MEMBERSHIP"
            assert substantive["physicalMissingPercentDevelopment"] == 100.0
            assert substantive["coveragePercentDevelopment"] == 0.0
            assert set(substantive["coveragePercentBySeason"]) == {
                "2023",
                "2024",
                "2025",
                "2026",
            }
            assert set(substantive["coveragePercentBySeason"].values()) == {0.0}
            assert substantive["pairedMissingnessIndicator"] == (
                f"{side}_rivalry_mapping_missing"
            )
            expected_duplicate = {
                "rivalry_pair": f"{side}_rivalry_game",
                "rivalry_game": f"{side}_rivalry_pair",
                "home_rivalry_game": f"{side}_home_x_rivalry",
                "home_x_rivalry": f"{side}_home_rivalry_game",
            }.get(suffix)
            assert substantive["exactDuplicateFeature"] == expected_duplicate
            assert substantive["overlapsObjectiveOpponentContext"] is True
            assert substantive["requiresCuratedRivalryTable"] is True

    adversarial_source = deepcopy(governed.source)
    adversarial_source["features"][0]["feature_name"] = (
        str(adversarial_source["features"][0]["feature_name"]).upper()
    )
    rejected = specification._reconcile(governed.authority, adversarial_source)
    assert rejected["exactMatchCount"] == 239
    assert rejected["fuzzyMatchesUsed"] == 0
    assert rejected["fatal"] is True
    assert rejected["unmatchedWorkbookFeatures"] == [
        "FIELDING_HOURS_SINCE_PREVIOUS_GAME_END"
    ]


def test_project_compilation_is_deterministic_and_inventory_changes_identity(
    governed: SimpleNamespace,
) -> None:
    """Backend properties 7 and 28 (feature-inventory identity side)."""

    first = specification.compile_project_specification(governed.database, PROJECT_ID)
    second = specification.compile_project_specification(governed.database, PROJECT_ID)
    assert second["manifestChecksum"] == first["manifestChecksum"]
    assert second["specificationId"] == first["specificationId"]
    assert second == first

    changed_rows = deepcopy(governed.inventory_holder["rows"])
    changed_rows[0]["sourceLineage"] += "::new-governed-version"
    governed.inventory_holder["rows"] = changed_rows
    changed = specification.compile_project_specification(governed.database, PROJECT_ID)
    assert changed["reconciliation"]["fatal"] is False
    assert changed["featureInventorySha256"] != first["featureInventorySha256"]
    assert changed["manifestChecksum"] != first["manifestChecksum"]
    assert changed["specificationId"] != first["specificationId"]


def test_unexpected_backend_extra_fails_closed_without_compiler_exception(
    governed: SimpleNamespace,
) -> None:
    unexpected = deepcopy(governed.inventory_holder["rows"][-1])
    unexpected.update(
        {
            "order": len(governed.inventory_holder["rows"]) + 1,
            "name": "unexpected_backend_feature",
            "sourceLineage": "unexpected-test-authority",
        }
    )
    governed.inventory_holder["rows"].append(unexpected)

    compiled = specification.compile_project_specification(
        governed.database, PROJECT_ID
    )

    assert compiled["canCreateDraft"] is False
    assert compiled["reconciliation"]["fatal"] is True
    unexpected_row = next(
        row
        for row in compiled["reconciliation"]["additionalBackendFeatures"]
        if row["featureName"] == "unexpected_backend_feature"
    )
    assert unexpected_row["classification"] == "UNMATCHED_BACKEND_FEATURE"
    assert unexpected_row["executableSelectable"] is False
    assert "authorityAudit" not in unexpected_row
    assert compiled["rivalryAuthorityAudit"]["fieldCount"] == 14
    assert set(compiled["rivalryAuthorityAudit"]["fields"]) == set(
        specification.EXPECTED_ADDITIONAL_FEATURES
    )


def test_semantic_execution_mapping_is_fail_closed_and_never_uses_blanket_iqr(
    governed: SimpleNamespace,
) -> None:
    """Backend properties 11 through 20."""

    compiled = specification.compile_project_specification(governed.database, PROJECT_ID)
    rows = _by_name(compiled)
    outlier_input, missing_input = specification._policy_inputs(compiled)
    assert outlier_input["globalMethod"] == "NO_STATISTICAL_DETECTION"
    assert outlier_input["globalTreatment"] == "RETAIN"

    source_blocked = {
        "batting_hours_since_previous_game_end",
        "fielding_hours_since_previous_game_end",
        "batting_altitude_change_feet",
        "fielding_altitude_change_feet",
    }
    coverage_blocked = {
        "batting_previous_game_duration_minutes",
        "fielding_previous_game_duration_minutes",
    }
    for name in source_blocked | coverage_blocked:
        assert rows[name]["executable_selectable"] is False
        assert missing_input["overrides"][name]["method"] == "DROP_VARIABLE"
        assert missing_input["overrides"][name]["method"] not in {
            "MEAN",
            "MEDIAN",
            "MODE",
            "CONSTANT",
        }

    doubleheader = rows["doubleheader_game_number"]
    holiday = rows["holiday_name"]
    assert doubleheader["executable_status"] == "BLOCKED"
    doubleheader_audit = doubleheader["execution_contract"]["semanticAuthorityDecision"]
    assert doubleheader_audit == {
        "status": "BLOCKED_POINT_IN_TIME_UNPROVEN",
        "resolution": "HOLD_UNTIL_PREGAME_SCHEDULE_SEQUENCE_AUTHORITY_IS_REGISTERED",
        "blockedReason": (
            "Schedule authority MLB_CURRENT_STATCAST_SCHEDULE_AUTHORITY_V1 does not "
            "preserve original pregame knowledge timestamps or an independent "
            "{0,1,2} doubleheader-sequence invariant tied to scheduled start order."
        ),
        "sourceAuthority": "MLB_CURRENT_STATCAST_SCHEDULE_AUTHORITY_V1",
        "sourceAuthoritySha256": (
            "44c43a48dda5895dcc19417a1c51b18ccdaed9dd83b3c1b8d50826e42be69c4d"
        ),
        "requiredSemantics": {
            "0": "VERIFIED_NOT_A_DOUBLEHEADER",
            "1": "VERIFIED_FIRST_DOUBLEHEADER_GAME",
            "2": "VERIFIED_SECOND_DOUBLEHEADER_GAME",
            "unknown": "HELD_NONEXECUTABLE",
        },
        "forbiddenTreatment": "MEDIAN_IMPUTATION",
    }
    assert missing_input["overrides"]["doubleheader_game_number"]["method"] != "MEDIAN"
    assert holiday["executable_status"] == "BLOCKED"
    holiday_audit = holiday["execution_contract"]["semanticAuthorityDecision"]
    assert holiday_audit == {
        "status": "BLOCKED_CALENDAR_AUTHORITY_UNAVAILABLE",
        "resolution": (
            "HOLD_UNTIL_EFFECTIVE_DATED_VENUE_REGION_CALENDAR_AUTHORITY_IS_REGISTERED"
        ),
        "blockedReason": (
            "US_CA_ON_BASEBALL_CALENDAR_V1 is embedded code rather than a separately "
            "registered, sourced, effective-dated holiday calendar; null currently "
            "cannot prove NONE rather than UNKNOWN_HOLIDAY for every venue region."
        ),
        "sourceAuthority": "MLB_CONTEXT_STATIC_AUTHORITY_V1",
        "sourceAuthoritySha256": (
            "da9740dd90a161202d42ba18d3c89f2569e3711e33b48f86bf3bd0a8a6f2d56d"
        ),
        "embeddedRulesVersion": "US_CA_ON_BASEBALL_CALENDAR_V1",
        "requiredSemantics": {
            "holiday": "EXACT_REGISTERED_HOLIDAY_NAME",
            "none": "VERIFIED_NON_HOLIDAY",
            "unknown": "UNKNOWN_HOLIDAY_HELD_NONEXECUTABLE",
        },
        "forbiddenTreatment": "MODE_IMPUTATION",
    }
    assert missing_input["overrides"]["holiday_name"]["method"] != "MODE"

    binary = [row for row in compiled["features"] if row["semantic_class"] == "Binary indicator"]
    categorical = [row for row in compiled["features"] if row["data_type"] == "categorical"]
    assert len(binary) == 73
    assert all(
        row["execution_contract"]["outlier"]["detector"]
        == "REGISTERED_BOOLEAN_DOMAIN_VALIDATION"
        for row in binary
    )
    assert all(
        (outlier_input["overrides"].get(row["feature_name"]) or {}).get("treatment")
        not in {"CAP_IQR", "WINSORIZE"}
        for row in categorical
    )
    ordinary_categories = [
        row
        for row in categorical
        if row["semantic_class"]
        not in {"Binary indicator", "Missingness indicator"}
        and row["recommended_inclusion"]
        not in {
            "EXCLUDE_UNTIL_SOURCE_REPAIRED",
            "EXCLUDE_UNTIL_COVERAGE_REPAIRED",
        }
    ]
    assert all(
        row["execution_contract"]["outlier"]["detector"]
        == "ALLOWED_CATEGORY_SEMANTIC_VALIDATION"
        for row in ordinary_categories
    )

    cyclic = [row for row in compiled["features"] if row["semantic_class"] == "Cyclic encoding"]
    assert len(cyclic) == 6
    assert all(row["logical_domain"] == "[-1,1]" for row in cyclic)
    assert all(
        row["execution_contract"]["outlier"]["detector"]
        == "HARD_BOUNDS_MINUS_ONE_TO_ONE"
        for row in cyclic
    )
    rates = [row for row in compiled["features"] if row["semantic_class"] == "Rate / skill estimate"]
    assert len(rates) == 21
    assert all(
        str(row["logical_domain"]).startswith(("[0,", ">=0"))
        or str(row["logical_domain"]).startswith("Registered semantic domain")
        for row in rates
    )
    assert all(
        row["execution_contract"]["outlier"]["detector"]
        == "PHYSICAL_BOUNDS_AND_DENOMINATOR_VALIDATION"
        for row in rates
    )

    simulator = [row for row in compiled["features"] if row["recommended_inclusion"] == "SIMULATOR_REQUIRED"]
    assert len(simulator) == 5
    assert all(row["executable_resolution"] == "REQUIRE_VALID_SIMULATOR_STATE_NO_IMPUTATION" for row in simulator)
    assert all(
        row["execution_contract"]["outlier"]["detector"]
        == "STATE_MACHINE_VALIDATION"
        for row in simulator
    )
    assert all(
        missing_input["overrides"][row["feature_name"]]["method"]
        not in {"MEAN", "MEDIAN", "MODE", "CONSTANT", "PREDICTED"}
        for row in simulator
    )

    interactions = [
        row
        for row in compiled["features"]
        if row["recommended_inclusion"] == "RECOMPUTE_AFTER_PARENT_TREATMENT"
    ]
    assert len(interactions) == 8
    assert tuple(row["feature_name"] for row in interactions) == (
        interaction_authority.INTERACTION_FEATURES
    )
    assert all(row["interaction_parent_lineage"] for row in interactions)
    assert all(len(row["execution_contract"]["lineage"]["parentFeatures"]) == 2 for row in interactions)
    assert all(
        row["execution_contract"]["lineage"]["transformationOrder"]
        == [
            "TREAT_PARENTS_INSIDE_TRAINING_FOLD",
            "RECOMPUTE_INTERACTION",
            "APPLY_FROZEN_PIPELINE",
        ]
        for row in interactions
    )
    assert all(
        row["executable_status"] == "READY_GOVERNED_INTERACTION_V1"
        and row["executable_selectable"] is True
        and row["executable_resolution"]
        == "EXECUTE_PARENT_FIRST_FOLD_LOCAL_INTERACTION_RECOMPUTATION_V1"
        and row["blocked_reason"] is None
        for row in interactions
    )
    for row in interactions:
        authority = row["execution_contract"]["interactionAuthority"]
        expected = interaction_authority.feature_binding(row["feature_name"])
        assert expected is not None
        assert authority["formula"] == expected["formula"]
        assert authority["orderedParents"] == expected["orderedParents"]
        assert authority["parentIdentities"] == expected["orderedParents"]
        assert authority["parentTreatmentAuthorities"] == [
            parent["parentTreatmentAuthority"] for parent in expected["parents"]
        ]
        assert authority["operationOrder"] == list(interaction_authority.OPERATION_ORDER)
        assert authority["independentImputation"] == "PROHIBITED"
        assert authority["independentCapping"] == "PROHIBITED"
        assert authority["authorityId"] == compiled["interactionAuthority"]["authorityId"]
        assert missing_input["overrides"][row["feature_name"]][
            "parentFirstRecomputationRequired"
        ] is True
        assert missing_input["overrides"][row["feature_name"]][
            "independentInteractionImputation"
        ] is False
        assert outlier_input["overrides"][row["feature_name"]][
            "independentInteractionCapping"
        ] is False

    _save(governed)
    for blocked_name in (
        "fielding_hours_since_previous_game_end",
        "batting_previous_game_duration_minutes",
    ):
        with pytest.raises(feature_engineering.Stage3FeatureEngineeringError) as caught:
            feature_engineering.save_feature_set_draft(
                PROJECT_ID,
                governed.database,
                {
                    "authorityKey": governed.authority_key,
                    "featureSetKey": f"blocked-{blocked_name}",
                    "name": f"Blocked {blocked_name}",
                    "orderedFeatures": [blocked_name],
                    "expectedChecksum": NO_CURRENT_FEATURE_SET,
                },
                "scientist-1",
            )
        assert caught.value.code == "FEATURE_SET_INACTIVE_FEATURE"


def test_fold_local_model_family_and_indicator_contracts_are_explicit(
    governed: SimpleNamespace,
) -> None:
    """Backend properties 21 through 25."""

    compiled = specification.compile_project_specification(governed.database, PROJECT_ID)
    outlier_input, missing_input = specification._policy_inputs(compiled)
    _outlier_method, _outlier_policy, outlier_fitting, _ = treatments._validate_policy_payload(
        governed.database, governed.authority, "OUTLIER", outlier_input
    )
    _missing_method, _missing_policy, missing_fitting, _ = treatments._validate_policy_payload(
        governed.database, governed.authority, "IMPUTATION", missing_input
    )
    for fitting in (outlier_fitting, missing_fitting):
        assert fitting["fitScope"] == "TRAINING_ROWS_OF_EACH_CHRONOLOGICAL_FOLD_ONLY"
        assert fitting["validationTransform"] == "PARAMETERS_FROM_CORRESPONDING_TRAINING_FOLD_ONLY"
        assert fitting["finalOotParameterFittingProhibited"] is True
        assert fitting["targetOutcomeInputsProhibited"] is True

    fallbacks = [
        row
        for row in compiled["features"]
        if row["executable_status"] == "READY_WITH_EXPLICIT_FALLBACK"
    ]
    assert len(fallbacks) == 22
    assert all(
        row["model_family_resolution"]["tree_native_missing"]
        != row["model_family_resolution"]["dense_matrix"]
        for row in fallbacks
    )
    assert all(
        row["model_family_resolution"]["fallback_label"]
        == "INTERIM_TECHNICAL_FALLBACK"
        for row in fallbacks
    )

    governed_prior_rows = [
        row
        for row in compiled["features"]
        if row["recommended_inclusion"]
        in {
            "HIERARCHICAL_PRIOR_PLUS_INDICATOR",
            "ROLE_OR_LEAGUE_PRIOR_PLUS_INDICATOR",
        }
    ]
    assert len(governed_prior_rows) == 24
    ready_prior_rows = [
        row
        for row in governed_prior_rows
        if row["executable_status"] == "READY_GOVERNED_PRIOR_V1"
    ]
    blocked_prior_rows = [
        row for row in governed_prior_rows if row["executable_status"] == "BLOCKED"
    ]
    assert len(ready_prior_rows) == 11
    assert len(blocked_prior_rows) == 13
    assert all(
        row["execution_contract"]["missing"]["strategy"]
        == "GOVERNED_PLAYER_TO_MLB_LEAGUE_BETA_BINOMIAL_PRIOR_V1"
        for row in ready_prior_rows
    )
    assert all(row["executable_selectable"] for row in ready_prior_rows)
    assert all(
        row["execution_contract"]["missing"]["strategy"]
        == "GOVERNED_PRIOR_REQUIRED_EXACT_SUFFICIENT_STATISTIC_UNAVAILABLE"
        for row in blocked_prior_rows
    )
    assert all(not row["executable_selectable"] for row in blocked_prior_rows)
    assert all(row["blocked_reason"] for row in blocked_prior_rows)

    indicators = [
        row
        for row in compiled["features"]
        if row["recommended_inclusion"] == "INCLUDE_DERIVED_INDICATOR"
    ]
    assert len(indicators) == 42
    assert all(row["exact_name_match"] for row in indicators)
    assert all(row["source_or_lineage_identity"] for row in indicators)
    assert all(
        row["executable_resolution"]
        == "DETERMINISTIC_BOOLEAN_RECOMPUTATION_NO_STATISTICAL_IMPUTATION"
        for row in indicators
    )
    # A physically complete indicator column is not evidence that the flag is
    # beneficial.  Preserve the workbook's exact semantics: prior-state flags
    # can describe genuine baseball state or cold start, while every indicator
    # remains subject to source-era/extraction/pipeline drift review.
    assert all(row["missing_percent"] == 0 for row in indicators)
    assert all(
        "No statistical imputation" in row["missing_treatment"]
        for row in indicators
    )
    assert all(
        "monitor source-era leakage and drift" in row["model_handling"]
        for row in indicators
    )
    indicator_names = {row["feature_name"] for row in indicators}
    assert {
        "batting_previous_game_missing",
        "fielding_previous_game_missing",
    } <= indicator_names
    prior_state_or_cold_start = {
        name
        for name in indicator_names
        if "_previous_" in name or "recent_bullpen_workload" in name
    }
    source_extraction_or_pipeline = indicator_names - prior_state_or_cold_start
    assert len(prior_state_or_cold_start) == 20
    assert len(source_extraction_or_pipeline) == 22
    paired = {
        indicator
        for row in compiled["features"]
        for indicator in row["execution_contract"]["lineage"]["pairedIndicators"]
    }
    exact_names = {row["feature_name"] for row in compiled["features"]}
    assert paired
    assert paired <= exact_names

    unsupported = specification._execution_contract(
        "unregistered_feature", "UNREGISTERED_RECOMMENDATION"
    )
    assert unsupported["executableStatus"] == "BLOCKED"
    assert unsupported["selectable"] is False
    assert unsupported["fallbackLabel"] is None
    assert "no registered executable mapping" in unsupported["blockedReason"]


def test_combined_draft_is_append_only_idempotent_and_does_not_touch_gold_or_oot(
    governed: SimpleNamespace,
) -> None:
    """Backend properties 6, 9, 10, and 27."""

    gold_before = hashlib.sha256(governed.gold_path.read_bytes()).hexdigest()
    first = _save(governed)["data"]
    draft_summary = first["draft"]
    draft = Stage3TreatmentRepository(governed.database).combined_spec_state(
        PROJECT_ID, governed.authority.review_digest
    )["latestDraft"]
    assert draft_summary["status"] == "DRAFT"
    assert draft_summary["specificationDraftId"] == draft["specificationId"]
    assert draft["sourceWorkbookSha256"] == WORKBOOK_SHA256
    assert draft["manifestChecksum"] == first["specification"]["manifestChecksum"]
    assert draft["goldManifestSha"] == governed.authority.gold_manifest_sha
    assert draft["finalOot"]["locked"] is True
    assert draft["manifest"]["finalOotIdentity"]["targetsAccessed"] is False
    assert draft["immutable"] is True
    assert first["audit"] == {
        "sourceRowsMutated": 0,
        "goldRowsMutated": 0,
        "finalOotRowsOpened": 0,
        "finalOotTargetsAccessed": False,
        "preprocessingFits": 0,
        "modelFits": 0,
        "scores": 0,
        "calibrations": 0,
        "promotions": 0,
        "featureSetsRegistered": 0,
        "gameScoringJobs": 0,
        "historicalArtifactMutations": 0,
        "providerCalls": 0,
    }

    second = _save(governed)["data"]
    assert second["draft"]["specificationDraftId"] == draft["specificationId"]
    assert second["draft"]["checksum"] == draft["checksum"]
    with governed.database.connection() as connection:
        spec_count = connection.execute(
            "SELECT COUNT(*) FROM stage3_feature_treatment_spec_versions"
        ).fetchone()[0]
        event_count = connection.execute(
            "SELECT COUNT(*) FROM stage3_feature_treatment_spec_events"
        ).fetchone()[0]
    assert spec_count == 1
    assert event_count == 1
    assert hashlib.sha256(governed.gold_path.read_bytes()).hexdigest() == gold_before

    with governed.database.connection() as connection:
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            connection.execute(
                "UPDATE stage3_feature_treatment_spec_versions SET notes = 'changed' "
                "WHERE feature_treatment_spec_id = ?",
                (draft["specificationId"],),
            )
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            connection.execute(
                "DELETE FROM stage3_feature_treatment_spec_versions "
                "WHERE feature_treatment_spec_id = ?",
                (draft["specificationId"],),
            )


def test_existing_adopted_policies_remain_exactly_immutable(
    governed: SimpleNamespace,
) -> None:
    """Backend property 8."""

    repository = Stage3TreatmentRepository(governed.database)
    authority = treatments._persistence_authority(governed.authority)
    fitting = {
        "fitScope": "TRAINING_ROWS_OF_EACH_CHRONOLOGICAL_FOLD_ONLY",
        "validationTransform": "PARAMETERS_FROM_CORRESPONDING_TRAINING_FOLD_ONLY",
        "finalOotParameterFittingProhibited": True,
        "targetOutcomeInputsProhibited": True,
    }
    adopted_ids: dict[str, str] = {}
    for kind in ("OUTLIER", "IMPUTATION"):
        draft = repository.save_draft(
            authority=authority,
            policy_kind=kind,
            global_method="DOMAIN_BOUNDS" if kind == "OUTLIER" else "DATA_TYPE_AWARE",
            policy={"legacyAdoptedSentinel": kind},
            fitting_rule=fitting,
            notes="pre-existing adopted policy",
            actor_user_id="prior-scientist",
        )
        adopted = repository.adopt(
            project_id=PROJECT_ID,
            policy_kind=kind,
            authority_review_digest=governed.authority.review_digest,
            draft_policy_id=str(draft["policyId"]),
            actor_user_id="prior-scientist",
        )
        adopted_ids[kind] = str(adopted["policyId"])

    with governed.database.connection() as connection:
        before = {
            row["treatment_policy_id"]: tuple(row)
            for row in connection.execute(
                "SELECT * FROM stage3_treatment_policy_versions "
                "WHERE policy_status = 'ADOPTED' ORDER BY treatment_policy_id"
            ).fetchall()
        }
    result = _save(governed)["data"]
    assert result["adoptionPerformed"] is False
    assert result["adoptedSpecification"] is None
    assert {
        result["priorPolicies"]["outlier"]["policyId"],
        result["priorPolicies"]["missing"]["policyId"],
    } == set(adopted_ids.values())
    with governed.database.connection() as connection:
        after = {
            row["treatment_policy_id"]: tuple(row)
            for row in connection.execute(
                "SELECT * FROM stage3_treatment_policy_versions "
                "WHERE policy_status = 'ADOPTED' ORDER BY treatment_policy_id"
            ).fetchall()
        }
    assert after == before


def test_combined_draft_blocks_every_individual_policy_adoption(
    governed: SimpleNamespace,
) -> None:
    saved = _save(governed)["data"]["draft"]
    repository = Stage3TreatmentRepository(governed.database)
    child_ids = {
        "OUTLIER": saved["outlierPolicyRef"]["policyId"],
        "IMPUTATION": saved["missingPolicyRef"]["policyId"],
    }

    for kind, policy_id in child_ids.items():
        with pytest.raises(
            Stage3TreatmentPersistenceError, match="unadopted combined"
        ):
            repository.adopt(
                project_id=PROJECT_ID,
                policy_kind=kind,
                authority_review_digest=governed.authority.review_digest,
                draft_policy_id=policy_id,
                actor_user_id="scientist-1",
            )

    with governed.database.connection() as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM stage3_treatment_policy_versions "
            "WHERE policy_status = 'ADOPTED'"
        ).fetchone()[0] == 0

    authority = treatments._persistence_authority(governed.authority)
    fitting_rule = {
        "fitScope": "TRAINING_ROWS_OF_EACH_CHRONOLOGICAL_FOLD_ONLY",
        "validationTransform": "PARAMETERS_FROM_CORRESPONDING_TRAINING_FOLD_ONLY",
        "finalOotParameterFittingProhibited": True,
        "targetOutcomeInputsProhibited": True,
    }
    for kind in ("OUTLIER", "IMPUTATION"):
        legacy = repository.save_draft(
            authority=authority,
            policy_kind=kind,
            global_method=("DOMAIN_BOUNDS" if kind == "OUTLIER" else "DATA_TYPE_AWARE"),
            policy={"unrelatedLegacyDraft": kind},
            fitting_rule=fitting_rule,
            notes="unrelated legacy draft",
            actor_user_id="legacy-scientist",
        )
        with pytest.raises(
            Stage3TreatmentPersistenceError, match="combined feature-treatment"
        ):
            repository.adopt(
                project_id=PROJECT_ID,
                policy_kind=kind,
                authority_review_digest=governed.authority.review_digest,
                draft_policy_id=str(legacy["policyId"]),
                actor_user_id="legacy-scientist",
            )

    with governed.database.connection() as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM stage3_treatment_policy_versions "
            "WHERE policy_status = 'ADOPTED'"
        ).fetchone()[0] == 0


def test_feature_engineering_draft_overlay_exposes_prior_outputs_without_authorizing_them(
    governed: SimpleNamespace,
) -> None:
    saved = _save(governed)["data"]
    assert saved["draft"]["status"] == "DRAFT"
    assert saved["canAdopt"] is False
    overlay, reference = specification.current_draft_execution_overlay(
        governed.database, governed.authority
    )
    assert reference is not None
    assert reference["status"] == "DRAFT_NEGATIVE_GATES_ONLY"
    assert reference["priorAuthority"] == saved["priorAuthority"]
    assert reference["priorAuthority"]["authorityName"] == player_prior.AUTHORITY_NAME
    assert reference["priorAuthority"]["readyFeatureCount"] == 11
    assert reference["priorAuthority"]["blockedFeatureCount"] == 13
    assert reference["interactionAuthority"] == saved["interactionAuthority"]
    assert reference["interactionAuthority"]["authorityName"] == (
        interaction_authority.AUTHORITY_NAME
    )
    assert reference["interactionAuthority"]["readyFeatureCount"] == 8
    assert reference["interactionAuthority"]["blockedFeatureCount"] == 0

    ready_interaction_name = interaction_authority.INTERACTION_FEATURES[0]
    ready_interaction = overlay[ready_interaction_name]
    assert ready_interaction["executableStatus"] == "READY_GOVERNED_INTERACTION_V1"
    assert ready_interaction["executableSelectable"] is True
    assert ready_interaction["interactionAuthority"]["formula"] == (
        interaction_authority.feature_binding(ready_interaction_name)["formula"]
    )
    assert ready_interaction["interactionAuthority"]["authorityId"] == (
        reference["interactionAuthority"]["authorityId"]
    )

    ready_name = "batter_so_per_pa"
    blocked_name = "batter_avg"
    ready = overlay[ready_name]
    blocked = overlay[blocked_name]
    expected_ready_binding = player_prior.feature_binding(ready_name)
    expected_blocked_binding = player_prior.feature_binding(blocked_name)
    assert ready["executableStatus"] == "READY_GOVERNED_PRIOR_V1"
    assert ready["executableSelectable"] is True
    assert ready["priorAuthority"]["authorityId"] == reference["priorAuthority"]["authorityId"]
    assert ready["priorAuthority"]["status"] == "READY_GOVERNED_PRIOR_V1"
    assert ready["priorAuthority"]["rawFeatureDecision"] == (
        "RETAIN_RAW_AND_ADD_GOVERNED_DERIVED_FAMILY"
    )
    assert ready["derivedPriorOutputs"] == expected_ready_binding["derivedOutputs"]
    assert ready["derivedPriorOutputs"] == ready["priorAuthority"]["derivedOutputMap"]
    assert set(ready["derivedPriorOutputs"]) == {
        "posteriorMean",
        "effectiveN",
        "priorStrength",
        "posteriorStrength",
        "posteriorVariance",
        "posteriorStd",
        "posteriorReliability",
        "historyState",
        "priorLevel",
        "observedSuccesses",
        "observedTrials",
        "priorAlpha",
        "priorBeta",
        "posteriorAlpha",
        "posteriorBeta",
        "decayWeight",
        "observationAgeDays",
    }
    # Derived preprocessing columns are lineage, not silently registered raw candidates.
    assert not (set(ready["derivedPriorOutputs"].values()) & set(overlay))

    assert blocked["executableStatus"] == "BLOCKED"
    assert blocked["executableSelectable"] is False
    assert blocked["blockedReason"] == expected_blocked_binding["blockedReason"]
    assert blocked["priorAuthority"]["status"] == "BLOCKED_REQUIRES_PRIOR_AUTHORITY"
    assert blocked["priorAuthority"]["blockedReason"] == expected_blocked_binding["blockedReason"]

    # A positive DRAFT resolution is informational only: it cannot reactivate a
    # row already held closed by the underlying feature authority.
    initially_closed = {
        "name": ready_name,
        "eligible": False,
        "featureState": "INACTIVE",
        "governanceStatus": "HELD_CLOSED_BY_SOURCE_AUTHORITY",
        "inactiveReason": "Pre-existing source hold.",
        "stateReason": "Pre-existing source hold.",
    }
    returned_reference = feature_engineering._apply_treatment_spec_draft_gates(
        [initially_closed], governed.database, governed.authority
    )
    assert returned_reference == reference
    assert initially_closed["treatmentSpecificationDraft"] == reference
    assert initially_closed["treatmentSpecificationResolution"] == ready
    assert initially_closed["eligible"] is False
    assert initially_closed["featureState"] == "INACTIVE"
    assert initially_closed["governanceStatus"] == "HELD_CLOSED_BY_SOURCE_AUTHORITY"
    assert initially_closed["inactiveReason"] == "Pre-existing source hold."

    # A negative DRAFT resolution still performs its existing fail-closed duty.
    blocked_row = deepcopy(
        next(row for row in governed.inventory_holder["rows"] if row["name"] == blocked_name)
    )
    feature_engineering._apply_treatment_spec_draft_gates(
        [blocked_row], governed.database, governed.authority
    )
    assert blocked_row["eligible"] is False
    assert blocked_row["featureState"] == "INACTIVE"
    assert blocked_row["governanceStatus"] == "HELD_CLOSED_BY_TREATMENT_SPEC_DRAFT"
    assert blocked_row["inactiveReason"] == expected_blocked_binding["blockedReason"]

    assert saved["audit"]["preprocessingFits"] == 0
    assert saved["audit"]["modelFits"] == 0
    with governed.database.connection() as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM stage3_feature_set_versions"
        ).fetchone()[0] == 0
        assert connection.execute(
            "SELECT COUNT(*) FROM stage3_feature_set_events"
        ).fetchone()[0] == 0


def test_feature_engineering_overlay_rejects_a_stale_interaction_authority(
    governed: SimpleNamespace, monkeypatch: pytest.MonkeyPatch
) -> None:
    _save(governed)
    original_configuration = interaction_authority.authority_configuration

    def changed_configuration(*, overrides=None):
        configuration = original_configuration(overrides=overrides)
        configuration["thresholds"]["longTravelMiles"] = 1_501.0
        return configuration

    monkeypatch.setattr(
        interaction_authority, "authority_configuration", changed_configuration
    )
    with pytest.raises(specification.Stage3FeatureTreatmentSpecError) as caught:
        specification.current_draft_execution_overlay(
            governed.database, governed.authority
        )
    assert caught.value.code == "TREATMENT_SPEC_DRAFT_STALE"


def test_stale_gold_authority_cannot_apply_prior_compilation(
    governed: SimpleNamespace,
) -> None:
    """Backend property 26."""

    compiled = specification.compile_project_specification(governed.database, PROJECT_ID)
    stale_key = governed.authority_key
    replacement = deepcopy(governed.authority)
    replacement.gold_manifest_sha = "9" * 64
    replacement_public = replacement.public()
    replacement_public["goldManifestSha"] = replacement.gold_manifest_sha
    replacement.public = lambda: dict(replacement_public)
    governed.authority_holder["authority"] = replacement

    with pytest.raises(treatments.Stage3TreatmentError) as caught:
        specification.save_treatment_specification_draft(
            governed.database,
            PROJECT_ID,
            {
                "authorityKey": stale_key,
                "expectedManifestChecksum": compiled["manifestChecksum"],
            },
            actor_user_id="scientist-1",
        )
    assert caught.value.code == "TREATMENT_AUTHORITY_STALE"


def test_disposition_counts_and_adoption_gate_are_source_derived(
    governed: SimpleNamespace,
) -> None:
    compiled = specification.compile_project_specification(governed.database, PROJECT_ID)
    assert compiled["dispositionCounts"] == dict(
        sorted(specification.EXPECTED_DISPOSITION_COUNTS.items())
    )
    assert sum(compiled["dispositionCounts"].values()) == 240
    assert Counter(
        row["scientific_recommendation"] for row in compiled["features"]
    ) == Counter(specification.EXPECTED_DISPOSITION_COUNTS)
    assert compiled["canCreateDraft"] is True
    assert compiled["canAdopt"] is False
    summary = compiled["executionSummary"]
    assert summary["totalReady"] == 219
    assert summary["interactionBlockersBefore"] == 8
    assert summary["interactionBlockersCleared"] == 8
    assert summary["interactionBlockersRemaining"] == 0
    assert summary["semanticFieldsCleared"] == 0
    assert summary["semanticFieldsRemaining"] == 2
    assert summary["rivalryFieldsHeldClosed"] == 14
    assert summary["currentTotalBlockers"] == 29
    assert compiled["interactionAuthority"]["featureCount"] == 8
    assert compiled["interactionAuthority"]["readyFeatureCount"] == 8
    assert compiled["interactionAuthority"]["blockedFeatureCount"] == 0
    with pytest.raises(specification.Stage3FeatureTreatmentSpecError) as caught:
        specification.adopt_treatment_specification(
            governed.database,
            PROJECT_ID,
            {"authorityKey": governed.authority_key},
            actor_user_id="scientist-1",
        )
    assert caught.value.code == "TREATMENT_SPEC_DRAFT_REQUIRED"

    saved = _save(governed)["data"]["draft"]
    base = {
        "authorityKey": governed.authority_key,
        "expectedManifestChecksum": saved["manifestChecksum"],
    }
    with pytest.raises(specification.Stage3FeatureTreatmentSpecError) as caught:
        specification.adopt_treatment_specification(
            governed.database,
            PROJECT_ID,
            {**base, "specificationDraftId": "not-the-current-draft"},
            actor_user_id="scientist-1",
        )
    assert caught.value.code == "TREATMENT_SPEC_DRAFT_NOT_CURRENT"

    with pytest.raises(specification.Stage3FeatureTreatmentSpecError) as caught:
        specification.adopt_treatment_specification(
            governed.database,
            PROJECT_ID,
            {
                "authorityKey": governed.authority_key,
                "specificationDraftId": saved["specificationDraftId"],
                "expectedManifestChecksum": "0" * 64,
            },
            actor_user_id="scientist-1",
        )
    assert caught.value.code == "TREATMENT_SPEC_DRAFT_CHECKSUM_MISMATCH"

    with pytest.raises(specification.Stage3FeatureTreatmentSpecError) as caught:
        specification.adopt_treatment_specification(
            governed.database,
            PROJECT_ID,
            {
                **base,
                "specificationDraftId": saved["specificationDraftId"],
            },
            actor_user_id="scientist-1",
        )
    assert caught.value.code == "TREATMENT_SPEC_ADOPTION_BLOCKED"
